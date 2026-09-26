"""The Studio (D-94): what it makes, the figure check on every item — distractors
and table cells included —, the generation that runs after the response, the
quota, the exports and their notices, and one student never reaching another's
artifacts."""

from __future__ import annotations

import csv
import io
from datetime import UTC, datetime, timedelta
from xml.etree import ElementTree

import pytest
from docx import Document
from openpyxl import load_workbook
from sqlalchemy import select

from app.ai.provider import AIProvider, AIUnavailableError
from app.ai.studio import CATALOG, build_tree, read_studio
from app.config import settings
from app.exporters import studio as studio_exporter
from app.exporters.report import LIMITATION_NOTICE
from app.exporters.studio import (
    AI_NOTICE,
    NO_ATTRIBUTION_LABEL,
    NO_URL_LABEL,
    render,
)
from app.models.notebook import NotebookSource, StudioArtifact
from app.notebooks import mindmap
from app.schemas.notebook import ArtifactOut
from app.services import studio_service

TEXT = (
    "O aço carbono tem densidade de 7850 kg/m³ e módulo de elasticidade de 210 GPa. "
    "É o material estrutural mais usado na construção civil.\n\n"
    "O alumínio 6061 tem densidade de 2700 kg/m³, cerca de um terço da do aço. "
    "Por isso aparece em estruturas leves, como quadros de bicicleta."
)


def _notebook(client, title="Materiais estruturais", text=TEXT) -> int:
    notebook = client.post("/api/notebooks", json={"title": title}).json()
    response = client.post(
        f"/api/notebooks/{notebook['id']}/sources/text", json={"title": "Aula 1", "text": text}
    )
    assert response.status_code == 201, response.text
    return notebook["id"]


def _generate(client, notebook_id: int, **payload) -> dict:
    """Start a generation and read it back. The test client runs the background
    task before ``post`` returns, so the artifact read after it is finished."""
    response = client.post(f"/api/notebooks/{notebook_id}/studio", json=payload)
    assert response.status_code == 202, response.text
    assert response.json()["status"] == "gerando"
    artifact_id = response.json()["id"]
    return client.get(f"/api/notebooks/{notebook_id}/studio/{artifact_id}").json()


class _Scripted(AIProvider):
    """Answers the Studio from a script, in the model's JSON shape, and reads
    it the way a real provider does."""

    name = "roteiro"
    simulated = True

    def __init__(self, answers: list[dict | Exception]) -> None:
        self.answers = answers
        self.calls = 0
        self.notes: list[str | None] = []

    def interpret(self, context):  # pragma: no cover - not used here
        raise NotImplementedError

    def explain(self, context):  # pragma: no cover - not used here
        raise NotImplementedError

    def studio(self, request):
        self.notes.append(request.retry_note)
        reply = self.answers[min(self.calls, len(self.answers) - 1)]
        self.calls += 1
        if isinstance(reply, Exception):
            raise reply
        return read_studio(request, reply)


def _script(monkeypatch, answers: list[dict | Exception]) -> _Scripted:
    provider = _Scripted(answers)
    monkeypatch.setattr(studio_service, "get_provider", lambda _settings: provider)
    return provider


# --- catalogue -------------------------------------------------------------------


def test_the_catalogue_is_the_one_truth_the_modal_reads(client):
    response = client.get("/api/notebooks/studio-catalog")
    assert response.status_code == 200
    tools = {tool["slug"]: tool for tool in response.json()["tools"]}
    assert list(tools) == ["report", "flashcards", "quiz", "table", "mindmap"]
    guide = next(t for t in tools["report"]["templates"] if t["slug"] == "guia_estudo")
    # What the pencil shows is what the model receives.
    assert guide["instructions"] == CATALOG["report"].template("guia_estudo").instructions
    properties = next(t for t in tools["table"]["templates"] if t["slug"] == "propriedades")
    assert "Valor" in properties["columns"]
    assert [c["amount"] for c in tools["flashcards"]["counts"]] == [10, 15, 25]
    assert tools["mindmap"]["exports"] == ["svg"]


# --- making each tool ------------------------------------------------------------


@pytest.mark.parametrize("tool", ["report", "flashcards", "quiz", "table", "mindmap"])
def test_every_tool_is_made_and_cited_with_the_simulated_provider(client, tool):
    notebook_id = _notebook(client)
    artifact = _generate(client, notebook_id, tool=tool)
    assert artifact["status"] == "pronto", artifact["error"]
    assert artifact["item_count"] and artifact["item_count"] > 0
    assert artifact["withheld"] == []
    # Renumbered 1..k in reading order, and every citation is a copy.
    assert [c["number"] for c in artifact["citations"]] == list(
        range(1, len(artifact["citations"]) + 1)
    )
    assert all(c["excerpt"] and c["source_title"] == "Aula 1" for c in artifact["citations"])
    assert artifact["exports"] == list(CATALOG[tool].exports)

    listed = client.get(f"/api/notebooks/{notebook_id}/studio").json()
    assert [a["id"] for a in listed["artifacts"]] == [artifact["id"]]
    assert listed["usage"]["used"] == 1


def test_the_mind_map_comes_with_the_layout_the_screen_and_the_svg_share(client):
    notebook_id = _notebook(client)
    artifact = _generate(client, notebook_id, tool="mindmap", format="detalhado")
    layout = artifact["layout"]
    nodes = {node["id"]: node for node in layout["nodes"]}
    assert nodes[1]["parent"] is None and nodes[1]["depth"] == 0
    assert len(layout["edges"]) == len(nodes) - 1
    for edge in layout["edges"]:
        parent, child = nodes[edge["source"]], nodes[edge["target"]]
        # An edge leaves its parent's right side and reaches its child's left.
        assert edge["x1"] == pytest.approx(parent["x"] + parent["width"], abs=0.2)
        assert edge["x2"] == pytest.approx(child["x"], abs=0.2)
        assert child["x"] > parent["x"]


def test_a_template_s_instruction_reaches_the_model_as_edited(client, monkeypatch):
    provider = _script(
        monkeypatch,
        [
            {
                "title": "Guia",
                "sections": [
                    {
                        "heading": "Aço",
                        "paragraphs": [{"text": "O aço tem 7850 kg/m³.", "citations": [1]}],
                    }
                ],
            }
        ],
    )
    captured = {}
    original = provider.studio

    def spy(request):
        captured["request"] = request
        return original(request)

    provider.studio = spy
    notebook_id = _notebook(client)
    artifact = _generate(
        client,
        notebook_id,
        tool="report",
        template="guia_estudo",
        instructions="Escreva para calouros.",
        topic="aço",
    )
    assert artifact["status"] == "pronto"
    request = captured["request"]
    assert request.instructions == "Escreva para calouros."
    assert request.topic == "aço"
    assert artifact["options"]["template"] == "guia_estudo"


# --- the figure check, item by item ----------------------------------------------

_GOOD_QUESTION = {
    "prompt": "Qual a densidade do aço carbono?",
    "options": ["7850 kg/m³", "2700 kg/m³", "Não consta", "Depende"],
    "answer_index": 0,
    "hint": "Veja a primeira aula.",
    "explanation": "A fonte diz 7850 kg/m³.",
    "citations": [1],
}
_BAD_DISTRACTOR = {
    **_GOOD_QUESTION,
    "prompt": "Qual o módulo do aço?",
    # 999 GPa is a wrong answer — and a figure no source states.
    "options": ["210 GPa", "999 GPa", "2700 kg/m³", "Nenhum"],
}


def test_a_distractor_with_an_invented_figure_gets_one_retry_naming_it(client, monkeypatch):
    provider = _script(
        monkeypatch,
        [
            {"title": "Teste", "questions": [_GOOD_QUESTION, _BAD_DISTRACTOR]},
            {"title": "Teste", "questions": [_GOOD_QUESTION]},
        ],
    )
    artifact = _generate(client, _notebook(client), tool="quiz")
    assert provider.calls == 2
    assert provider.notes[0] is None and "999" in provider.notes[1]
    assert artifact["status"] == "pronto"
    assert artifact["withheld"] == []
    assert len(artifact["content"]["questions"]) == 1


def test_a_distractor_invented_twice_is_left_out_and_said(client, monkeypatch):
    _script(monkeypatch, [{"title": "Teste", "questions": [_GOOD_QUESTION, _BAD_DISTRACTOR]}])
    artifact = _generate(client, _notebook(client), tool="quiz")
    assert [q["prompt"] for q in artifact["content"]["questions"]] == [_GOOD_QUESTION["prompt"]]
    assert artifact["withheld"] == [
        "Uma questão foi omitida porque citava números que não aparecem nos trechos "
        "citados: 999."
    ]


def test_a_figure_grounded_only_in_an_uncited_passage_does_not_pass(client, monkeypatch):
    card = {"front": "Densidade do aço?", "back": "7850 kg/m³.", "citations": []}
    _script(monkeypatch, [{"title": "Cartões", "cards": [card]}])
    artifact = _generate(client, _notebook(client), tool="flashcards")
    assert artifact["status"] == "falhou"
    assert "7850" in artifact["error"]


def test_a_table_cell_is_labelled_never_blank(client, monkeypatch):
    rows = [
        {
            "cells": [
                {"text": "Aço carbono", "citations": [1]},
                {"text": "7850 kg/m³", "citations": [1]},
                {"text": "", "citations": []},
            ]
        },
        {
            "cells": [
                {"text": "Alumínio 6061", "citations": [1]},
                {"text": "3100 kg/m³", "citations": [1]},
                {"text": "estruturas leves", "citations": [1]},
            ]
        },
    ]
    _script(monkeypatch, [{"title": "Densidades", "columns": [], "rows": rows}])
    artifact = _generate(
        client,
        _notebook(client),
        tool="table",
        template="personalizado",
        columns=["Material", "Densidade", "Uso"],
    )
    table = artifact["content"]
    assert table["columns"] == ["Material", "Densidade", "Uso"]
    first, second = (row["cells"] for row in table["rows"])
    assert first[2] == {"text": None, "citations": [], "status": "ausente"}
    assert second[1] == {"text": None, "citations": [], "status": "omitida"}
    assert second[2]["status"] == "ok"
    assert artifact["withheld"] == [
        "Uma célula ficou sem valor porque citava números que não aparecem nos trechos "
        "citados: 3100."
    ]


def test_a_mind_map_node_with_an_invented_figure_goes_with_its_branch(client, monkeypatch):
    nodes = [
        {"id": 1, "parent": 0, "label": "Materiais", "citations": []},
        {"id": 2, "parent": 1, "label": "Aço, 7850 kg/m³", "citations": [1]},
        {"id": 3, "parent": 1, "label": "Titânio, 4500 kg/m³", "citations": [1]},
        {"id": 4, "parent": 3, "label": "Aeronáutica", "citations": [1]},
    ]
    _script(monkeypatch, [{"title": "Mapa", "nodes": nodes}])
    artifact = _generate(client, _notebook(client), tool="mindmap")
    root = artifact["content"]["root"]
    assert [child["label"] for child in root["children"]] == ["Aço, 7850 kg/m³"]
    assert "4500" in artifact["withheld"][0]


def test_a_title_with_an_invented_figure_falls_back_to_the_tool_s_name(client, monkeypatch):
    card = {"front": "Densidade do aço?", "back": "7850 kg/m³.", "citations": [1]}
    _script(monkeypatch, [{"title": "Os 4321 fatos do aço", "cards": [card]}])
    artifact = _generate(client, _notebook(client), tool="flashcards")
    assert artifact["title"] == "Cartões didáticos"


def test_nothing_that_passes_is_a_failed_generation_and_costs_nothing(client, monkeypatch):
    _script(monkeypatch, [{"title": "Teste", "questions": [_BAD_DISTRACTOR]}])
    notebook_id = _notebook(client)
    artifact = _generate(client, notebook_id, tool="quiz")
    assert artifact["status"] == "falhou"
    assert "999" in artifact["error"]
    assert client.get(f"/api/notebooks/{notebook_id}/studio").json()["usage"]["used"] == 0


def test_the_provider_s_own_message_is_the_failure(client, monkeypatch):
    _script(monkeypatch, [AIUnavailableError("Limite do plano gratuito atingido (429).")])
    artifact = _generate(client, _notebook(client), tool="report")
    assert artifact["status"] == "falhou"
    assert artifact["error"] == "Limite do plano gratuito atingido (429)."


def test_a_busy_provider_on_the_retry_keeps_the_first_answer(client, monkeypatch):
    _script(
        monkeypatch,
        [
            {"title": "Teste", "questions": [_GOOD_QUESTION, _BAD_DISTRACTOR]},
            AIUnavailableError("sobrecarga"),
        ],
    )
    artifact = _generate(client, _notebook(client), tool="quiz")
    assert artifact["status"] == "pronto"
    assert len(artifact["content"]["questions"]) == 1
    assert "999" in artifact["withheld"][0]


# --- refusing choices ------------------------------------------------------------


@pytest.mark.parametrize(
    ("payload", "fragment"),
    [
        ({"tool": "flashcards", "template": "guia_estudo"}, "não tem modelo"),
        ({"tool": "report", "count": "mais"}, "não tem quantidade"),
        ({"tool": "report", "format": "slides"}, "Aceitos: corrido, topicos"),
        ({"tool": "report", "template": "personalizado"}, "precisa de instruções"),
        ({"tool": "quiz", "instructions": "Seja duro."}, "use o campo de foco"),
        ({"tool": "table", "columns": ["A", "a"]}, "mesmo nome"),
        ({"tool": "table", "columns": ["A", " "]}, "sem nome"),
        ({"tool": "report", "columns": ["A"]}, "não tem colunas"),
    ],
)
def test_a_choice_the_tool_does_not_have_is_refused(client, payload, fragment):
    response = client.post(f"/api/notebooks/{_notebook(client)}/studio", json=payload)
    assert response.status_code == 400
    assert fragment in response.json()["detail"]


def test_an_unknown_tool_is_refused(client):
    response = client.post(f"/api/notebooks/{_notebook(client)}/studio", json={"tool": "audio"})
    assert response.status_code == 422


def test_nothing_selected_is_refused_before_anything_runs(client):
    notebook_id = _notebook(client)
    client.put(f"/api/notebooks/{notebook_id}/sources/selection", json={"selected": False})
    response = client.post(f"/api/notebooks/{notebook_id}/studio", json={"tool": "report"})
    assert response.status_code == 400
    assert "Nenhuma fonte marcada" in response.json()["detail"]


# --- quota, running and stuck generations ----------------------------------------


def _running(db_session, notebook_id: int, minutes_ago: float = 0) -> StudioArtifact:
    artifact = StudioArtifact(
        notebook_id=notebook_id,
        tool="report",
        title="Relatório",
        status="gerando",
        source_ids=[],
        created_at=datetime.now(UTC) - timedelta(minutes=minutes_ago),
    )
    db_session.add(artifact)
    db_session.flush()
    return artifact


def test_the_daily_limit(client, monkeypatch):
    monkeypatch.setattr(settings, "notebook_daily_artifacts", 1)
    notebook_id = _notebook(client)
    assert _generate(client, notebook_id, tool="flashcards")["status"] == "pronto"
    response = client.post(f"/api/notebooks/{notebook_id}/studio", json={"tool": "quiz"})
    assert response.status_code == 429
    assert "gerações de hoje" in response.json()["detail"]


def test_a_running_generation_already_counts_against_the_day(client, db_session, monkeypatch):
    monkeypatch.setattr(settings, "notebook_daily_artifacts", 1)
    notebook_id = _notebook(client)
    _running(db_session, notebook_id)
    listed = client.get(f"/api/notebooks/{notebook_id}/studio").json()
    assert listed["usage"] == {"used": 0, "limit": 1, "remaining": 0}
    response = client.post(f"/api/notebooks/{notebook_id}/studio", json={"tool": "quiz"})
    assert response.status_code == 429


def test_too_many_at_once_is_a_conflict(client, db_session, monkeypatch):
    monkeypatch.setattr(settings, "notebook_studio_in_flight", 1)
    notebook_id = _notebook(client)
    _running(db_session, notebook_id)
    response = client.post(f"/api/notebooks/{notebook_id}/studio", json={"tool": "quiz"})
    assert response.status_code == 409
    assert "em andamento" in response.json()["detail"]


def test_a_stuck_generation_reads_as_failed_and_stops_counting(client, db_session, monkeypatch):
    monkeypatch.setattr(settings, "notebook_studio_in_flight", 1)
    notebook_id = _notebook(client)
    stuck = _running(db_session, notebook_id, minutes_ago=60)
    artifact = client.get(f"/api/notebooks/{notebook_id}/studio/{stuck.id}").json()
    assert artifact["status"] == "falhou"
    assert "interrompida" in artifact["error"]
    # Derived on read, never written by the GET.
    db_session.refresh(stuck)
    assert stuck.status == "gerando"
    assert _generate(client, notebook_id, tool="flashcards")["status"] == "pronto"


# --- editing, notes and deleting -------------------------------------------------


def test_rename_save_as_note_and_delete(client):
    notebook_id = _notebook(client)
    artifact = _generate(client, notebook_id, tool="flashcards")
    base = f"/api/notebooks/{notebook_id}/studio/{artifact['id']}"

    renamed = client.patch(base, json={"title": "Revisão de sexta"}).json()
    assert renamed["title"] == "Revisão de sexta"

    note = client.post(f"{base}/note").json()
    assert note["origin"] == "estudio"
    assert note["title"] == "Revisão de sexta"
    assert note["body"].startswith("Frente:") and "[1]" in note["body"]
    assert note["citations"][0]["number"] == 1

    assert client.delete(base).status_code == 204
    assert client.get(base).status_code == 404


def test_deleting_a_notebook_deletes_its_artifacts(client, db_session):
    notebook_id = _notebook(client)
    _generate(client, notebook_id, tool="report")
    assert client.delete(f"/api/notebooks/{notebook_id}").status_code == 204
    left = db_session.execute(
        select(StudioArtifact).where(StudioArtifact.notebook_id == notebook_id)
    ).all()
    assert left == []


# --- exports ---------------------------------------------------------------------


def _export(client, notebook_id: int, artifact: dict, fmt: str):
    return client.get(f"/api/notebooks/{notebook_id}/studio/{artifact['id']}/export.{fmt}")


def test_the_report_docx_carries_both_notices_and_the_references(client):
    notebook_id = _notebook(client)
    artifact = _generate(client, notebook_id, tool="report", template="guia_estudo")
    response = _export(client, notebook_id, artifact, "docx")
    assert response.status_code == 200
    assert (
        response.headers["content-disposition"]
        == 'attachment; filename="guia-de-estudo-materiais-estruturais.docx"'
    )
    text = "\n".join(p.text for p in Document(io.BytesIO(response.content)).paragraphs)
    assert LIMITATION_NOTICE in text and AI_NOTICE in text
    assert "Referências" in text and "[1] Aula 1" in text


def test_the_quiz_docx_puts_the_answer_key_after_the_questions(client):
    notebook_id = _notebook(client)
    artifact = _generate(client, notebook_id, tool="quiz", format="vf", count="menos")
    text = "\n".join(
        p.text
        for p in Document(
            io.BytesIO(_export(client, notebook_id, artifact, "docx").content)
        ).paragraphs
    )
    assert text.index("Questões") < text.index("Gabarito")
    assert "a) Verdadeiro" in text


def test_the_flashcards_csv_neutralises_a_formula_and_carries_the_notice(client):
    notebook_id = _notebook(client, text='=HYPERLINK("http://mal.example") é um texto hostil.')
    artifact = _generate(client, notebook_id, tool="flashcards")
    body = _export(client, notebook_id, artifact, "csv").content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(body)))
    assert [LIMITATION_NOTICE] in rows and [AI_NOTICE] in rows
    backs = [row[1] for row in rows if len(row) == 3 and row[0] != "Frente"]
    assert backs and all(not back.startswith("=") for back in backs)
    assert backs[0].startswith("'=")


def test_the_table_xlsx_writes_absence_in_words(client):
    notebook_id = _notebook(client)
    artifact = _generate(client, notebook_id, tool="table", template="propriedades")
    workbook = load_workbook(io.BytesIO(_export(client, notebook_id, artifact, "xlsx").content))
    cover = [row[0] for row in workbook["Aviso"].iter_rows(values_only=True)]
    assert LIMITATION_NOTICE in cover and AI_NOTICE in cover
    values = [cell for row in workbook["Tabela"].iter_rows(values_only=True) for cell in row]
    # The simulated provider fills two columns; the other three are silence.
    assert "não consta nas fontes" in values
    assert "Referências" in workbook.sheetnames


def test_the_mind_map_svg_escapes_everything_and_is_served_inert(client):
    notebook_id = _notebook(client, title='<script>alert("x")</script>')
    artifact = _generate(client, notebook_id, tool="mindmap")
    response = _export(client, notebook_id, artifact, "svg")
    assert response.status_code == 200
    assert response.headers["content-security-policy"] == "default-src 'none'"
    assert response.headers["content-type"].startswith("image/svg+xml")
    svg = response.text
    assert "<script>" not in svg and "&lt;script&gt;" in svg
    assert "Esta ferramenta destina-se a apoio didático" in svg


def test_a_format_the_tool_does_not_export_is_refused_with_the_accepted_ones(client):
    notebook_id = _notebook(client)
    artifact = _generate(client, notebook_id, tool="report")
    response = _export(client, notebook_id, artifact, "csv")
    assert response.status_code == 400
    assert response.json()["detail"] == "Relatório não se exporta em CSV. Aceitos: DOCX."


def test_an_artifact_that_is_not_ready_does_not_export(client, db_session):
    notebook_id = _notebook(client)
    running = _running(db_session, notebook_id)
    response = _export(client, notebook_id, {"id": running.id}, "docx")
    assert response.status_code == 400


# --- exports: an external source's address and credit (D-97) ---------------------

WIKI_URL = "https://pt.wikipedia.org/wiki/A%C3%A7o"
WIKI_CREDIT = (
    "Texto de “Aço”, da Wikipédia em português, escrito por colaboradores da "
    "Wikipédia e publicado sob a licença CC BY-SA 4.0 "
    "(https://creativecommons.org/licenses/by-sa/4.0/); revisão 123, obtido em 25/09/2026."
)
#: Formula, markup and an XML-illegal control character, all at once.
HOSTILE = '=HYPERLINK("http://mal.example")<script>alert(1)</script></w:t>\x0b&amp;'


def _external(client, db_session, tool: str, url=WIKI_URL, attribution=WIKI_CREDIT, **extra):
    """An artifact made from a source carrying a ``meta`` as the external-source
    service stores it — nothing leaves for the network."""
    notebook_id = _notebook(client)
    source = db_session.scalars(
        select(NotebookSource).where(NotebookSource.notebook_id == notebook_id)
    ).one()
    source.meta = {"url": url, "attribution": attribution, "license": "CC BY-SA 4.0"}
    db_session.flush()
    return notebook_id, _generate(client, notebook_id, tool=tool, **extra)


def _with_credit(artifact: dict, url=None, attribution=None, only_first=False) -> ArtifactOut:
    """The artifact, its citations given (or stripped of) an address and credit."""
    citations = []
    for index, citation in enumerate(artifact["citations"]):
        credited = not only_first or index == 0
        citations.append(
            {
                **citation,
                "source_url": url if credited else None,
                "source_attribution": attribution if credited else None,
            }
        )
    return ArtifactOut.model_validate({**artifact, "citations": citations})


def _csv_rows(data: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))


def _references(rows: list[list[str]]) -> list[list[str]]:
    """The rows after the "Referências" title of a CSV export."""
    start = rows.index(["Referências"])
    return rows[start + 1 :]


def _docx_text(data: bytes) -> str:
    return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)


def test_the_address_and_credit_survive_into_the_stored_artifact_and_its_docx(client, db_session):
    notebook_id, artifact = _external(client, db_session, "report", template="guia_estudo")
    assert artifact["citations"]
    assert all(c["source_url"] == WIKI_URL for c in artifact["citations"])
    assert all(c["source_attribution"] == WIKI_CREDIT for c in artifact["citations"])
    text = _docx_text(_export(client, notebook_id, artifact, "docx").content)
    assert LIMITATION_NOTICE in text and AI_NOTICE in text
    references = text[text.index("Referências") :]
    # On lines of their own, after the reference and before the excerpt.
    first = artifact["citations"][0]
    assert f"[1] Aula 1\n{WIKI_URL}\n{WIKI_CREDIT}\n{first['excerpt']}" in references


def test_the_spreadsheets_add_an_address_and_a_credit_column(client, db_session):
    notebook_id, artifact = _external(client, db_session, "flashcards")
    rows = _csv_rows(_export(client, notebook_id, artifact, "csv").content)
    assert [LIMITATION_NOTICE] in rows and [AI_NOTICE] in rows
    header, *body = _references(rows)
    assert header == ["Nº", "Fonte", "Trecho citado", "Endereço", "Atribuição"]
    assert body and all(row[3:] == [WIKI_URL, WIKI_CREDIT] for row in body)

    notebook_id, artifact = _external(client, db_session, "table", template="propriedades")
    workbook = load_workbook(io.BytesIO(_export(client, notebook_id, artifact, "xlsx").content))
    cover = [row[0] for row in workbook["Aviso"].iter_rows(values_only=True)]
    assert LIMITATION_NOTICE in cover and AI_NOTICE in cover
    header, *body = list(workbook["Referências"].iter_rows(values_only=True))
    assert list(header) == ["Nº", "Fonte", "Trecho citado", "Endereço", "Atribuição"]
    assert body and all(list(row[3:]) == [WIKI_URL, WIKI_CREDIT] for row in body)


def test_a_citation_without_them_is_written_in_words_beside_one_with_them(client):
    notebook_id = _notebook(client)
    artifact = _generate(client, notebook_id, tool="quiz", count="menos")
    first = artifact["citations"][0]
    # A second citation from a phase 1 source, beside the external one.
    plain = {**first, "number": 2, "source_id": first["source_id"] + 1, "source_title": "Aula 2"}
    mixed = ArtifactOut.model_validate(
        {
            **artifact,
            "citations": [
                {**first, "source_url": WIKI_URL, "source_attribution": WIKI_CREDIT},
                {**plain, "source_url": None, "source_attribution": None},
            ],
        }
    )
    header, *body = _references(_csv_rows(render(mixed, "csv", "Caderno")[0]))
    assert header[3:] == ["Endereço", "Atribuição"]
    assert body[0][3:] == [WIKI_URL, WIKI_CREDIT]
    # Never a blank cell (D-24).
    assert body[1][3:] == [NO_URL_LABEL, NO_ATTRIBUTION_LABEL]
    text = _docx_text(render(mixed, "docx", "Caderno")[0])
    assert f"[2] Aula 2\n{first['excerpt']}" in text


def test_a_citation_from_before_phase_3_renders_as_it_always_did(client):
    notebook_id = _notebook(client)
    for tool in ("report", "flashcards", "mindmap"):
        artifact = _generate(client, notebook_id, tool=tool)
        # A citation stored without the two keys at all, as phase 2 stored it.
        stored = {
            **artifact,
            "citations": [
                {k: v for k, v in c.items() if k not in ("source_url", "source_attribution")}
                for c in artifact["citations"]
            ],
        }
        old = ArtifactOut.model_validate(stored)
        assert old.citations
        assert all(c.source_url is None and c.source_attribution is None for c in old.citations)
        for fmt in CATALOG[tool].exports:
            data = render(old, fmt, "Caderno")[0]
            if fmt == "csv":
                header, *body = _references(_csv_rows(data))
                assert header == ["Nº", "Fonte", "Trecho citado"]
                assert body and all(len(row) == 3 for row in body)
            elif fmt == "xlsx":
                sheet = load_workbook(io.BytesIO(data))["Referências"]
                assert sheet.max_column == 3
            elif fmt == "docx":
                text = _docx_text(data)
                references = text[text.index("\nReferências") :]
                expected = "\n".join(
                    f"[{c.number}] {studio_exporter._reference(c)}\n{c.excerpt}"
                    for c in old.citations
                )
                assert references == f"\nReferências\n{expected}"
            else:
                below = data.decode("utf-8").split("</g>")[-1]
                layout = old.layout
                chars = int((max(layout.width, 640.0) - 32) // 6.4)
                expected = [
                    line
                    for c in old.citations
                    for line in mindmap.wrap(
                        f"[{c.number}] {studio_exporter._reference(c)}", chars, 2
                    )
                ]
                assert below.count("<text") == len(expected)


def test_a_hostile_credit_is_inert_in_the_csv_and_the_xlsx(client):
    notebook_id = _notebook(client)
    artifact = _with_credit(
        _generate(client, notebook_id, tool="flashcards"),
        url="=cmd|'/c calc'!A1",
        attribution=HOSTILE,
    )
    rows = _csv_rows(render(artifact, "csv", "Caderno")[0])
    assert [LIMITATION_NOTICE] in rows and [AI_NOTICE] in rows
    body = _references(rows)[1:]
    assert body
    for row in body:
        # One apostrophe, visible: the writer's, never a second one from here.
        assert row[3] == "'=cmd|'/c calc'!A1"
        assert row[4].startswith("'=HYPERLINK(") and not row[4].startswith("''")
        assert "\x0b" not in row[4]
    for prefix in ("+", "-", "@"):
        rows = _csv_rows(
            render(_with_credit(artifact.model_dump(), attribution=f"{prefix}1+1"), "csv", "C")[0]
        )
        assert all(row[-1] == f"'{prefix}1+1" for row in _references(rows)[1:])

    workbook = load_workbook(io.BytesIO(render(artifact, "xlsx", "Caderno")[0]))
    cover = [row[0] for row in workbook["Aviso"].iter_rows(values_only=True)]
    assert LIMITATION_NOTICE in cover and AI_NOTICE in cover
    _, *body = list(workbook["Referências"].iter_rows(values_only=True))
    for row in body:
        assert row[3] == "'=cmd|'/c calc'!A1"
        assert row[4].startswith("'=HYPERLINK(")


def test_a_hostile_credit_is_text_in_the_docx(client):
    notebook_id = _notebook(client)
    artifact = _with_credit(
        _generate(client, notebook_id, tool="report"), url=WIKI_URL, attribution=HOSTILE
    )
    data = render(artifact, "docx", "Caderno")[0]
    text = _docx_text(data)
    assert LIMITATION_NOTICE in text and AI_NOTICE in text
    # Every character but the control one, literally, as text.
    assert HOSTILE.replace("\x0b", " ") in text
    document = Document(io.BytesIO(data))
    xml = document.element.xml
    assert "<script>" not in xml and "&lt;script&gt;" in xml
    assert "&amp;amp;" in xml


def test_a_hostile_credit_is_escaped_in_the_svg_and_kept_whole(client):
    notebook_id = _notebook(client)
    generated = _generate(client, notebook_id, tool="mindmap")
    artifact = _with_credit(
        generated, url='https://x.example/"><script>alert(1)</script>', attribution=HOSTILE
    )
    svg = render(artifact, "svg", "Caderno")[0].decode("utf-8")
    assert "Esta ferramenta destina-se a apoio didático" in svg
    assert "<script>" not in svg and "&lt;script&gt;" in svg
    assert '"><' not in svg.split("</g>")[-1]
    assert "\x0b" not in svg
    # Well-formed: an XML parser reads it, and every text reads back literally.
    root = ElementTree.fromstring(svg)
    texts = "".join(t.text or "" for t in root.iter("{http://www.w3.org/2000/svg}text"))
    assert "</w:t>" in texts and "&amp;" in texts

    # A long credit is wrapped, never cut with "…"; an address is never hyphenated.
    credited = _with_credit(generated, url=WIKI_URL + "x" * 200, attribution=WIKI_CREDIT * 2)
    root = ElementTree.fromstring(render(credited, "svg", "Caderno")[0].decode("utf-8"))
    lines = [t.text or "" for t in root.iter("{http://www.w3.org/2000/svg}text")]
    joined = " ".join(lines)
    assert "obtido em 25/09/2026." in joined and not joined.endswith("…")
    assert (WIKI_URL + "x" * 200) in "".join(lines)


# --- isolation -------------------------------------------------------------------


def test_another_students_artifact_is_not_found_anywhere(client, login_as, other_user):
    with login_as(other_user):
        notebook_id = _notebook(client, "Caderno do Bruno")
        artifact = _generate(client, notebook_id, tool="report")
    base = f"/api/notebooks/{notebook_id}/studio"
    responses = [
        client.get(base),
        client.post(base, json={"tool": "report"}),
        client.get(f"{base}/{artifact['id']}"),
        client.patch(f"{base}/{artifact['id']}", json={"title": "Meu"}),
        client.post(f"{base}/{artifact['id']}/note"),
        client.get(f"{base}/{artifact['id']}/export.docx"),
        client.delete(f"{base}/{artifact['id']}"),
    ]
    assert [r.status_code for r in responses] == [404] * len(responses)


# --- the pieces, alone -----------------------------------------------------------


def test_the_tree_drops_orphans_cycles_and_what_is_too_deep():
    tree = build_tree(
        [
            {"id": 1, "parent": 0, "label": "Tema", "citations": []},
            {"id": 2, "parent": 1, "label": "A", "citations": [1]},
            {"id": 3, "parent": 2, "label": "A1", "citations": [1]},
            {"id": 4, "parent": 3, "label": "fundo demais", "citations": [1]},
            {"id": 5, "parent": 9, "label": "órfão", "citations": []},
            {"id": 6, "parent": 7, "label": "ciclo", "citations": []},
            {"id": 7, "parent": 6, "label": "ciclo", "citations": []},
            {"id": 0, "parent": 1, "label": "id zero", "citations": []},
        ],
        depth=2,
    )
    assert tree == {
        "label": "Tema",
        "citations": [],
        "children": [
            {
                "label": "A",
                "citations": [1],
                "children": [{"label": "A1", "citations": [1], "children": []}],
            }
        ],
    }


def test_a_map_without_a_root_is_nothing():
    assert build_tree([{"id": 2, "parent": 1, "label": "solto", "citations": []}], 2) is None


def test_wrap_breaks_at_words_and_marks_what_it_cuts():
    assert mindmap.wrap("um dois três", width=7) == ["um dois", "três"]
    assert mindmap.wrap("a b c d e f g h", width=3, max_lines=2) == ["a b", "c…"]
    assert mindmap.wrap("supercalifragilístico", width=8)[0] == "superca-"


def test_a_parent_is_centred_on_its_children():
    drawn = mindmap.layout(
        {
            "label": "Tema",
            "citations": [],
            "children": [
                {"label": "A", "citations": [], "children": []},
                {"label": "B", "citations": [], "children": []},
            ],
        }
    )
    root, first, second = drawn.nodes
    centre = lambda node: node.y + node.height / 2  # noqa: E731
    assert centre(root) == pytest.approx((centre(first) + centre(second)) / 2)
    assert first.y < second.y
    assert drawn.width > root.width + first.width
