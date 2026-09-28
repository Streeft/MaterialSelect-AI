"""The Studio deck as .pptx (D-98): structure, notes, notices, escape."""

from __future__ import annotations

from io import BytesIO

from pptx import Presentation

from app.exporters.pptx import SLIDE_HEIGHT, SLIDE_WIDTH
from app.exporters.report import LIMITATION_NOTICE
from app.exporters.studio_pptx import (
    DECK_SUBTITLE,
    NARRATION_LABEL,
    NO_CITATIONS,
    REFERENCES_PER_SLIDE,
    WITHHELD_TITLE,
    deck_to_pptx,
)
from app.schemas.notebook import CitationOut

AI = "Gerado por IA a partir das fontes do caderno. Confira nas fontes."
NOTICES = [LIMITATION_NOTICE, AI]


def _citation(number: int, **extra) -> CitationOut:
    fields = {
        "number": number,
        "chunk_id": number,
        "source_id": 1,
        "source_title": f"Fonte {number}",
        "excerpt": f"Trecho {number}",
    }
    fields.update(extra)
    return CitationOut(**fields)


def _deck(n: int = 2) -> dict:
    return {
        "title": "Ligas leves",
        "slides": [
            {
                "title": f"Slide {i}",
                "bullets": [f"Tópico {i}.a", f"Tópico {i}.b"],
                "notes": f"Nota do slide {i}.",
                "citations": [1, 2] if i == 1 else [2],
            }
            for i in range(1, n + 1)
        ],
    }


def _open(data: bytes):
    return list(Presentation(BytesIO(data)).slides)


def _text(slide) -> str:
    return "\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)


def _notes(slide) -> str:
    return slide.notes_slide.notes_text_frame.text


def _all_text(slides) -> str:
    return "\n".join(_text(s) + "\n" + _notes(s) for s in slides)


def test_deck_is_16_9_with_title_content_references_and_notices():
    citations = [_citation(1, page_start=3), _citation(2)]
    data = deck_to_pptx("Ligas leves", _deck(3), citations, [], notices=NOTICES)
    presentation = Presentation(BytesIO(data))
    assert presentation.slide_width == SLIDE_WIDTH
    assert presentation.slide_height == SLIDE_HEIGHT
    slides = list(presentation.slides)
    # 1 title + 3 content + 1 references + 1 notices, nothing withheld.
    assert len(slides) == 1 + 3 + 1 + 1
    assert "Ligas leves" in _text(slides[0])
    assert DECK_SUBTITLE in _text(slides[0])
    assert "Slide 1" in _text(slides[1])
    assert "Tópico 1.a" in _text(slides[1])
    assert "[1] [2]" in _text(slides[1])
    assert "Referências" in _text(slides[4])
    assert "Avisos" in _text(slides[5])
    assert "didático" in _text(slides[5])
    assert AI in _text(slides[5])


def test_speaker_notes_carry_the_notes_and_the_sources():
    citations = [
        _citation(1, page_start=3, page_end=5, source_url="https://exemplo.org/a"),
        _citation(2),
    ]
    slides = _open(deck_to_pptx("T", _deck(2), citations, [], notices=NOTICES))
    notes = _notes(slides[1])
    assert "Nota do slide 1." in notes
    assert NARRATION_LABEL not in notes
    assert "Fontes:" in notes
    assert "[1] Fonte 1, p. 3–5 — https://exemplo.org/a" in notes
    assert "[2] Fonte 2" in notes
    # Slide 2 cites only [2].
    assert "[1]" not in _notes(slides[2])


def test_video_notes_are_labelled_narration():
    slides = _open(deck_to_pptx("T", _deck(1), [_citation(1)], [], notices=NOTICES, narration=True))
    assert _notes(slides[1]).startswith(f"{NARRATION_LABEL} Nota do slide 1.")


def test_a_number_without_a_citation_is_not_listed_in_the_notes():
    deck = {
        "title": "T",
        "slides": [{"title": "S", "bullets": ["b"], "notes": "", "citations": [9]}],
    }
    slides = _open(deck_to_pptx("T", deck, [_citation(1)], [], notices=NOTICES))
    assert "Fontes:" not in _notes(slides[1])


def test_references_are_paged_and_carry_url_and_attribution():
    n = REFERENCES_PER_SLIDE + 3
    citations = [_citation(i) for i in range(1, n + 1)]
    citations[0] = _citation(
        1, source_url="https://pt.wikipedia.org/wiki/A%C3%A7o", source_attribution="CC BY-SA 4.0"
    )
    slides = _open(deck_to_pptx("T", _deck(1), citations, [], notices=NOTICES))
    # 1 title + 1 content + 2 reference pages + notices.
    assert len(slides) == 1 + 1 + 2 + 1
    first, second = _text(slides[2]), _text(slides[3])
    assert "Referências (1/2)" in first and "Referências (2/2)" in second
    assert "https://pt.wikipedia.org/wiki/A%C3%A7o" in first
    assert "CC BY-SA 4.0" in first
    assert f"[{REFERENCES_PER_SLIDE}] " in first
    assert f"[{REFERENCES_PER_SLIDE + 1}] " in second


def test_withheld_sentences_get_a_final_slide():
    slides = _open(deck_to_pptx("T", _deck(1), [_citation(1)], ["Frase omitida."], notices=NOTICES))
    assert len(slides) == 1 + 1 + 1 + 1 + 1
    assert WITHHELD_TITLE in _text(slides[-1])
    assert "Frase omitida." in _text(slides[-1])


def test_an_empty_deck_still_has_title_references_and_notices():
    for content in (None, {"title": "T", "slides": []}):
        slides = _open(deck_to_pptx("T", content, [], [], notices=NOTICES))
        assert len(slides) == 3
        assert NO_CITATIONS in _text(slides[1])
        assert "didático" in _text(slides[2])


def test_control_characters_are_stripped_and_markup_stays_literal():
    deck = {
        "title": "T",
        "slides": [
            {
                "title": "Tí\x00tulo",
                "bullets": ["<script>alert(1)</script>\x0b", "a\x1fb"],
                "notes": "no\x08ta",
                "citations": [1],
            }
        ],
    }
    citations = [
        _citation(
            1,
            source_title="Fo\x0cnte",
            source_url="https://x.org/\x00a",
            source_attribution="CC\x0e BY",
        )
    ]
    data = deck_to_pptx("Ti\x00t", deck, citations, ["om\x00itida"], notices=NOTICES)
    slides = _open(data)
    text = _all_text(slides)
    assert "_x00" not in text
    assert "Título" in text and "Titulo" not in text
    assert "<script>alert(1)</script>" in text
    assert "ab" in text and "nota" in text and "Fonte" in text
    assert "https://x.org/a" in text and "CC BY" in text and "omitida" in text


def test_a_narration_only_scene_has_no_empty_body_placeholder():
    deck = {"title": "T", "slides": [{"title": "Cena", "bullets": [], "notes": "Fala."}]}
    slides = _open(deck_to_pptx("T", deck, [], [], notices=NOTICES, narration=True))
    scene = slides[1]
    assert [shape.text_frame.text for shape in scene.shapes if shape.has_text_frame] == ["Cena"]
    assert _notes(scene) == f"{NARRATION_LABEL} Fala."
