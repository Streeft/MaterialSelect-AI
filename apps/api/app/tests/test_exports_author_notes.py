"""The author's notes in the reports and the laudo (D-122, D-161).

A note is the author's plain text and each format neutralises it its own way:
the spreadsheet with a leading apostrophe, the HTML with ``html.escape``. These
tests store hostile notes through the real API and read them back out of every
real export, because a hand-built string would not prove the wiring.
"""

from __future__ import annotations

import csv
import io

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation

from app.exporters.pptx import _chunk_rows, to_pptx
from app.exporters.report import Report, Sheet

MARKUP_NOTE = '<script>alert("x")</script> E < 70 GPa & "ok"'
FORMULA_NOTE = '=HYPERLINK("http://exemplo","clique")'
SHEET = "Notas do autor"
NO_NOTE = "Nenhuma nota foi escrita neste estudo, nem em seus estágios."
ONLY_NOTED = "Só aparecem aqui o estudo e os estágios que têm nota."


def _stage(label: str, notes: str | None) -> dict:
    return {
        "kind": "limit",
        "label": label,
        "notes": notes,
        "root_group": {
            "operator": "AND",
            "constraints": [
                {"operator": "gt", "property_slug": "modulo_young", "value": 1.0, "unit": "GPa"}
            ],
        },
    }


def _study(client: TestClient, *, notes: str | None, stage_notes: list[str | None]) -> int:
    response = client.post(
        "/api/selection/studies",
        json={
            "name": "Estudo com notas",
            "universe": "material",
            "stages": [_stage(f"Etapa {i + 1}", n) for i, n in enumerate(stage_notes)],
            "index": {
                "name": "Viga leve",
                "expression": "sqrt(modulo_young) / densidade",
                "goal": "maximize",
            },
            "criteria": [{"key": "__index__", "weight": 1.0}],
            "notes": notes,
        },
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def _docx_text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _pptx_text(content: bytes) -> str:
    parts: list[str] = []
    for slide in Presentation(io.BytesIO(content)).slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if shape.has_table:
                parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return "\n".join(parts)


class TestStudyReportNotes:
    def test_html_escapes_markup_and_keeps_text(self, client: TestClient) -> None:
        sid = _study(client, notes=MARKUP_NOTE, stage_notes=["Por que 1 GPa?"])
        text = client.get(f"/api/exports/estudos/{sid}.html").text
        assert SHEET in text
        assert "&lt;script&gt;alert(&quot;x&quot;)&lt;/script&gt; E &lt; 70 GPa &amp;" in text
        assert "<script>" not in text
        assert "Por que 1 GPa?" in text
        assert ONLY_NOTED in text

    def test_html_does_not_use_the_spreadsheet_apostrophe(self, client: TestClient) -> None:
        sid = _study(client, notes=FORMULA_NOTE, stage_notes=[None])
        text = client.get(f"/api/exports/estudos/{sid}.html").text
        assert "<td>=HYPERLINK(" in text
        assert "'=HYPERLINK(" not in text

    def test_html_keeps_the_line_breaks_of_a_note(self, client: TestClient) -> None:
        sid = _study(client, notes="linha um\nlinha dois", stage_notes=[None])
        text = client.get(f"/api/exports/estudos/{sid}.html").text
        assert "linha um\nlinha dois" in text
        assert "white-space: pre-line" in text

    def test_csv_neutralises_a_formula_note(self, client: TestClient) -> None:
        sid = _study(client, notes=FORMULA_NOTE, stage_notes=["=1+1"])
        text = client.get(f"/api/exports/estudos/{sid}.csv").text
        rows = list(csv.reader(io.StringIO(text.lstrip("﻿"))))
        assert ["Estudo", "'" + FORMULA_NOTE] in rows
        assert not any(cell.startswith("=") for row in rows for cell in row)
        stage_row = next(r for r in rows if r and r[0].startswith("Estágio 1"))
        assert stage_row[1] == "'=1+1"

    def test_csv_keeps_markup_as_literal_text(self, client: TestClient) -> None:
        sid = _study(client, notes=MARKUP_NOTE, stage_notes=[None])
        text = client.get(f"/api/exports/estudos/{sid}.csv").text
        rows = list(csv.reader(io.StringIO(text.lstrip("﻿"))))
        assert ["Estudo", MARKUP_NOTE] in rows
        assert "&lt;" not in text

    def test_xlsx_neutralises_a_formula_note(self, client: TestClient) -> None:
        sid = _study(client, notes=FORMULA_NOTE, stage_notes=["@SUM(A1)"])
        response = client.get(f"/api/exports/estudos/{sid}.xlsx")
        sheet = load_workbook(io.BytesIO(response.content))[SHEET]
        cells = [c.value for row in sheet.iter_rows() for c in row if c.value is not None]
        assert "'" + FORMULA_NOTE in cells
        assert "'@SUM(A1)" in cells
        assert not any(isinstance(v, str) and v.startswith(("=", "@")) for v in cells)

    def test_docx_and_pptx_carry_the_notes(self, client: TestClient) -> None:
        sid = _study(client, notes=MARKUP_NOTE, stage_notes=["Por que 1 GPa?"])
        docx = client.get(f"/api/exports/estudos/{sid}.docx")
        pptx = client.get(f"/api/exports/estudos/{sid}.pptx")
        for text in (_docx_text(docx.content), _pptx_text(pptx.content)):
            assert SHEET in text
            assert MARKUP_NOTE in text
            assert "Por que 1 GPa?" in text
            assert "Estágio 1 — Limites — Etapa 1" in text

    def test_only_the_stage_with_a_note_gets_a_row(self, client: TestClient) -> None:
        sid = _study(client, notes=None, stage_notes=[None, "só a segunda"])
        text = client.get(f"/api/exports/estudos/{sid}.csv").text
        assert "Estágio 2 — Limites — Etapa 2,só a segunda" in text
        assert "Estágio 1 — Limites — Etapa 1" not in text

    @pytest.mark.parametrize("fmt", ["csv", "html", "docx", "pptx"])
    def test_a_study_without_notes_says_so_in_words(self, client: TestClient, fmt: str) -> None:
        sid = _study(client, notes=None, stage_notes=[None])
        response = client.get(f"/api/exports/estudos/{sid}.{fmt}")
        assert response.status_code == 200, response.text
        if fmt == "docx":
            text = _docx_text(response.content)
        elif fmt == "pptx":
            text = _pptx_text(response.content)
        else:
            text = response.text
        assert SHEET in text
        assert NO_NOTE in text
        assert ONLY_NOTED not in text

    def test_a_study_without_notes_has_no_empty_row(self, client: TestClient) -> None:
        sid = _study(client, notes=None, stage_notes=[None])
        sheet = load_workbook(io.BytesIO(client.get(f"/api/exports/estudos/{sid}.xlsx").content))[
            SHEET
        ]
        values = [[c.value for c in row if c.value is not None] for row in sheet.iter_rows()]
        assert values == [[NO_NOTE], ["Onde", "Nota"]]


class TestLaudoNotes:
    def test_laudo_html_escapes_the_notes(self, client: TestClient) -> None:
        sid = _study(client, notes=MARKUP_NOTE, stage_notes=["=1+1"])
        text = client.get(f"/api/exports/estudos/{sid}/laudo.html").text
        assert SHEET in text
        assert "&lt;script&gt;" in text
        assert "<script>" not in text
        assert "<td>=1+1</td>" in text

    def test_laudo_docx_and_pptx_carry_the_notes(self, client: TestClient) -> None:
        sid = _study(client, notes="Nota do laudo.", stage_notes=["Nota do estágio."])
        docx = client.get(f"/api/exports/estudos/{sid}/laudo.docx")
        pptx = client.get(f"/api/exports/estudos/{sid}/laudo.pptx")
        for text in (_docx_text(docx.content), _pptx_text(pptx.content)):
            assert "Nota do laudo." in text
            assert "Nota do estágio." in text

    def test_laudo_without_notes_says_so(self, client: TestClient) -> None:
        sid = _study(client, notes=None, stage_notes=[None])
        text = client.get(f"/api/exports/estudos/{sid}/laudo.html").text
        assert NO_NOTE in text


class TestLongNotesInPptx:
    def _rows(self, n: int, size: int) -> list[list[object]]:
        return [[f"Estágio {i}", "x" * size] for i in range(n)]

    def test_short_rows_keep_ten_per_slide(self) -> None:
        assert [len(c) for c in _chunk_rows(self._rows(25, 10))] == [10, 10, 5]

    def test_a_two_thousand_character_note_takes_its_own_slide(self) -> None:
        rows = self._rows(1, 5) + self._rows(1, 2000) + self._rows(1, 5)
        assert [len(c) for c in _chunk_rows(rows)] == [1, 1, 1]

    def test_order_is_kept(self) -> None:
        rows = self._rows(3, 150) + self._rows(3, 5)
        assert [r for chunk in _chunk_rows(rows) for r in chunk] == rows

    def test_rendered_deck_has_one_slide_per_long_note(self) -> None:
        sheet = Sheet(name=SHEET, header=["Onde", "Nota"], rows=self._rows(3, 2000))
        report = Report(title="T", subtitle="", notices=[], sheets=[sheet])
        assert len(Presentation(io.BytesIO(to_pptx(report))).slides) == 1 + 3
