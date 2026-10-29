"""Import wizard hardening (D-176): single commit, bounded reading, curator gate."""

from __future__ import annotations

import csv
import io
from datetime import UTC, datetime

import pytest
from openpyxl import Workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from sqlalchemy import func, select

from app.config import settings
from app.domain.errors import ConflictError, ValidationError
from app.importers import readers
from app.importers.service import ImportService
from app.models.import_job import ImportJob
from app.models.material import Material
from app.models.subscription import Subscription

CSV_CONTENT = (
    b"nome;classe;densidade [g/cm3]\nLiga Corrida A;Metais;2,70\nLiga Corrida B;Metais;7,80\n"
)

MAPPING = {
    "name_column": "nome",
    "class_column": "classe",
    "source_label": "Planilha de teste",
    "source_license_label": "Uso interno autorizado",
    "columns": [
        {"column": "densidade [g/cm3]", "property_slug": "densidade", "unit": "g/cm**3"},
    ],
}


def _validated_job(client) -> int:
    upload = client.post(
        "/api/imports/upload", files={"file": ("corrida.csv", CSV_CONTENT, "text/csv")}
    )
    assert upload.status_code == 200, upload.text
    job_id = upload.json()["job_id"]
    validated = client.post(f"/api/imports/{job_id}/validate", json={"mapping": MAPPING})
    assert validated.status_code == 200, validated.text
    return job_id


def _materials_of(db_session, job_id: int) -> int:
    stmt = select(func.count()).select_from(Material).where(Material.import_job_id == job_id)
    return db_session.execute(stmt).scalar_one()


# --- A1: one job, one commit -------------------------------------------------


def test_a_commit_working_from_a_stale_status_cannot_import_twice(client, db_session, test_user):
    job_id = _validated_job(client)
    # A second request read the job while it was still VALIDADO (double click,
    # client retry, another worker) and keeps that stale copy.
    stale = db_session.get(ImportJob, job_id)
    assert stale is not None and stale.status.value == "VALIDADO"

    first = client.post(f"/api/imports/{job_id}/commit")
    assert first.status_code == 200, first.text
    assert first.json()["imported_count"] == 2

    with pytest.raises(ConflictError, match="já foi efetivada"):
        ImportService(db_session, test_user).commit(job_id)
    assert _materials_of(db_session, job_id) == 2


def test_a_second_commit_over_http_is_a_409(client, db_session):
    job_id = _validated_job(client)
    assert client.post(f"/api/imports/{job_id}/commit").status_code == 200
    again = client.post(f"/api/imports/{job_id}/commit")
    assert again.status_code == 409
    assert _materials_of(db_session, job_id) == 2


# --- B2: limits enforced while reading ---------------------------------------


def test_csv_over_the_row_limit_is_refused_without_reading_it_all(monkeypatch):
    real_reader = csv.reader
    consumed = {"rows": 0}

    def counting_reader(*args, **kwargs):
        for row in real_reader(*args, **kwargs):
            consumed["rows"] += 1
            yield row

    monkeypatch.setattr(readers.csv, "reader", counting_reader)
    data = b"nome\n" + b"a\n" * 200_000
    with pytest.raises(ValidationError, match="mais de 50 linhas"):
        readers.read_csv(data, max_rows=50)
    assert consumed["rows"] <= 52


def _xlsx(rows: int, cols: int = 3) -> bytes:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Dados")
    ws.append([f"c{i}" for i in range(cols)])
    for r in range(rows):
        ws.append([f"m{r}"] + [r] * (cols - 1))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_xlsx_over_the_row_limit_is_refused_without_reading_it_all(monkeypatch):
    real_iter_rows = ReadOnlyWorksheet.iter_rows
    consumed = {"rows": 0}

    def counting_iter_rows(self, *args, **kwargs):
        for row in real_iter_rows(self, *args, **kwargs):
            consumed["rows"] += 1
            yield row

    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", counting_iter_rows)
    with pytest.raises(ValidationError, match="A aba tem mais de 50 linhas"):
        readers.read_xlsx(_xlsx(5_000), max_rows=50)
    assert consumed["rows"] <= 52


def test_xlsx_that_inflates_past_the_cap_is_refused_before_opening(monkeypatch):
    monkeypatch.setattr(readers, "MAX_XLSX_UNCOMPRESSED_BYTES", 1024)
    with pytest.raises(ValidationError, match="descompactado passa de"):
        readers.read_xlsx(_xlsx(200), max_rows=5000)


def test_too_many_columns_is_refused():
    header = ";".join(f"c{i}" for i in range(readers.MAX_COLUMNS + 1))
    with pytest.raises(ValidationError, match="colunas; o limite é"):
        readers.read_csv(f"{header}\n".encode() + b"x\n", max_rows=10)


def test_files_within_the_limits_read_exactly_as_before():
    table = readers.read_csv(b"nome;densidade\n\nA;1,0\nB;2,0\n", max_rows=2)
    assert table.headers == ["nome", "densidade"]
    assert table.rows == [["A", "1,0"], ["B", "2,0"]]

    sheet = readers.read_xlsx(_xlsx(3), max_rows=3)
    assert sheet.headers == ["c0", "c1", "c2"]
    assert len(sheet.rows) == 3 and sheet.sheet_name == "Dados"


def test_sqlite_table_over_the_limit_is_refused(tmp_path):
    import sqlite3

    path = tmp_path / "t.sqlite"
    conn = sqlite3.connect(path)
    conn.execute("CREATE TABLE m (nome TEXT)")
    conn.executemany("INSERT INTO m VALUES (?)", [(f"m{i}",) for i in range(100)])
    conn.commit()
    conn.close()
    data = path.read_bytes()
    with pytest.raises(ValidationError, match="A tabela tem mais de 10 linhas"):
        readers.read_sqlite(data, max_rows=10)
    assert len(readers.read_sqlite(data, max_rows=100).rows) == 100


# --- B5: the curator gate covers templates, history and reports ---------------


def _subscribe(db_session, user) -> None:
    db_session.add(
        Subscription(
            user_id=user.id,
            stripe_customer_id="cus_import",
            stripe_subscription_id="sub_import",
            status="active",
            created_at=datetime.now(UTC),
            updated_at=datetime.now(UTC),
        )
    )
    db_session.commit()


TEMPLATE = {"name": "Padrão da turma", "mapping": MAPPING}


def test_open_mode_student_cannot_touch_templates_history_or_reports(client, monkeypatch):
    job_id = _validated_job(client)  # created by a curator (closed mode)
    monkeypatch.setattr(settings, "access_mode", "open")
    assert client.post("/api/import-templates", json=TEMPLATE).status_code == 403
    assert client.get("/api/imports").status_code == 403
    assert client.get(f"/api/imports/{job_id}/report").status_code == 403


def test_open_mode_subscriber_creates_templates(client, db_session, test_user, monkeypatch):
    monkeypatch.setattr(settings, "access_mode", "open")
    _subscribe(db_session, test_user)
    assert client.post("/api/import-templates", json=TEMPLATE).status_code == 201
    assert client.get("/api/imports").status_code == 200


def test_closed_mode_without_subscription_is_refused(client_without_subscription):
    assert (
        client_without_subscription.post("/api/import-templates", json=TEMPLATE).status_code == 403
    )
    assert client_without_subscription.get("/api/imports").status_code == 403
    assert client_without_subscription.get("/api/imports/1/report").status_code == 403
