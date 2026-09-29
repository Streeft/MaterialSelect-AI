"""Exporting the D-98 Studio tools: the audio script (DOCX, TXT), slides and
video (PPTX) and the infographic (SVG) — each file with its notices, its
references and what the grounding check withheld, escaped by format."""

from __future__ import annotations

import io
from datetime import UTC, datetime
from xml.etree import ElementTree

import pytest
from docx import Document
from pptx import Presentation

from app.exporters.report import LIMITATION_NOTICE
from app.exporters.studio import (
    AI_NOTICE,
    NO_CITATIONS,
    SPEAKER_LABEL,
    WITHHELD_TITLE,
    infographic_svg,
    render,
)
from app.exporters.studio_pptx import NARRATION_LABEL
from app.notebooks import infographic
from app.schemas.notebook import ArtifactOut, CitationOut

SVG = "{http://www.w3.org/2000/svg}"
URL = "https://pt.wikipedia.org/wiki/A%C3%A7o"
CREDIT = "Texto de “Aço”, da Wikipédia, sob a licença CC BY-SA 4.0; revisão 123."
HOSTILE = "<script>alert(1)</script> & aço"
WITHHELD = (
    "Um dado em destaque foi omitido porque sua unidade não aparece no trecho citado: 210 MPa"
)
NOTEBOOK = "Materiais estruturais"


def _citation(number: int, **extra) -> CitationOut:
    fields = {
        "number": number,
        "chunk_id": number,
        "source_id": 1,
        "source_title": f"Fonte {number}",
        "excerpt": f"Trecho {number}",
    }
    fields.update(extra)
    return CitationOut(**fields)


def _artifact(tool: str, content: dict, **extra) -> ArtifactOut:
    now = datetime.now(UTC)
    fields = {
        "id": 1,
        "tool": tool,
        "title": "Ligas leves",
        "status": "pronto",
        "source_count": 1,
        "created_at": now,
        "updated_at": now,
        "content": content,
        "citations": [
            _citation(1, page_start=3),
            _citation(2, source_url=URL, source_attribution=CREDIT),
        ],
        "withheld": [WITHHELD],
    }
    fields.update(extra)
    return ArtifactOut(**fields)


AUDIO = {
    "title": "Ligas leves",
    "lines": [
        {"speaker": 1, "text": "Vamos ver o que a fonte diz.", "citations": [1]},
        {"speaker": 2, "text": "O alumínio tem 2700 kg/m³.", "citations": [2]},
    ],
}

DECK = {
    "title": "Ligas leves",
    "slides": [
        {
            "title": "Densidade",
            "bullets": ["Alumínio: 2700 kg/m³"],
            "notes": "Narra.",
            "citations": [2],
        },
    ],
}

INFOGRAPHIC = {
    "title": f"Aço {HOSTILE}",
    "subtitle": "O que as fontes dizem",
    "stats": [{"value": "210 GPa", "label": "módulo do aço", "citations": [1]}],
    "points": [{"heading": "Leveza", "text": "O alumínio é leve.", "citations": [2]}],
    "steps": [
        {"text": "Definir a função", "citations": [1]},
        {"text": "Comparar materiais", "citations": [2]},
    ],
}


# --- audio -----------------------------------------------------------------------


def _docx_text(data: bytes) -> str:
    return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)


def test_audio_docx_is_the_script_with_notices_references_and_withheld():
    data, media = render(_artifact("audio", AUDIO), "docx", NOTEBOOK)
    assert media.endswith("wordprocessingml.document")
    text = _docx_text(data)
    assert "Roteiro" in text
    assert f"{SPEAKER_LABEL} 1: Vamos ver o que a fonte diz. [1]" in text
    assert f"{SPEAKER_LABEL} 2: O alumínio tem 2700 kg/m³. [2]" in text
    for expected in (LIMITATION_NOTICE, AI_NOTICE, WITHHELD, "Referências", "[1] Fonte 1, p. 3"):
        assert expected in text
    assert URL in text and CREDIT in text
    # One paragraph per spoken line.
    lines = [p.text for p in Document(io.BytesIO(data)).paragraphs]
    assert sum(line.startswith(SPEAKER_LABEL) for line in lines) == 2


def test_audio_txt_is_notices_then_the_script_then_the_references():
    data, media = render(_artifact("audio", AUDIO), "txt", NOTEBOOK)
    assert media == "text/plain; charset=utf-8"
    text = data.decode("utf-8")
    notice = text.index(LIMITATION_NOTICE)
    ai = text.index(AI_NOTICE)
    script = text.index(f"{SPEAKER_LABEL} 1: Vamos ver o que a fonte diz. [1]")
    references = text.index("Referências")
    assert notice < ai < script < references
    assert f"\n\n{SPEAKER_LABEL} 1:" in text  # a blank line before the script
    assert WITHHELD in text and WITHHELD_TITLE in text
    assert "[2] Fonte 2" in text and URL in text and CREDIT in text
    assert "Trecho 1" in text


def test_audio_without_citations_says_so():
    artifact = _artifact("audio", AUDIO, citations=[], withheld=[])
    text = render(artifact, "txt", NOTEBOOK)[0].decode("utf-8")
    assert NO_CITATIONS in text
    assert WITHHELD_TITLE not in text
    assert NO_CITATIONS in _docx_text(render(artifact, "docx", NOTEBOOK)[0])


def test_a_control_character_in_a_line_does_not_break_the_docx():
    content = {"title": "x", "lines": [{"speaker": 1, "text": "a\x0bb", "citations": []}]}
    data, _ = render(_artifact("audio", content), "docx", NOTEBOOK)
    assert f"{SPEAKER_LABEL} 1: a b" in _docx_text(data)


# --- slides and video -------------------------------------------------------------


@pytest.mark.parametrize("tool", ["slides", "video"])
def test_slides_and_video_export_a_deck_that_reopens(tool):
    data, media = render(_artifact(tool, DECK), "pptx", NOTEBOOK)
    assert media == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    slides = list(Presentation(io.BytesIO(data)).slides)
    text = "\n".join(
        shape.text_frame.text for s in slides for shape in s.shapes if shape.has_text_frame
    )
    notes = "\n".join(s.notes_slide.notes_text_frame.text for s in slides if s.has_notes_slide)
    assert "Densidade" in text
    assert NOTEBOOK in text  # the subtitle line reaches the title slide
    assert LIMITATION_NOTICE in text and AI_NOTICE in text
    assert WITHHELD in text
    assert URL in text and CREDIT in text
    assert (NARRATION_LABEL in notes) is (tool == "video")


# --- infographic -------------------------------------------------------------------


def _svg(artifact: ArtifactOut) -> tuple[str, ElementTree.Element]:
    data, media = render(artifact, "svg", NOTEBOOK)
    assert media == "image/svg+xml"
    text = data.decode("utf-8")
    return text, ElementTree.fromstring(text)


def test_infographic_svg_parses_escapes_and_has_no_foreign_object():
    artifact = _artifact("infographic", INFOGRAPHIC, format="paisagem", title=INFOGRAPHIC["title"])
    text, root = _svg(artifact)
    assert "<script>" not in text
    assert "&lt;script&gt;" in text and "&amp; aço" in text
    assert "foreignObject" not in text
    assert "href" not in text and "url(" not in text and "@import" not in text
    # The escaped text reads back as the original.
    strings = [t.text or "" for t in root.iter(f"{SVG}text")]
    assert any(HOSTILE in s for s in strings)


def test_infographic_svg_has_size_and_a_background_covering_the_canvas():
    _, root = _svg(_artifact("infographic", INFOGRAPHIC, format="retrato"))
    width, height = int(root.get("width")), int(root.get("height"))
    assert root.get("viewBox") == f"0 0 {width} {height}"
    drawn = infographic.layout(INFOGRAPHIC, "retrato")
    assert width == drawn.width == 800
    assert height > drawn.height  # the footer sits below the poster
    background = root.find(f"{SVG}rect")
    assert background is not None
    assert (background.get("x"), background.get("y")) == ("0", "0")
    assert (int(background.get("width")), int(background.get("height"))) == (width, height)
    assert background.get("fill") not in (None, "none", "transparent")


def test_infographic_svg_draws_the_same_layout_the_screen_gets():
    artifact = _artifact("infographic", INFOGRAPHIC, format="quadrado")
    _, root = _svg(artifact)
    drawn = infographic.layout(INFOGRAPHIC, "quadrado")
    cards = [b for b in drawn.blocks if b.kind in ("stat", "point", "step")]
    rects = [
        (int(r.get("x")), int(r.get("y")), int(r.get("width")), int(r.get("height")))
        for r in root.iter(f"{SVG}rect")
    ][1:]
    assert rects == [(b.x, b.y, b.width, b.height) for b in cards]
    assert len(list(root.iter(f"{SVG}line"))) == len(drawn.connectors) + 1  # + footer rule
    assert len(list(root.iter(f"{SVG}polygon"))) == len(drawn.connectors)
    strings = [t.text or "" for t in root.iter(f"{SVG}text")]
    assert "210 GPa" in strings and "[1]" in strings


def test_infographic_footer_has_notices_references_and_withheld():
    _, root = _svg(_artifact("infographic", INFOGRAPHIC))
    joined = " ".join(t.text or "" for t in root.iter(f"{SVG}text"))
    for expected in ("Referências", "[1] Fonte 1, p. 3", "[2] Fonte 2", URL, WITHHELD_TITLE):
        assert expected in joined
    # Notices, credit and withheld sentences are wrapped: compare word by word.
    for sentence in (LIMITATION_NOTICE, AI_NOTICE, CREDIT, WITHHELD):
        assert " ".join(sentence.split()) in " ".join(joined.split())
    assert NOTEBOOK in joined
    footer_ys = [
        int(t.get("y")) for t in root.iter(f"{SVG}text") if (t.text or "") == "Referências"
    ]
    assert footer_ys[0] > infographic.layout(INFOGRAPHIC, None).height


def test_infographic_without_citations_says_so_in_the_footer():
    svg = infographic_svg(_artifact("infographic", INFOGRAPHIC, citations=[], withheld=[]), "sub")
    assert NO_CITATIONS in svg
    assert WITHHELD_TITLE not in svg


def test_the_mind_map_still_goes_to_its_own_svg():
    # "svg" means the mind map for every tool but the infographic.
    with pytest.raises(ValueError, match="Mapa mental sem layout"):
        render(_artifact("mindmap", {"root": None}), "svg", NOTEBOOK)


# --- end to end: every new tool exports its formats, and only those ----------------

TEXT = (
    "O aço carbono tem densidade de 7850 kg/m³ e módulo de elasticidade de 210 GPa. "
    "É o material estrutural mais usado na construção civil.\n\n"
    "O alumínio 6061 tem densidade de 2700 kg/m³, cerca de um terço da do aço. "
    "Por isso aparece em estruturas leves, como quadros de bicicleta."
)


def _generate(client, tool: str, **payload) -> tuple[int, dict]:
    notebook = client.post("/api/notebooks", json={"title": NOTEBOOK}).json()
    response = client.post(
        f"/api/notebooks/{notebook['id']}/sources/text", json={"title": "Aula 1", "text": TEXT}
    )
    assert response.status_code == 201, response.text
    response = client.post(
        f"/api/notebooks/{notebook['id']}/studio", json={"tool": tool, **payload}
    )
    assert response.status_code == 202, response.text
    artifact = client.get(f"/api/notebooks/{notebook['id']}/studio/{response.json()['id']}")
    assert artifact.json()["status"] == "pronto", artifact.json()
    return notebook["id"], artifact.json()


@pytest.mark.parametrize(
    ("tool", "fmt", "media", "refused"),
    [
        ("audio", "docx", "application/vnd.openxmlformats", "pptx"),
        ("audio", "txt", "text/plain; charset=utf-8", "svg"),
        ("slides", "pptx", "application/vnd.openxmlformats", "docx"),
        ("video", "pptx", "application/vnd.openxmlformats", "txt"),
        ("infographic", "svg", "image/svg+xml", "pptx"),
    ],
)
def test_each_new_tool_exports_its_formats_and_refuses_the_others(
    client, tool, fmt, media, refused
):
    notebook_id, artifact = _generate(client, tool)
    base = f"/api/notebooks/{notebook_id}/studio/{artifact['id']}"
    response = client.get(f"{base}/export.{fmt}")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(media)
    assert response.headers["content-disposition"].endswith(f'.{fmt}"')
    if fmt == "svg":
        assert "default-src 'none'" in response.headers["content-security-policy"]
        ElementTree.fromstring(response.content)
    if fmt == "txt":
        assert LIMITATION_NOTICE in response.content.decode("utf-8")

    other = client.get(f"{base}/export.{refused}")
    assert other.status_code == 400
    assert other.json()["detail"].startswith(f"{artifact_label(tool)} não se exporta em")


def artifact_label(tool: str) -> str:
    from app.ai.studio import CATALOG

    return CATALOG[tool].label
