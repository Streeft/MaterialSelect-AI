"""Hardening of the D-98 Studio after the final review of PR #83.

Three holes, each reproduced here as the review found it:

* **A headline is a finding.** The infographic's title and subtitle — and the
  cover of a deck or a video — are printed big and cite nothing; "45% das
  falhas são por fadiga" there is the invented headline the strict rule of a
  statistic exists to stop.
* **A line never breaks inside a figure.** The layout of the infographic and of
  the mind map wraps text by character count; cutting or hyphenating "1 200 MPa"
  or "12.345.678.901 kWh" prints numbers no source states.
* **A control character never reaches a file.** Text extracted from a PDF
  carries some (``\\x02`` from a ligature), the model copies them, and python-docx
  refuses the string while an SVG holding one does not parse.
"""

from __future__ import annotations

import io
import itertools
from datetime import UTC, datetime
from xml.etree import ElementTree

import pytest
from docx import Document
from openpyxl import load_workbook

from app.ai.guardrails import _NUMBER_TOKEN, numeric_tokens
from app.ai.notebook import Passage
from app.ai.studio import (
    MAX_LINE,
    MAX_SHORT,
    MAX_TITLE,
    StudioRequest,
    _cap,
    read_audio,
    read_deck,
    read_infographic,
    read_studio,
    strip_markup,
)
from app.exporters.studio import infographic_svg, mindmap_svg, render
from app.notebooks import infographic, mindmap
from app.notebooks.infographic import STYLES, layout
from app.notebooks.studio_content import check, retry_note, withheld_sentences
from app.schemas.notebook import ArtifactOut, CitationOut, MindMapLayoutOut

STEEL = Passage(
    number=1,
    source_title="Apostila de materiais",
    heading="Tabela 45 — Aços",
    page_start=3,
    page_end=3,
    text="O aço tem módulo de 210 GPa.",
)
POLYMER = Passage(
    number=2,
    source_title="Notas de polímeros",
    heading=None,
    page_start=None,
    page_end=None,
    text="Em 80% dos casos o polietileno amolece perto de 120 °C.",
)
PASSAGES = (STEEL, POLYMER)

TITLE = "45% das falhas são por fadiga"
SUBTITLE = "Em 12 anos, 80% dos casos"


def _request(tool: str) -> StudioRequest:
    return StudioRequest(
        tool=tool,
        notebook_title="Materiais",
        source_titles=tuple(p.source_title for p in PASSAGES),
        passages=PASSAGES,
    )


def _board(title: str = TITLE, subtitle: str = SUBTITLE, cites=(1,)) -> dict:
    return {
        "title": title,
        "subtitle": subtitle,
        "stats": [{"value": "210 GPa", "label": "módulo do aço", "citations": list(cites)}],
        "points": [],
        "steps": [],
    }


# --- I-1: a headline is held to the strict rule --------------------------------------


def test_the_review_s_headline_is_replaced_and_said_so():
    checked = check(_request("infographic"), _board(), set(), "Infográfico")
    assert checked.title == "Infográfico"
    assert checked.body["title"] == "Infográfico"
    assert checked.body["subtitle"] == ""
    assert checked.items == 1  # the statistic stays: only the headline failed
    assert withheld_sentences(checked) == [
        "O título gerado foi trocado por um título neutro porque trazia números que não "
        "aparecem nos trechos citados: 45.",
        "O subtítulo gerado foi omitido porque trazia números que não aparecem nos "
        "trechos citados: 12, 80.",
    ]


def test_the_student_s_words_do_not_ground_a_headline():
    # "45" in the topic grounds a point; never a headline.
    checked = check(_request("infographic"), _board(), {45.0, 12.0, 80.0}, "Infográfico")
    assert checked.title == "Infográfico"
    assert checked.body["subtitle"] == ""


def test_a_headline_is_checked_against_the_passages_the_items_cite():
    # 80 is in passage 2, which no kept item cites: the reader could not find it.
    subtitle = "Em 80% dos casos"
    only_steel = check(_request("infographic"), _board(subtitle=subtitle), set(), "R")
    assert only_steel.body["subtitle"] == ""
    both = _board(subtitle=subtitle)
    both["points"] = [{"heading": "PE", "text": "Amolece.", "citations": [2]}]
    cited = check(_request("infographic"), both, set(), "R")
    assert cited.body["subtitle"] == subtitle
    assert "subtitle" not in cited.withheld


def test_a_passage_s_label_does_not_ground_a_headline():
    # "Tabela 45" is where a figure is, never a figure.
    checked = check(_request("infographic"), _board(title="As 45 ligas", subtitle=""), set(), "R")
    assert checked.title == "R"


def test_a_grounded_headline_stays_as_written():
    read = _board(title="Aço de 210 GPa", subtitle="Módulo do aço")
    checked = check(_request("infographic"), read, set(), "R")
    assert checked.title == "Aço de 210 GPa"
    assert checked.body["title"] == "Aço de 210 GPa"
    assert checked.body["subtitle"] == "Módulo do aço"
    assert checked.withheld == {}


@pytest.mark.parametrize("tool", ["slides", "video"])
def test_a_deck_s_cover_title_is_a_headline_too(tool):
    slide = {"title": "Aço", "bullets": ["210 GPa"], "notes": "Rígido.", "citations": [1]}
    checked = check(_request(tool), {"title": TITLE, "slides": [slide]}, set(), "Apresentação")
    assert checked.title == checked.body["title"] == "Apresentação"
    assert checked.items == 1
    assert withheld_sentences(checked)[0].startswith("O título gerado foi trocado")


def test_an_audio_title_keeps_the_d94_structural_rule():
    # An audio has no cover: its title is structural, as a report's.
    line = {"speaker": 1, "text": "Oi.", "citations": [1]}
    checked = check(_request("audio"), {"title": TITLE, "lines": [line]}, set(), "R")
    assert checked.title == TITLE
    assert checked.withheld == {}


def test_the_retry_note_says_what_went_wrong_by_kind():
    read = _board()
    read["stats"].append({"value": "210 MPa", "label": "módulo", "citations": [1]})
    read["points"] = [{"heading": "X", "text": "Custa 999 por kg.", "citations": [1]}]
    note = retry_note(check(_request("infographic"), read, set(), "R"))
    assert "números que não aparecem nos trechos citados: 999" in note
    assert "cada item cita" in note and "parágrafo" not in note
    assert "unidade que o trecho não escreve" in note and "210 MPa" in note
    assert "O título e o subtítulo não citam trechos" in note
    assert "45, 12, 80" in note


# --- I-2: a figure is never split -------------------------------------------------------


def _figures(text: str) -> list[set[float]]:
    return numeric_tokens(text)


def _only_stated(lines: list[str], original: str) -> bool:
    """Every figure the lines print is one the original text writes."""
    stated = {frozenset(readings) for readings in _figures(original)}
    return all(frozenset(readings) in stated for line in lines for readings in _figures(line))


STEP_TEXT = (
    "Separar os corpos de prova conforme a norma, medir a seção de cada um com o "
    "paquímetro, montar o extensômetro na garra da máquina e aplicar a carga aos poucos, "
    "anotando cada leitura; repetir para três amostras de cada lote antes de comparar os "
    "resultados entre si e com a especificação do fornecedor para estruturas; e o valor "
    "medido foi 1 200 MPa no ensaio"
)


def test_a_long_step_ending_in_a_figure_is_printed_whole():
    steps = [{"text": STEP_TEXT, "citations": [1]}] + [
        {"text": "Comparar.", "citations": [1]} for _ in range(2)
    ]
    drawn = layout({"title": "T", "steps": steps}, "paisagem")
    first = next(b for b in drawn.blocks if b.kind == "step")
    assert " ".join(first.body_lines) == STEP_TEXT  # nothing cut, nothing hyphenated
    assert any("1 200 MPa" in line for line in first.body_lines)
    assert not any(line.endswith("…") for line in first.body_lines)
    assert _only_stated(first.body_lines, STEP_TEXT)
    # Past its budget side by side, the flow is stacked, where a line is wider.
    assert len(first.body_lines) <= STYLES["step"].body_max_lines
    assert len({b.x for b in drawn.blocks if b.kind == "step"}) == 1


def test_a_long_statistic_widens_its_band_instead_of_being_hyphenated():
    stats = [
        {"value": "12.345.678.901 kWh", "label": "energia", "citations": [1]},
        *({"value": f"{n} GPa", "label": "módulo", "citations": [1]} for n in (210, 70, 45)),
    ]
    drawn = layout({"title": "T", "stats": stats}, "paisagem")
    blocks = [b for b in drawn.blocks if b.kind == "stat"]
    big = blocks[0]
    assert "12.345.678.901" in " ".join(big.heading_lines)
    assert not any(line.endswith("-") for line in big.heading_lines)
    style = STYLES["stat"]
    chars = int((big.width - 2 * style.pad_x) // style.heading_char)
    assert all(len(line) <= chars for line in big.heading_lines)
    assert len([b for b in blocks if b.y == big.y]) < infographic.STAT_COLUMNS["paisagem"]
    assert all(
        _only_stated(b.heading_lines, s["value"]) for b, s in zip(blocks, stats, strict=True)
    )


def test_a_band_that_fits_keeps_its_columns():
    stats = [{"value": f"{n} GPa", "label": "módulo", "citations": [1]} for n in (1, 2, 3, 4)]
    drawn = layout({"title": "T", "stats": stats}, "paisagem")
    assert len({b.y for b in drawn.blocks if b.kind == "stat"}) == 1


def test_a_long_point_is_never_cut():
    text = " ".join(["propriedade"] * 60) + " de 1.200 kg/m³"
    drawn = layout({"points": [{"heading": "H", "text": text, "citations": [1]}]}, "quadrado")
    point = drawn.blocks[0]
    assert " ".join(point.body_lines) == text


@pytest.mark.parametrize(
    "text",
    [
        "o valor medido foi 1 200 MPa no ensaio",
        "12.345.678.901 kWh",
        "a densidade de 7 850 kg/m³ e 3,5 % de carbono, 12 345 678 ciclos",
        "entre 1.200 MPa e 1.500 MPa, cerca de -196 °C, e 2,70 g/cm³",
    ],
)
def test_wrap_never_prints_a_figure_the_text_does_not_write(text):
    for width, most in itertools.product(range(3, 30), (1, 2, 3, 50)):
        lines = mindmap.wrap(text, width, most)
        assert _only_stated(lines, text), (width, most, lines)
        if lines[-1].endswith("…"):
            # A cut never leaves a figure right before the "…".
            words = lines[-1][:-1].split()
            assert not words or not any(ch.isdecimal() for ch in words[-1]), lines


def test_wrap_keeps_a_figure_whole_and_hyphenates_only_words():
    assert mindmap.wrap("12.345.678.901 kWh", 13, 2) == ["12.345.678.901", "kWh"]
    assert mindmap.wrap("foi 1 200 MPa", 9, 3) == ["foi", "1 200 MPa"]
    assert mindmap.wrap("supercalifragilístico", 8)[0] == "superca-"
    # A cut drops a trailing figure: "1 200…" would read as a number that goes on.
    assert mindmap.wrap("valor 1 200 MPa e mais texto", 9, 1) == ["valor…"]


def test_a_mind_map_node_grows_to_hold_a_long_figure():
    label = "Energia de 12.345.678.901.234.567.890.123.456 kWh"
    drawn = mindmap.layout({"label": label, "citations": [], "children": []})
    node = drawn.nodes[0]
    assert any("12.345.678.901.234.567.890.123.456" in line for line in node.lines)
    longest = max(len(line) for line in node.lines)
    assert node.width >= longest * mindmap.CHAR_WIDTH + 2 * mindmap.PAD_X


# --- I-3: control characters ------------------------------------------------------------


def test_the_reader_strips_control_characters_from_every_tool():
    audio = read_audio({"title": "T\x01x", "lines": [{"speaker": 1, "text": "a\x02b"}]})
    assert audio["title"] == "T x"
    assert audio["lines"][0]["text"] == "a b"
    deck = read_deck({"title": "D\x00", "slides": [{"title": "S\x07", "bullets": ["b\x1b"]}]})
    assert deck["title"] == "D" and deck["slides"][0]["title"] == "S"
    assert deck["slides"][0]["bullets"] == ["b"]
    board = read_infographic(
        {
            "title": "I\x0b",
            "subtitle": "s\x7f",
            "stats": [{"value": "210\x03 GPa", "label": "l\x9f"}],
            "points": [{"heading": "h\x0e", "text": "a\x03b"}],
            "steps": [{"text": "p\ud800"}],
        }
    )
    assert board["title"] == "I" and board["subtitle"] == "s"
    assert board["stats"][0] == {"value": "210 GPa", "label": "l", "citations": []}
    assert board["points"][0]["heading"] == "h" and board["points"][0]["text"] == "a b"
    assert board["steps"][0]["text"] == "p"
    report = read_studio(
        StudioRequest(tool="report", notebook_title="N", source_titles=(), passages=()),
        {"title": "R\x02", "sections": [{"heading": "H\x02", "paragraphs": [{"text": "x\x02y"}]}]},
    )
    assert report["title"] == "R"
    assert report["sections"][0]["paragraphs"][0]["text"] == "x y"
    assert strip_markup("<b>a</b>\x01\tb\nc") == "a b c"


def _citation() -> CitationOut:
    return CitationOut(number=1, chunk_id=1, source_id=1, source_title="F\x02", excerpt="e\x02")


def _artifact(tool: str, content: dict, **extra) -> ArtifactOut:
    now = datetime.now(UTC)
    fields = {
        "id": 1,
        "tool": tool,
        "title": "T\x01x",
        "status": "pronto",
        "source_count": 1,
        "created_at": now,
        "updated_at": now,
        "content": content,
        "citations": [_citation()],
        "withheld": ["Um ponto\x02 foi omitido."],
    }
    fields.update(extra)
    return ArtifactOut(**fields)


def test_a_control_character_stored_before_the_reader_does_not_break_the_docx():
    lines = [{"speaker": 1, "text": "a\x03b", "citations": [1]}]
    data, _ = render(_artifact("audio", {"title": "T", "lines": lines}), "docx", "N\x04")
    text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    assert "T x" in text and "Um ponto  foi omitido." in text and "a b" in text


def test_a_control_character_stored_before_the_reader_does_not_break_the_xlsx():
    """N-1: title, subtitle, withheld line, source title and excerpt all carry
    one; openpyxl raised on each of them before the second wall covered the
    spreadsheet."""
    cells = [
        {"status": "ok", "text": "f\x05", "citations": [1]},
        {"status": "ok", "text": "b", "citations": []},
    ]
    content = {"title": "T", "columns": ["C\x06", "D"], "rows": [{"cells": cells}]}
    artifact = _artifact("table", content)
    data, _ = render(artifact, "xlsx", "N\x04")
    workbook = load_workbook(io.BytesIO(data))
    cover = [row[0] for row in workbook["Aviso"].iter_rows(values_only=True)]
    assert "T x" in cover
    values = [
        cell
        for sheet in workbook.worksheets
        for row in sheet.iter_rows(values_only=True)
        for cell in row
        if isinstance(cell, str)
    ]
    assert "Um ponto  foi omitido." in values and "f " in values and "e " in values
    assert "C " in values


def test_a_control_character_does_not_break_the_infographic_svg():
    board = {
        "title": "T",
        "subtitle": "s\x05",
        "stats": [{"value": "210 GPa", "label": "l\x06", "citations": [1]}],
        "points": [{"heading": "h", "text": "a\x03b", "citations": [1]}],
        "steps": [],
    }
    svg = infographic_svg(_artifact("infographic", board, format="paisagem"), "sub\x08")
    root = ElementTree.fromstring(svg)  # well-formed: a browser opens it
    strings = [t.text or "" for t in root.iter("{http://www.w3.org/2000/svg}text")]
    assert "T x" in strings and "a b" in strings


def test_a_control_character_does_not_break_the_mind_map_svg():
    drawn = mindmap.layout({"label": "r\x02", "citations": [1], "children": []})
    node = drawn.nodes[0]
    shape = MindMapLayoutOut(
        width=drawn.width,
        height=drawn.height,
        nodes=[
            {
                "id": node.id,
                "parent": node.parent,
                "label": node.label,
                "lines": node.lines,
                "depth": node.depth,
                "x": node.x,
                "y": node.y,
                "width": node.width,
                "height": node.height,
                "citations": node.citations,
            }
        ],
        edges=[],
    )
    artifact = _artifact("mindmap", {"root": {}}, layout=shape)
    ElementTree.fromstring(mindmap_svg(artifact, "sub"))


# --- M-8: where the citation marks sit travels with the layout -----------------------------


def test_the_marks_are_placed_by_the_layout_s_style():
    data = layout(_board(title="T", subtitle=""), "paisagem").to_dict()
    for kind in ("stat", "point", "step"):
        assert all(
            data["styles"][kind][k] > 0 for k in ("marks_size", "marks_right", "marks_bottom")
        )
    board = _board(title="T", subtitle="")
    svg = infographic_svg(_artifact("infographic", board, format="paisagem", title="T"), "s")
    root = ElementTree.fromstring(svg)
    stat = next(b for b in layout(board, "paisagem").blocks if b.kind == "stat")
    style = STYLES["stat"]
    marks = next(
        t for t in root.iter("{http://www.w3.org/2000/svg}text") if (t.text or "") == "[1]"
    )
    assert int(marks.get("x")) == stat.x + stat.width - style.marks_right
    assert int(marks.get("y")) == stat.y + stat.height - style.marks_bottom
    assert marks.get("font-size") == str(style.marks_size)


# --- The reader's cap never cuts inside a figure -----------------------------------------


def test_the_reader_s_cap_never_cuts_a_figure():
    # A title whose 120th character falls inside "1 200": a character cap kept
    # "… 1 20" — figures no passage states.
    head = ("Aço " * 29).strip()  # 115 chars
    title = f"{head} 1 200 MPa"
    assert title[:MAX_TITLE].endswith("1 20")  # the old behaviour
    deck = read_deck({"title": title, "slides": []})
    assert deck["title"] == head
    assert not any(ch.isdigit() for ch in deck["title"])


def test_the_reader_s_cap_keeps_a_figure_that_fits_with_its_unit():
    text = ("palavra " * 48).strip() + " chega a 12.345.678 kWh"  # 383 + 23 = 406 chars
    board = read_infographic({"points": [{"text": text}]})
    kept = board["points"][0]["text"]
    assert len(kept) <= MAX_SHORT
    # "12.345.678 kWh" does not fit, so it goes whole: no stray "12.345".
    assert kept.endswith("chega a")
    assert "12" not in kept
    # And a text that fits is untouched.
    fits = read_infographic({"points": [{"text": "O aço chega a 1 200 MPa."}]})
    assert fits["points"][0]["text"] == "O aço chega a 1 200 MPa."


def test_a_single_piece_longer_than_the_cap_is_dropped_not_split():
    assert _cap("x" * (MAX_TITLE + 5), MAX_TITLE) == ""
    assert _cap("1" * 30, 10) == ""
    assert read_audio({"title": "9" * (MAX_TITLE + 1), "lines": []})["title"] == ""


def test_a_spoken_line_is_capped_on_a_boundary_too():
    text = "<b>" + "fala " * 119 + "1 200 MPa</b>"
    audio = read_audio({"title": "T", "lines": [{"speaker": 1, "text": text}]})
    line = audio["lines"][0]["text"]
    assert len(line) <= MAX_LINE
    assert not any(ch.isdigit() for ch in line)


@pytest.mark.parametrize("limit", range(1, 60))
def test_the_cap_never_prints_a_figure_the_text_does_not_write(limit):
    text = "A liga 7075 resiste a 1 200 MPa e pesa 2,81 g/cm³ a 20 °C"
    cut = _cap(text, limit)
    assert len(cut) <= limit
    assert text.startswith(cut)
    written = [m.group() for m in _NUMBER_TOKEN.finditer(text)]
    printed = [m.group() for m in _NUMBER_TOKEN.finditer(cut)]
    # Every figure printed is one the text writes, whole and in order.
    assert printed == written[: len(printed)], (limit, cut)


# --- the quiz's answer key follows the options that survive (N-2) ---------------------


def _quiz(options: list, answer: int) -> list[dict]:
    return read_studio(
        StudioRequest(tool="quiz", notebook_title="N", source_titles=(), passages=()),
        {
            "title": "Q",
            "questions": [
                {"prompt": "Qual?", "options": options, "answer_index": answer, "citations": [1]}
            ],
        },
    )["questions"]


def test_an_empty_option_before_the_answer_moves_the_answer_index():
    (question,) = _quiz(["A", "", "C", "D"], 2)
    assert question["options"] == ["A", "C", "D"]
    assert question["options"][question["answer_index"]] == "C"


def test_an_option_that_is_not_text_moves_the_answer_index_too():
    (question,) = _quiz([7, "B", "C"], 2)
    assert question["options"] == ["B", "C"]
    assert question["options"][question["answer_index"]] == "C"


def test_a_question_whose_correct_option_came_back_empty_is_dropped():
    # Before N-2 this was read as ["A", "C", "D"] with answer 1: "C" marked correct.
    assert _quiz(["A", "", "C", "D"], 1) == []
    assert _quiz(["A", "\x07", "C"], 1) == []


def test_a_question_left_with_one_option_is_dropped():
    assert _quiz(["", "B", ""], 1) == []
