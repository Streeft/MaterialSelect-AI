"""What changed between two releases of the official catalogue (D-108), end to end.

Two **fictitious** releases of one fictitious catalogue go through the real
importer of the D-102 — the bundle, the dry-run, the transactional commit — so
the diff reads exactly what an import stores. Between them, one record of each
kind: unchanged, removed, a number changed together with its unit, a value that
goes from declared missing to present (and a property that goes from not
registered to registered), a renamed record, and a new one whose name is a
spreadsheet formula. None of it is real material data.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.catalog.bundle import BundleValidationError, verify_bundle
from app.catalog.importer import (
    OfficialCatalogImporter,
    OfficialCatalogImportError,
    validate_semantics,
)
from app.db.clear_demo import clear_demo_data
from app.domain.curves import PointInput, SeriesInput, build_curve
from app.exporters.report import DEMO_DATA_NOTICE, LIMITATION_NOTICE
from app.models.catalog import CatalogDataset, CatalogImportRun, CatalogRecordRef
from app.models.enums import CurveKind
from app.models.material import Material
from app.models.material_class import MaterialClass
from app.models.material_property_value import MaterialPropertyValue
from app.models.property_definition import PropertyDefinition
from app.models.source import Source
from app.models.user import User
from app.repositories.curve_repository import curve_rows

LINEAGE = "catalogo-ficticio-teste"
FORMULA_NAME = "=1+1 Liga Fictícia E"


def _hash(value: str) -> str:
    return hashlib.sha256(value.encode()).hexdigest()


PROPERTY_DEFINITIONS = [
    # The seed's own definition, repeated exactly (the importer refuses drift).
    {
        "slug": "densidade",
        "name": "Densidade",
        "symbol": "ρ",
        "category": "FISICA",
        "physical_dimension": "[mass] / [length] ** 3",
        "canonical_unit": "kg/m**3",
        "accepted_units": ["kg/m**3", "g/cm**3"],
        "display_unit": "g/cm**3",
        "is_interval": False,
        "better_direction": "LOWER",
    },
    {
        "slug": "temp_minima_teste_diff",
        "name": "Temperatura mínima fictícia",
        "category": "TERMICA",
        "physical_dimension": "[temperature]",
        "canonical_unit": "K",
        "accepted_units": ["K", "degC"],
        "display_unit": "degC",
        "is_interval": False,
        "better_direction": "LOWER",
    },
]

CLASSES = [
    {"external_id": "mc-metais", "name": "Metais", "slug": "metais", "parent_external_id": None},
    {
        "external_id": "mc-ceram",
        "name": "Cerâmicas fictícias do diff",
        "slug": "ceramicas-ficticias-diff",
        "parent_external_id": None,
    },
]


def _material(external_id: str, name: str, klass: str = "mc-metais") -> dict:
    return {
        "external_id": external_id,
        "external_table": "MaterialUniverse",
        "gruid": f"G-{external_id}",
        "raw_sha256": _hash(f"{external_id}:{name}"),
        "name": name,
        "class_external_id": klass,
    }


def _scalar(material: str, slug: str, value: float, unit: str) -> dict:
    return {
        "material_external_id": material,
        "property_slug": slug,
        "value_kind": "scalar",
        "value": value,
        "original_unit": unit,
    }


def _missing(material: str, slug: str) -> dict:
    return {"material_external_id": material, "property_slug": slug, "value_kind": "missing"}


TEMP = "temp_minima_teste_diff"

RELEASE_1 = {
    "materials.ndjson": [
        _material("mat-a", "Liga Fictícia A"),
        _material("mat-b", "Liga Fictícia B"),
        _material("mat-c", "Liga Fictícia C"),
        _material("mat-d", "Liga Fictícia D"),
        _material("mat-r", "Liga Fictícia R"),
    ],
    "material_values.ndjson": [
        _scalar("mat-a", "densidade", 7850.0, "kg/m**3"),
        _scalar("mat-a", TEMP, -40.0, "degC"),
        _scalar("mat-b", "densidade", 4500.0, "kg/m**3"),
        _scalar("mat-c", "densidade", 7850.0, "kg/m**3"),
        _scalar("mat-c", TEMP, -40.0, "degC"),
        _missing("mat-d", "densidade"),
        _scalar("mat-r", "densidade", 1000.0, "kg/m**3"),
    ],
}

RELEASE_2 = {
    "materials.ndjson": [
        _material("mat-a", "Liga Fictícia A"),
        _material("mat-c", "Liga Fictícia C"),
        _material("mat-d", "Liga Fictícia D"),
        # Renamed: the same identity, so "alterado", never removed + new.
        {**_material("mat-r", "Liga Fictícia R renomeada")},
        _material("mat-e", FORMULA_NAME, klass="mc-ceram"),
    ],
    "material_values.ndjson": [
        _scalar("mat-a", "densidade", 7850.0, "kg/m**3"),
        _scalar("mat-a", TEMP, -40.0, "degC"),
        # The number and the unit change: 7850 kg/m³ → 7,9 g/cm³ (7900 kg/m³).
        _scalar("mat-c", "densidade", 7.9, "g/cm**3"),
        _scalar("mat-c", TEMP, -50.0, "degC"),
        # Declared missing → a number; not registered → registered.
        _scalar("mat-d", "densidade", 2700.0, "kg/m**3"),
        _scalar("mat-d", TEMP, -10.0, "degC"),
        _scalar("mat-r", "densidade", 1000.0, "kg/m**3"),
        _scalar("mat-e", "densidade", 3900.0, "kg/m**3"),
    ],
}


def write_bundle(
    tmp_path: Path,
    slug: str,
    release: str,
    files: dict[str, list[dict]],
    *,
    lineage: str | None = LINEAGE,
) -> Path:
    contents = {
        "material_classes.ndjson": CLASSES,
        "property_definitions.ndjson": PROPERTY_DEFINITIONS,
        **files,
    }
    encoded = {
        name: "".join(
            json.dumps(r, ensure_ascii=False, sort_keys=True) + "\n" for r in rows
        ).encode()
        for name, rows in contents.items()
    }
    dataset = {
        "slug": slug,
        "name": "Catálogo Fictício de Teste",
        "release": release,
        "source_sha256": _hash(f"source:{slug}"),
        "license_label": "Fixture fictícia de teste",
        "provenance": "Dados inventados para testar o diff entre releases.",
    }
    if lineage is not None:
        dataset["lineage"] = lineage
    manifest = {
        "schema_version": 1,
        "dataset": dataset,
        "files": {
            name: {"sha256": hashlib.sha256(raw).hexdigest(), "count": len(contents[name])}
            for name, raw in encoded.items()
        },
    }
    path = tmp_path / f"{slug}.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, sort_keys=True))
        for name, raw in encoded.items():
            archive.writestr(name, raw)
    return path


def import_bundle(db: Session, path: Path, reviewer: User) -> dict:
    bundle = verify_bundle(path)
    validate_semantics(bundle)
    return OfficialCatalogImporter(db, bundle, reviewer_email=reviewer.email).run()


@pytest.fixture()
def reviewer(db_session: Session) -> User:
    clear_demo_data(db_session)
    user = User(google_sub="diff-reviewer", email="diff-reviewer@example.com", name="Revisor")
    db_session.add(user)
    db_session.flush()
    return user


@pytest.fixture()
def two_releases(tmp_path: Path, db_session: Session, reviewer: User) -> tuple[str, str]:
    import_bundle(db_session, write_bundle(tmp_path, "ficticio-r1", "R1", RELEASE_1), reviewer)
    import_bundle(db_session, write_bundle(tmp_path, "ficticio-r2", "R2", RELEASE_2), reviewer)
    return "ficticio-r1", "ficticio-r2"


def _items(body: dict) -> dict[str, dict]:
    return {item["external_record_id"]: item for item in body["items"]}


def _change(item: dict, field: str) -> dict:
    return next(change for change in item["changes"] if change["field"] == field)


# -- the releases ----------------------------------------------------------------


def test_the_release_list_carries_lineage_order_and_provenance(client, two_releases) -> None:
    response = client.get("/api/catalogo/releases")
    assert response.status_code == 200
    releases = {r["slug"]: r for r in response.json()}
    r1, r2 = releases["ficticio-r1"], releases["ficticio-r2"]
    assert r1["lineage"] == r2["lineage"] == LINEAGE
    assert (r1["material_count"], r2["material_count"]) == (5, 5)
    assert r1["previous_slug"] is None and r2["previous_slug"] == "ficticio-r1"
    assert r2["imported_at"] is not None and len(r2["manifest_sha256"]) == 64
    assert r2["is_demo"] is False and r2["license_label"] == "Fixture fictícia de teste"


def test_the_importer_keeps_each_release_in_its_own_rows(db_session: Session, two_releases) -> None:
    """The fact the diff stands on: release 2 wrote new rows, release 1's are intact."""
    refs = db_session.execute(
        select(
            CatalogDataset.slug, CatalogRecordRef.external_record_id, CatalogRecordRef.material_id
        )
        .join(CatalogDataset, CatalogDataset.id == CatalogRecordRef.dataset_id)
        .where(CatalogDataset.lineage == LINEAGE)
    ).all()
    by_release: dict[str, dict[str, int]] = {}
    for slug, external_id, material_id in refs:
        by_release.setdefault(slug, {})[external_id] = material_id
    assert by_release["ficticio-r1"]["mat-c"] != by_release["ficticio-r2"]["mat-c"]
    old_c = db_session.execute(
        select(MaterialPropertyValue.value_scalar, MaterialPropertyValue.original_unit)
        .join(PropertyDefinition, PropertyDefinition.id == MaterialPropertyValue.property_id)
        .where(
            MaterialPropertyValue.material_id == by_release["ficticio-r1"]["mat-c"],
            PropertyDefinition.slug == "densidade",
        )
    ).one()
    assert tuple(old_c) == (7850.0, "kg/m**3")


# -- the diff --------------------------------------------------------------------


def test_the_diff_classifies_every_record(client, two_releases) -> None:
    response = client.get("/api/catalogo/releases/ficticio-r1/diff/ficticio-r2")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["lineage"] == LINEAGE and body["is_demo"] is False
    assert [(c["status"], c["count"]) for c in body["counts"]] == [
        ("alterado", 3),
        ("novo", 1),
        ("desativado", 1),
        ("inalterado", 1),
    ]
    assert body["total"] == 6 and body["filtered_total"] == 6
    items = _items(body)
    assert items["mat-a"]["status"] == "inalterado" and items["mat-a"]["changes"] == []
    assert items["mat-b"]["status"] == "desativado" and items["mat-b"]["target"] is None
    assert items["mat-e"]["status"] == "novo" and items["mat-e"]["base"] is None
    assert items["mat-e"]["target"]["name"] == FORMULA_NAME
    # Order: changed, new, removed, unchanged — then by current name.
    assert [i["external_record_id"] for i in body["items"]] == [
        "mat-c",
        "mat-d",
        "mat-r",
        "mat-e",
        "mat-b",
        "mat-a",
    ]
    assert {c["slug"]: c["count"] for c in body["classes"]} == {
        "metais": 5,
        "ceramicas-ficticias-diff": 1,
    }


def test_a_number_that_changed_with_its_unit_keeps_all_three_units(client, two_releases) -> None:
    body = client.get("/api/catalogo/releases/ficticio-r1/diff/ficticio-r2").json()
    density = _change(_items(body)["mat-c"], "propriedade:densidade")
    assert density["kind"] == "valor"
    before, after = density["before"], density["after"]
    assert before["original"]["value"] == 7850.0 and before["original"]["unit"] == "kg/m**3"
    assert after["original"]["value"] == 7.9 and after["original"]["unit"] == "g/cm**3"
    assert after["canonical"]["value"] == pytest.approx(7900.0)
    assert after["canonical"]["unit"] == "kg/m**3"
    # Reading unit: the property's convention (g/cm³), via from_canonical.
    assert density["reading_unit"] == "g/cm**3" and density["reading_unit_label"] == "g/cm³"
    assert before["reading"]["value"] == pytest.approx(7.85)
    assert after["reading"]["value"] == pytest.approx(7.9)
    assert after["conversion_method"] and "g/cm**3" in after["conversion_method"]

    temperature = _change(_items(body)["mat-c"], f"propriedade:{TEMP}")
    assert temperature["kind"] == "valor"
    assert temperature["before"]["reading"]["value"] == pytest.approx(-40.0)
    assert temperature["after"]["canonical"]["value"] == pytest.approx(223.15)


def test_the_reading_unit_is_chosen_in_the_url(client, two_releases) -> None:
    body = client.get(
        "/api/catalogo/releases/ficticio-r1/diff/ficticio-r2",
        params={"unidades": f"densidade:kg/m**3,{TEMP}:K"},
    ).json()
    item = _items(body)["mat-c"]
    density = _change(item, "propriedade:densidade")
    assert density["reading_unit"] == "kg/m**3"
    assert density["after"]["reading"]["value"] == pytest.approx(7900.0)
    assert _change(item, f"propriedade:{TEMP}")["before"]["reading"]["value"] == pytest.approx(
        233.15
    )
    refused = client.get(
        "/api/catalogo/releases/ficticio-r1/diff/ficticio-r2",
        params={"unidades": "densidade:lb"},
    )
    assert refused.status_code == 400 and "não é admitida" in refused.json()["detail"]


def test_absence_is_a_state_written_in_words_never_zero(client, two_releases) -> None:
    item = _items(client.get("/api/catalogo/releases/ficticio-r1/diff/ficticio-r2").json())["mat-d"]
    density = _change(item, "propriedade:densidade")
    assert density["kind"] == "ausencia"
    assert density["before"]["state"] == "ausente"
    assert density["before"]["state_label"] == "declarado ausente pela fonte"
    assert density["before"]["original"] is None
    assert density["before"]["canonical"] is None
    assert density["before"]["reading"] is None
    assert density["after"]["reading"]["value"] == pytest.approx(2.7)

    temperature = _change(item, f"propriedade:{TEMP}")
    assert temperature["kind"] == "ausencia"
    assert temperature["before"]["state"] == "nao_cadastrado"
    assert temperature["before"]["reading"] is None


def test_a_renamed_record_is_changed_not_removed_and_new(client, two_releases) -> None:
    item = _items(client.get("/api/catalogo/releases/ficticio-r1/diff/ficticio-r2").json())["mat-r"]
    assert item["status"] == "alterado"
    assert [(c["field"], c["before_text"], c["after_text"]) for c in item["changes"]] == [
        ("nome", "Liga Fictícia R", "Liga Fictícia R renomeada")
    ]
    assert item["raw_record_changed"] is True


def test_filters_by_kind_and_class(client, two_releases) -> None:
    url = "/api/catalogo/releases/ficticio-r1/diff/ficticio-r2"
    body = client.get(url, params={"tipo": "novo"}).json()
    assert [i["external_record_id"] for i in body["items"]] == ["mat-e"]
    assert body["filtered_total"] == 1 and body["total"] == 6
    assert body["filters"] == {"tipo": "novo", "classe": None}

    body = client.get(url, params={"classe": "ceramicas-ficticias-diff"}).json()
    assert [i["external_record_id"] for i in body["items"]] == ["mat-e"]
    body = client.get(url, params={"classe": "metais", "tipo": "inalterado"}).json()
    assert [i["external_record_id"] for i in body["items"]] == ["mat-a"]


def test_pagination(client, two_releases) -> None:
    url = "/api/catalogo/releases/ficticio-r1/diff/ficticio-r2"
    pages = [client.get(url, params={"pagina": n, "por_pagina": 4}).json() for n in (1, 2)]
    assert [p["page_count"] for p in pages] == [2, 2]
    assert [len(p["items"]) for p in pages] == [4, 2]
    every = [i["external_record_id"] for p in pages for i in p["items"]]
    assert every == [i["external_record_id"] for i in client.get(url).json()["items"]]
    for params, message in (
        ({"pagina": 3, "por_pagina": 4}, "não existe"),
        ({"pagina": 0}, "Página inválida"),
        ({"por_pagina": 0}, "1 a 200"),
        ({"por_pagina": 201}, "1 a 200"),
    ):
        response = client.get(url, params=params)
        assert response.status_code == 400 and message in response.json()["detail"]


def test_one_record_by_its_external_identity(client, two_releases) -> None:
    url = "/api/catalogo/releases/ficticio-r1/diff/ficticio-r2/registro"
    item = client.get(url, params={"tabela": "MaterialUniverse", "id": "mat-c"}).json()
    assert item["status"] == "alterado" and item["change_count"] == 2
    missing = client.get(url, params={"tabela": "MaterialUniverse", "id": "nao-existe"})
    assert missing.status_code == 404 and "nenhuma das duas releases" in missing.json()["detail"]
    # The name is never an identity: asking by name finds nothing.
    by_name = client.get(url, params={"tabela": "MaterialUniverse", "id": "Liga Fictícia C"})
    assert by_name.status_code == 404
    incomplete = client.get(url, params={"tabela": "MaterialUniverse"})
    assert incomplete.status_code == 400 and "identidade externa" in incomplete.json()["detail"]


def test_the_diff_is_idempotent_and_writes_nothing(
    client, tmp_path: Path, db_session: Session, reviewer: User, two_releases
) -> None:
    def state() -> tuple:
        values = db_session.execute(
            select(
                MaterialPropertyValue.id,
                MaterialPropertyValue.value_scalar,
                MaterialPropertyValue.original_unit,
                MaterialPropertyValue.normalized_value,
                MaterialPropertyValue.is_missing,
            ).order_by(MaterialPropertyValue.id)
        ).all()
        counts = tuple(
            db_session.scalar(select(func.count()).select_from(model))
            for model in (Material, CatalogRecordRef, CatalogDataset, CatalogImportRun, Source)
        )
        return tuple(map(tuple, values)), counts

    url = "/api/catalogo/releases/ficticio-r1/diff/ficticio-r2"
    before = state()
    first = client.get(url).json()
    assert client.get(url).json() == first
    assert client.get("/api/exports/catalogo/releases/ficticio-r1/diff/ficticio-r2.csv")
    assert state() == before

    # Re-importing release 1 after release 2 vigorou is refused (D-114: a retired
    # release is not re-activated) and rolls back: nothing stored changes.
    with pytest.raises(OfficialCatalogImportError, match="substituída"):
        import_bundle(db_session, write_bundle(tmp_path, "ficticio-r1", "R1", RELEASE_1), reviewer)
    assert state() == before
    assert client.get(url).json() == first

    # Re-importing release 2 (the importer is idempotent) changes neither the
    # stored release nor the answer — only the audit of its imports (a new run).
    result = import_bundle(
        db_session, write_bundle(tmp_path, "ficticio-r2", "R2", RELEASE_2), reviewer
    )
    assert result["counts"].get("materials_created", 0) == 0
    again = client.get(url).json()
    assert again["target"]["imported_at"] != first["target"]["imported_at"]
    assert {k: v for k, v in again.items() if k != "target"} == {
        k: v for k, v in first.items() if k != "target"
    }


def test_a_release_cannot_change_its_lineage(
    tmp_path: Path, db_session: Session, reviewer: User, two_releases
) -> None:
    path = write_bundle(tmp_path, "ficticio-r1", "R1", RELEASE_1, lineage="outra-linha")
    with pytest.raises(OfficialCatalogImportError, match="linha"):
        import_bundle(db_session, path, reviewer)


def test_the_bundle_refuses_a_lineage_that_is_not_a_slug(tmp_path: Path) -> None:
    path = write_bundle(tmp_path, "ficticio-x", "X", RELEASE_1, lineage="Catálogo Fictício")
    with pytest.raises(BundleValidationError, match="lineage"):
        verify_bundle(path)


def test_the_reverse_direction_swaps_new_and_removed(client, two_releases) -> None:
    body = client.get("/api/catalogo/releases/ficticio-r2/diff/ficticio-r1").json()
    items = _items(body)
    assert items["mat-b"]["status"] == "novo" and items["mat-e"]["status"] == "desativado"


# -- refusals --------------------------------------------------------------------


def test_errors_in_portuguese(client, tmp_path: Path, db_session: Session, reviewer, two_releases):
    import_bundle(
        db_session,
        write_bundle(tmp_path, "outro-r1", "R1", RELEASE_1, lineage="outro-catalogo"),
        reviewer,
    )
    import_bundle(
        db_session, write_bundle(tmp_path, "sem-linha", "R1", RELEASE_1, lineage=None), reviewer
    )
    cases = [
        ("/api/catalogo/releases/nao-existe/diff/ficticio-r2", 404, "não encontrada"),
        ("/api/catalogo/releases/ficticio-r1/diff/ficticio-r1", 400, "duas releases diferentes"),
        ("/api/catalogo/releases/ficticio-r1/diff/outro-r1", 400, "catálogos diferentes"),
        ("/api/catalogo/releases/ficticio-r1/diff/sem-linha", 400, "não declara"),
        ("/api/catalogo/releases/ficticio-r1/diff/ficticio-r2?tipo=removido", 400, "Admitidos"),
        ("/api/catalogo/releases/ficticio-r1/diff/ficticio-r2?classe=nada", 400, "desconhecida"),
        ("/api/exports/catalogo/releases/ficticio-r1/diff/ficticio-r2.pdf", 400, "csv, xlsx"),
        ("/api/exports/catalogo/releases/ficticio-r1/diff/outro-r1.csv", 400, "catálogos"),
    ]
    for url, status, message in cases:
        response = client.get(url)
        assert response.status_code == status, (url, response.text)
        assert message in response.json()["detail"], url


def test_the_routes_require_login(anon_client) -> None:
    for url in (
        "/api/catalogo/releases",
        "/api/catalogo/releases/a/diff/b",
        "/api/exports/catalogo/releases/a/diff/b.csv",
    ):
        assert anon_client.get(url).status_code == 401


# -- visibility (D-62) -----------------------------------------------------------


PRIVATE = "Registro próprio confidencial do diff"


def test_a_private_record_never_enters_the_diff(
    client, login_as, other_user: User, db_session: Session, two_releases
) -> None:
    """Even an identity row pointing at somebody's own record is left out — for
    the owner too: the diff describes the shared catalogue, nothing else."""
    metals = db_session.execute(
        select(MaterialClass).where(MaterialClass.slug == "metais")
    ).scalar_one()
    private = Material(
        name=PRIVATE, class_id=metals.id, keywords=[], owner_id=other_user.id, is_demo=False
    )
    db_session.add(private)
    db_session.flush()
    target = db_session.execute(
        select(CatalogDataset).where(CatalogDataset.slug == "ficticio-r2")
    ).scalar_one()
    db_session.add(
        CatalogRecordRef(
            dataset_id=target.id,
            external_table="MaterialUniverse",
            external_record_id="mat-privado",
            raw_record_sha256=_hash("privado"),
            material_id=private.id,
        )
    )
    db_session.flush()

    urls = (
        "/api/catalogo/releases",
        "/api/catalogo/releases/ficticio-r1/diff/ficticio-r2",
        "/api/catalogo/releases/ficticio-r1/diff/ficticio-r2/registro?tabela=MaterialUniverse&id=mat-privado",
        "/api/exports/catalogo/releases/ficticio-r1/diff/ficticio-r2.csv",
    )
    for viewer in (None, other_user):
        if viewer is None:
            bodies = [client.get(url) for url in urls]
        else:
            with login_as(viewer):
                bodies = [client.get(url) for url in urls]
        assert all(PRIVATE not in b.text for b in bodies)
        # The 404 of the record route echoes the id that was asked; the others
        # must not know it exists.
        assert all("mat-privado" not in bodies[i].text for i in (0, 1, 3))
        assert bodies[1].json()["total"] == 6
        assert bodies[2].status_code == 404
        assert {r["slug"]: r for r in bodies[0].json()}["ficticio-r2"]["material_count"] == 5


# -- export ----------------------------------------------------------------------


def _csv_rows(text: str) -> list[list[str]]:
    return list(csv.reader(io.StringIO(text.lstrip("﻿"))))


def test_the_csv_carries_notice_provenance_and_safe_cells(client, two_releases) -> None:
    response = client.get("/api/exports/catalogo/releases/ficticio-r1/diff/ficticio-r2.csv")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["content-disposition"].startswith("attachment")
    text = response.text
    rows = _csv_rows(text)
    flat = [cell for row in rows for cell in row]

    assert LIMITATION_NOTICE in flat
    assert DEMO_DATA_NOTICE not in flat
    # Provenance: both releases, licence, lineage, the sources the values cite.
    assert ["Slug da release", "ficticio-r1", "ficticio-r2"] in rows
    assert ["Licença", "Fixture fictícia de teste", "Fixture fictícia de teste"] in rows
    sources = next(row for row in rows if row and row[0] == "Fontes citadas pelos valores")
    assert sources[1] == "Catálogo Fictício de Teste — R1"
    assert sources[2] == "Catálogo Fictício de Teste — R2"
    # A name written as a formula leaves as text.
    assert "'" + FORMULA_NAME in flat and FORMULA_NAME not in flat

    header = next(row for row in rows if row and row[0] == "Tabela externa" and "Campo" in row)
    changes = [dict(zip(header, row, strict=False)) for row in rows if len(row) == len(header)]
    by_field = {(c["ID externo"], c["Campo"]): c for c in changes if c is not None}
    temperature = by_field[("mat-c", "Temperatura mínima fictícia")]
    # Negative numbers stay numbers (no apostrophe), in the reading unit.
    assert temperature["Antes (leitura)"] == "-40.0"
    assert temperature["Depois (leitura)"] == "-50.0"
    assert temperature["Unidade de leitura"] == "°C"
    assert temperature["Antes: unidade original"] == "degC"
    assert temperature["Unidade canônica"] == "K"
    density = by_field[("mat-d", "Densidade")]
    assert density["Antes (leitura)"] == "declarado ausente pela fonte"
    assert density["Antes (canônico)"] == "declarado ausente pela fonte"
    assert density["Antes mín. (leitura)"] == "declarado ausente pela fonte"
    unregistered = by_field[("mat-d", "Temperatura mínima fictícia")]
    assert unregistered["Antes (leitura)"] == "não cadastrado nesta release"
    renamed = by_field[("mat-r", "Nome")]
    assert (renamed["Antes (leitura)"], renamed["Depois (leitura)"]) == (
        "Liga Fictícia R",
        "Liga Fictícia R renomeada",
    )
    records_header = next(row for row in rows if row and row[0] == "Situação" and "GRUID" in row)
    removed = next(row for row in rows if len(row) == len(records_header) and row[2] == "mat-b")
    assert removed[5] == "não está nesta release"


def test_the_export_applies_the_filters(client, two_releases) -> None:
    text = client.get(
        "/api/exports/catalogo/releases/ficticio-r1/diff/ficticio-r2.csv", params={"tipo": "novo"}
    ).text
    assert "mat-e" in text and "mat-c" not in text and "tipo de mudança = Novo" in text


def test_the_xlsx_keeps_negative_numbers_numeric(client, two_releases) -> None:
    response = client.get("/api/exports/catalogo/releases/ficticio-r1/diff/ficticio-r2.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert {"Aviso", "Releases", "Resumo", "Registros", "Alterações"} <= set(workbook.sheetnames)
    sheet = workbook["Alterações"]
    header = [
        cell.value for cell in next(r for r in sheet.iter_rows() if r[0].value == "Tabela externa")
    ]
    rows = [
        dict(zip(header, (c.value for c in row), strict=False))
        for row in sheet.iter_rows()
        if row[1].value == "mat-c" and row[3].value == "Temperatura mínima fictícia"
    ]
    assert rows and rows[0]["Antes (leitura)"] == pytest.approx(-40.0)
    assert isinstance(rows[0]["Depois (canônico)"], float)
    cover = [c.value for c in workbook["Aviso"]["A"]]
    assert LIMITATION_NOTICE in cover


# -- demo releases (D-108 contract for the seed) ---------------------------------


def _demo_release(db: Session, slug: str, density: float | None) -> CatalogDataset:
    """What the demo seed must write: a demo dataset, a demo material per release,
    a demo source, the identity row and the value — never through the importer."""
    metals = db.execute(select(MaterialClass).where(MaterialClass.slug == "metais")).scalar_one()
    prop = db.execute(
        select(PropertyDefinition).where(PropertyDefinition.slug == "densidade")
    ).scalar_one()
    source = Source(label=f"Fonte demo {slug}", is_demo=True)
    dataset = CatalogDataset(
        slug=slug,
        name="Catálogo Demo",
        release=slug[-2:].upper(),
        lineage="catalogo-demo-teste",
        source_sha256=_hash(slug),
        license_label="Dado fictício de demonstração",
        is_demo=True,
    )
    material = Material(name="Liga Demo Diff", class_id=metals.id, keywords=[], is_demo=True)
    db.add_all([source, dataset, material])
    db.flush()
    db.add(
        CatalogRecordRef(
            dataset_id=dataset.id,
            external_table="MaterialUniverse",
            external_record_id="demo-1",
            raw_record_sha256=_hash("demo-1"),
            material_id=material.id,
        )
    )
    db.add(
        MaterialPropertyValue(
            material_id=material.id,
            property_id=prop.id,
            value_scalar=density,
            original_unit="kg/m**3" if density is not None else None,
            normalized_value=density,
            canonical_unit="kg/m**3" if density is not None else None,
            is_missing=density is None,
            source_id=source.id,
        )
    )
    db.flush()
    return dataset


def test_two_demo_releases_compare_and_clear_demo_removes_them(client, db_session: Session) -> None:
    _demo_release(db_session, "demo-diff-r1", None)
    _demo_release(db_session, "demo-diff-r2", 2700.0)
    body = client.get("/api/catalogo/releases/demo-diff-r1/diff/demo-diff-r2").json()
    assert body["is_demo"] is True
    (item,) = body["items"]
    assert _change(item, "propriedade:densidade")["kind"] == "ausencia"
    csv_text = client.get("/api/exports/catalogo/releases/demo-diff-r1/diff/demo-diff-r2.csv").text
    assert DEMO_DATA_NOTICE in csv_text

    removed = clear_demo_data(db_session)
    db_session.flush()
    assert removed["catalog_releases"] == 2
    assert (
        db_session.scalar(
            select(func.count(CatalogDataset.id)).where(CatalogDataset.is_demo.is_(True))
        )
        == 0
    )
    assert (
        db_session.scalar(
            select(func.count(CatalogRecordRef.id)).where(
                CatalogRecordRef.external_record_id == "demo-1"
            )
        )
        == 0
    )
    assert client.get("/api/catalogo/releases/demo-diff-r1/diff/demo-diff-r2").status_code == 404


def test_a_demo_release_never_compares_with_a_real_one(
    client, db_session: Session, tmp_path: Path, reviewer: User
) -> None:
    import_bundle(db_session, write_bundle(tmp_path, "ficticio-r1", "R1", RELEASE_1), reviewer)
    demo = _demo_release(db_session, "demo-misto", 1.0)
    demo.lineage = LINEAGE
    db_session.flush()
    response = client.get("/api/catalogo/releases/ficticio-r1/diff/demo-misto")
    assert response.status_code == 400 and "fictícia" in response.json()["detail"]


def test_the_official_import_refuses_while_a_demo_release_exists(
    db_session: Session, tmp_path: Path, reviewer: User
) -> None:
    _demo_release(db_session, "demo-bloqueia", 1.0)
    # The demo material and source would block too; remove them to isolate the release.
    db_session.execute(
        CatalogRecordRef.__table__.delete().where(CatalogRecordRef.external_record_id == "demo-1")
    )
    for material in db_session.execute(
        select(Material).where(Material.name == "Liga Demo Diff")
    ).scalars():
        db_session.execute(
            MaterialPropertyValue.__table__.delete().where(
                MaterialPropertyValue.material_id == material.id
            )
        )
        db_session.delete(material)
    db_session.execute(Source.__table__.delete().where(Source.label == "Fonte demo demo-bloqueia"))
    db_session.flush()
    with pytest.raises(OfficialCatalogImportError, match="catalog_datasets"):
        import_bundle(db_session, write_bundle(tmp_path, "ficticio-r1", "R1", RELEASE_1), reviewer)


def test_clear_demo_refuses_a_demo_release_cited_by_a_real_curve(db_session: Session) -> None:
    dataset = _demo_release(db_session, "demo-curva", 1.0)
    metals = db_session.execute(
        select(MaterialClass.id).where(MaterialClass.slug == "metais")
    ).scalar_one()
    real = Material(name="Material real do teste", class_id=metals, keywords=[], is_demo=False)
    real_source = Source(label="Fonte real do teste de curva", is_demo=False)
    db_session.add_all([real, real_source])
    db_session.flush()
    db_session.add(
        curve_rows(
            build_curve(
                CurveKind.TENSAO_DEFORMACAO,
                x_quantity="deformacao",
                x_unit="%",
                y_quantity="tensao",
                y_unit="MPa",
                series=[SeriesInput(points=[PointInput(0, 0), PointInput(1, 300)])],
            ),
            material_id=real.id,
            title="Curva real ligada a release demo",
            source_id=real_source.id,
            is_demo=False,
            dataset_id=dataset.id,
            external_id="curva-1",
            raw_sha256=_hash("curva-1"),
        )
    )
    db_session.flush()
    with pytest.raises(RuntimeError, match="preservar proveniência"):
        clear_demo_data(db_session)
