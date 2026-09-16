"""Tests for the DOCX export renderer and endpoints."""

from __future__ import annotations

import io

from docx import Document
from fastapi.testclient import TestClient

from app.exporters.docx import to_docx
from app.exporters.report import (
    DEMO_DATA_NOTICE,
    LIMITATION_NOTICE,
    OWN_RECORD_NOTICE,
    REPRODUCIBILITY_NOTICE,
    Report,
    Sheet,
    standard_notices,
)
from app.routers.exports import DOCX_MEDIA_TYPE


def _sample_report(*, demo: bool = True, own: bool = False) -> Report:
    return Report(
        title="Relatório de Seleção Teste",
        subtitle="Subtítulo do estudo",
        notices=standard_notices(includes_demo_data=demo, includes_own_records=own),
        sheets=[
            Sheet(
                name="Candidatos",
                header=["Material", "Densidade", "Módulo de Young"],
                rows=[
                    ["Aço Demo A", 7800.0, 210.0],
                    ["Cerâmica Demo D", 3900.0, None],
                ],
                notes=["Estudo gerado para testes unitários."],
            ),
            Sheet(
                name="Seção Vazia",
                header=["Coluna"],
                rows=[],
                notes=[],
            ),
        ],
        responsible="Eng. Responsável",
    )


class TestDocxRendering:
    def test_to_docx_creates_valid_document(self) -> None:
        data = to_docx(_sample_report())
        assert isinstance(data, bytes)
        assert len(data) > 0

        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Relatório de Seleção Teste" in text
        assert "Subtítulo do estudo" in text
        assert "Eng. Responsável" in text

    def test_to_docx_includes_mandatory_notices(self) -> None:
        doc = Document(io.BytesIO(to_docx(_sample_report(demo=True, own=True))))
        text = ""
        for p in doc.paragraphs:
            text += p.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + "\n"

        assert LIMITATION_NOTICE in text
        assert REPRODUCIBILITY_NOTICE in text
        assert DEMO_DATA_NOTICE in text
        assert OWN_RECORD_NOTICE in text

    def test_to_docx_omits_demo_notice_when_not_applicable(self) -> None:
        doc = Document(io.BytesIO(to_docx(_sample_report(demo=False, own=False))))
        text = ""
        for p in doc.paragraphs:
            text += p.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + "\n"

        assert LIMITATION_NOTICE in text
        assert DEMO_DATA_NOTICE not in text
        assert OWN_RECORD_NOTICE not in text

    def test_to_docx_renders_tables_and_headers(self) -> None:
        doc = Document(io.BytesIO(to_docx(_sample_report())))
        data_tables = [t for t in doc.tables if len(t.rows) > 1]
        assert len(data_tables) >= 1

        table = data_tables[0]
        header_cells = [cell.text for cell in table.rows[0].cells]
        assert header_cells == ["Material", "Densidade", "Módulo de Young"]

        row1_cells = [cell.text for cell in table.rows[1].cells]
        assert row1_cells == ["Aço Demo A", "7800", "210"]

        # Missing value must be 'ausente', not zero or blank
        row2_cells = [cell.text for cell in table.rows[2].cells]
        assert row2_cells == ["Cerâmica Demo D", "3900", "ausente"]

    def test_to_docx_renders_empty_section_note(self) -> None:
        doc = Document(io.BytesIO(to_docx(_sample_report())))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Nenhum registro nesta seção." in text

    def test_to_docx_renders_ai_narrative_when_present(self) -> None:
        report = _sample_report()
        report.narrative = ["Interpretação da IA sobre o ranking."]
        report.narrative_caveats = ["Ressalva número um."]
        report.narrative_note = "Nota sobre modelo."

        doc = Document(io.BytesIO(to_docx(report)))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Interpretação técnica (IA)" in text
        assert "Interpretação da IA sobre o ranking." in text
        assert "Ressalva número um." in text
        assert "Nota sobre modelo." in text


class TestDocxEndpoints:
    def test_catalogue_docx_download_headers(self, client: TestClient) -> None:
        response = client.get("/api/exports/catalogo.docx")
        assert response.status_code == 200, response.text
        assert response.headers["content-type"].startswith(DOCX_MEDIA_TYPE)
        assert "attachment" in response.headers["content-disposition"]
        assert "catalogo.docx" in response.headers["content-disposition"]

        doc = Document(io.BytesIO(response.content))
        all_text = "\n".join(p.text for p in doc.paragraphs)
        for t in doc.tables:
            for r in t.rows:
                for c in r.cells:
                    all_text += "\n" + c.text
        assert LIMITATION_NOTICE in all_text

    def test_study_docx_download_headers(self, client: TestClient) -> None:
        study_id = client.post(
            "/api/selection/studies",
            json={
                "name": "Estudo DOCX Exportável",
                "function_text": "Tirante",
                "objective_text": "Minimizar massa",
                "combinator": "AND",
                "constraints": [],
            },
        ).json()["id"]

        response = client.get(f"/api/exports/estudos/{study_id}.docx")
        assert response.status_code == 200, response.text
        assert response.headers["content-type"].startswith(DOCX_MEDIA_TYPE)
        assert "attachment" in response.headers["content-disposition"]

        doc = Document(io.BytesIO(response.content))
        all_text = "\n".join(p.text for p in doc.paragraphs)
        for t in doc.tables:
            for r in t.rows:
                for c in r.cells:
                    all_text += "\n" + c.text
        assert LIMITATION_NOTICE in all_text
