"""Tests for the PPTX export renderer and endpoints."""

from __future__ import annotations

import io

from fastapi.testclient import TestClient
from pptx import Presentation

from app.exporters.pptx import to_pptx
from app.exporters.report import (
    DEMO_DATA_NOTICE,
    LIMITATION_NOTICE,
    OWN_RECORD_NOTICE,
    REPRODUCIBILITY_NOTICE,
    Report,
    Sheet,
    standard_notices,
)
from app.routers.exports import PPTX_MEDIA_TYPE


def _sample_report(*, demo: bool = True, own: bool = False) -> Report:
    return Report(
        title="Relatório de Seleção Teste",
        subtitle="Subtítulo do estudo",
        notices=standard_notices(includes_demo_data=demo, includes_own_records=own),
        sheets=[
            Sheet(
                name="Candidatos",
                header=["Material", "Densidade", "Módulo de Young"],
                rows=[
                    ["Aço Demo A", 7800.0, 210.0],
                    ["Cerâmica Demo D", 3900.0, None],
                ],
                notes=["Estudo gerado para testes unitários."],
            ),
            Sheet(
                name="Seção Vazia",
                header=["Coluna"],
                rows=[],
                notes=[],
            ),
        ],
        responsible="Eng. Responsável",
    )


def _extract_all_text(prs: Presentation) -> str:
    texts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        texts.append(cell.text)
    return "\n".join(texts)


class TestPptxRendering:
    def test_to_pptx_creates_valid_presentation(self) -> None:
        data = to_pptx(_sample_report())
        assert isinstance(data, bytes)
        assert len(data) > 0
        assert data.startswith(b"PK")

        prs = Presentation(io.BytesIO(data))
        text = _extract_all_text(prs)
        assert "Relatório de Seleção Teste" in text
        assert "Subtítulo do estudo" in text
        assert "Eng. Responsável" in text

    def test_to_pptx_includes_mandatory_notices(self) -> None:
        prs = Presentation(io.BytesIO(to_pptx(_sample_report(demo=True, own=True))))
        text = _extract_all_text(prs)

        assert LIMITATION_NOTICE in text
        assert REPRODUCIBILITY_NOTICE in text
        assert DEMO_DATA_NOTICE in text
        assert OWN_RECORD_NOTICE in text

    def test_to_pptx_omits_demo_notice_when_not_applicable(self) -> None:
        prs = Presentation(io.BytesIO(to_pptx(_sample_report(demo=False, own=False))))
        text = _extract_all_text(prs)

        assert LIMITATION_NOTICE in text
        assert DEMO_DATA_NOTICE not in text
        assert OWN_RECORD_NOTICE not in text

    def test_to_pptx_renders_tables_and_headers(self) -> None:
        prs = Presentation(io.BytesIO(to_pptx(_sample_report())))
        table_shapes = [
            shape
            for slide in prs.slides
            for shape in slide.shapes
            if shape.has_table
        ]
        assert len(table_shapes) >= 1

        table = table_shapes[0].table
        header_cells = [cell.text for cell in table.rows[0].cells]
        assert header_cells == ["Material", "Densidade", "Módulo de Young"]

        row1_cells = [cell.text for cell in table.rows[1].cells]
        assert row1_cells == ["Aço Demo A", "7800", "210"]

        row2_cells = [cell.text for cell in table.rows[2].cells]
        assert row2_cells == ["Cerâmica Demo D", "3900", "ausente"]

    def test_to_pptx_renders_empty_section_note(self) -> None:
        prs = Presentation(io.BytesIO(to_pptx(_sample_report())))
        text = _extract_all_text(prs)
        assert "Nenhum registro nesta seção." in text

    def test_to_pptx_renders_ai_narrative_when_present(self) -> None:
        report = _sample_report()
        report.narrative = ["Interpretação da IA sobre o ranking."]
        report.narrative_caveats = ["Ressalva número um."]
        report.narrative_note = "Nota sobre modelo."

        prs = Presentation(io.BytesIO(to_pptx(report)))
        text = _extract_all_text(prs)
        assert "Interpretação técnica (IA)" in text
        assert "Interpretação da IA sobre o ranking." in text
        assert "Ressalva número um." in text
        assert "Nota sobre modelo." in text

    def test_to_pptx_chunks_large_tables_across_slides(self) -> None:
        large_rows = [
            [f"Material {i}", float(i * 100), float(i * 10)]
            for i in range(25)
        ]
        report = Report(
            title="Relatório Extenso",
            subtitle="Teste de paginação de slides",
            notices=standard_notices(),
            sheets=[
                Sheet(
                    name="Muitos Candidatos",
                    header=["Material", "Densidade", "Módulo"],
                    rows=large_rows,
                )
            ],
        )
        prs = Presentation(io.BytesIO(to_pptx(report)))
        text = _extract_all_text(prs)
        assert "Muitos Candidatos (1/3)" in text
        assert "Muitos Candidatos (2/3)" in text
        assert "Muitos Candidatos (3/3)" in text


class TestPptxEndpoints:
    def test_catalogue_pptx_download_headers(self, client: TestClient) -> None:
        response = client.get("/api/exports/catalogo.pptx")
        assert response.status_code == 200, response.text
        assert response.headers["content-type"].startswith(PPTX_MEDIA_TYPE)
        assert "attachment" in response.headers["content-disposition"]
        assert ".pptx" in response.headers["content-disposition"]
        assert "catalogo-de-materiais.pptx" in response.headers["content-disposition"]

        prs = Presentation(io.BytesIO(response.content))
        text = _extract_all_text(prs)
        assert LIMITATION_NOTICE in text

    def test_study_pptx_download_headers(self, client: TestClient) -> None:
        response_post = client.post(
            "/api/selection/studies",
            json={
                "name": "Estudo PPTX Exportável",
                "function_text": "Tirante",
                "objective_text": "Minimizar massa",
                "combinator": "AND",
                "constraints": [],
            },
        )
        assert response_post.status_code == 201, response_post.text
        study_id = response_post.json()["id"]

        response = client.get(f"/api/exports/estudos/{study_id}.pptx")
        assert response.status_code == 200, response.text
        assert response.headers["content-type"].startswith(PPTX_MEDIA_TYPE)
        assert "attachment" in response.headers["content-disposition"]
        assert ".pptx" in response.headers["content-disposition"]
        assert "estudo-pptx-exportavel" in response.headers["content-disposition"]

        prs = Presentation(io.BytesIO(response.content))
        text = _extract_all_text(prs)
        assert LIMITATION_NOTICE in text

    def test_study_laudo_pptx_download_headers(self, client: TestClient) -> None:
        response_post = client.post(
            "/api/selection/studies",
            json={
                "name": "Estudo Laudo PPTX",
                "function_text": "Tirante",
                "objective_text": "Minimizar massa",
                "combinator": "AND",
                "constraints": [],
            },
        )
        assert response_post.status_code == 201, response_post.text
        study_id = response_post.json()["id"]

        response = client.get(
            f"/api/exports/estudos/{study_id}/laudo.pptx?responsavel=Eng.+Carlos"
        )
        assert response.status_code == 200, response.text
        assert response.headers["content-type"].startswith(PPTX_MEDIA_TYPE)
        assert "attachment" in response.headers["content-disposition"]
        assert ".pptx" in response.headers["content-disposition"]
        assert "laudo_estudo-laudo-pptx" in response.headers["content-disposition"]

        prs = Presentation(io.BytesIO(response.content))
        text = _extract_all_text(prs)
        assert LIMITATION_NOTICE in text
        assert "Eng. Carlos" in text

    def test_study_laudo_unsupported_format_rejected(self, client: TestClient) -> None:
        response = client.get("/api/exports/estudos/1/laudo.pdf")
        assert response.status_code == 400
        assert "Formato não suportado" in response.json()["detail"]
