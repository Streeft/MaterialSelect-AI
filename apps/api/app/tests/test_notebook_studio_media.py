"""The D-98 Studio tools end to end, through the HTTP API: the audio overview,
the video overview, the slide deck and the infographic — made in the background
by the simulated provider, cited, counted, laid out, saved as a note, exported
with their notices, checked figure by figure (a statistic by the strict rule and
by its unit) and kept from another student."""

from __future__ import annotations

import io
import json
from xml.etree import ElementTree

import pytest
from docx import Document
from pptx import Presentation

from app.ai.mock import MockAIProvider
from app.ai.provider import AIProvider
from app.ai.studio import CATALOG, read_studio, studio_system, studio_user
from app.exporters.report import LIMITATION_NOTICE
from app.exporters.studio import AI_NOTICE, SPEAKER_LABEL
from app.exporters.studio_pptx import NARRATION_LABEL
from app.notebooks import infographic
from app.services import studio_service
from app.services.studio_service import plain_text

NEW_TOOLS = ("audio", "video", "slides", "infographic")

STEEL = (
    "O aço carbono tem densidade de 7850 kg/m³ e módulo de elasticidade de 210 GPa. "
    "É o material estrutural mais usado na construção civil."
)
ALUMINIUM = (
    "O alumínio 6061 tem densidade de 2700 kg/m³, cerca de um terço da do aço. "
    "Por isso aparece em estruturas leves, como quadros de bicicleta."
)

MEDIA = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain; charset=utf-8",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "svg": "image/svg+xml",
}
EVERY_FORMAT = ("docx", "csv", "xlsx", "svg", "pptx", "txt")
SVG = "{http://www.w3.org/2000/svg}"


def _notebook(client, title: str = "Materiais estruturais", sources=None) -> int:
    """A notebook with one text source per ``(title, text)`` — one passage each."""
    notebook = client.post("/api/notebooks", json={"title": title}).json()
    for source_title, text in sources or [("Aula 1", STEEL), ("Aula 2", ALUMINIUM)]:
        response = client.post(
            f"/api/notebooks/{notebook['id']}/sources/text",
            json={"title": source_title, "text": text},
        )
        assert response.status_code == 201, response.text
    return notebook["id"]


def _one_passage(client) -> int:
    """Everything in passage [1]: what a scripted answer cites."""
    return _notebook(client, sources=[("Aula 1", f"{STEEL}\n\n{ALUMINIUM}")])


def _generate(client, notebook_id: int, **payload) -> dict:
    """202 now, made by the background task before ``post`` returns under the
    test client; the artifact read after it is finished."""
    response = client.post(f"/api/notebooks/{notebook_id}/studio", json=payload)
    assert response.status_code == 202, response.text
    assert response.json()["status"] == "gerando"
    artifact_id = response.json()["id"]
    return client.get(f"/api/notebooks/{notebook_id}/studio/{artifact_id}").json()


def _base(notebook_id: int, artifact: dict) -> str:
    return f"/api/notebooks/{notebook_id}/studio/{artifact['id']}"


def _json(value) -> object:
    return json.loads(json.dumps(value))


class _Spy(MockAIProvider):
    """The simulated provider, remembering every request it was handed."""

    def __init__(self) -> None:
        self.requests = []

    def studio(self, request):
        self.requests.append(request)
        return super().studio(request)


class _Scripted(AIProvider):
    """Answers from a script in the model's JSON shape, read as a real provider
    reads it; the last answer repeats."""

    name = "roteiro"
    simulated = True

    def __init__(self, answers: list[dict]) -> None:
        self.answers = answers
        self.calls = 0
        self.notes: list[str | None] = []

    def interpret(self, context):  # pragma: no cover - not used here
        raise NotImplementedError

    def explain(self, context):  # pragma: no cover - not used here
        raise NotImplementedError

    def studio(self, request):
        self.notes.append(request.retry_note)
        reply = self.answers[min(self.calls, len(self.answers) - 1)]
        self.calls += 1
        return read_studio(request, reply)


def _use(monkeypatch, provider: AIProvider) -> AIProvider:
    monkeypatch.setattr(studio_service, "get_provider", lambda _settings: provider)
    return provider


# --- making each tool with the simulated provider ---------------------------------


def _items(tool: str, content: dict) -> list[dict]:
    if tool == "audio":
        return content["lines"]
    if tool in ("slides", "video"):
        return content["slides"]
    return [*content["stats"], *content["points"], *content["steps"]]


@pytest.mark.parametrize("tool", NEW_TOOLS)
def test_each_new_tool_is_made_in_the_background_and_cited(client, tool):
    notebook_id = _notebook(client)
    artifact = _generate(client, notebook_id, tool=tool)
    assert artifact["status"] == "pronto", artifact["error"]
    assert artifact["error"] is None
    assert artifact["withheld"] == []
    assert artifact["exports"] == list(CATALOG[tool].exports)

    content = artifact["content"]
    items = _items(tool, content)
    assert items
    assert artifact["item_count"] == len(items)
    # The title travels in the body too, the same as the artifact's.
    assert content["title"] == artifact["title"]

    # Citations: renumbered 1..k in reading order, each one a copy of the passage.
    numbers = [c["number"] for c in artifact["citations"]]
    assert numbers == list(range(1, len(numbers) + 1))
    assert numbers == [1, 2]  # both sources are cited
    excerpts = {c["source_title"]: c["excerpt"] for c in artifact["citations"]}
    assert excerpts == {"Aula 1": STEEL, "Aula 2": ALUMINIUM}
    cited = [n for item in items for n in item["citations"]]
    assert list(dict.fromkeys(cited)) == [1, 2]  # first cited, first numbered

    listed = client.get(f"/api/notebooks/{notebook_id}/studio").json()
    (summary,) = listed["artifacts"]
    assert summary["id"] == artifact["id"]
    assert summary["item_count"] == artifact["item_count"]
    assert listed["usage"]["used"] == 1


def test_the_audio_is_a_script_between_two_hosts(client):
    artifact = _generate(client, _notebook(client), tool="audio")
    lines = artifact["content"]["lines"]
    assert [line["speaker"] for line in lines] == [1, 2, 1, 2]
    assert all(set(line) == {"speaker", "text", "citations"} for line in lines)
    # The second host reads the passage itself — copied, so grounded.
    assert lines[1]["text"].startswith("O aço carbono tem densidade de 7850 kg/m³")
    assert artifact["template"] == "conversa"


@pytest.mark.parametrize("tool", ["slides", "video"])
def test_slides_and_scenes_have_a_title_bullets_and_notes(client, tool):
    artifact = _generate(client, _notebook(client), tool=tool)
    slides = artifact["content"]["slides"]
    assert [s["title"] for s in slides] == ["Aula 1", "Aula 2"]
    for slide in slides:
        assert set(slide) == {"title", "bullets", "notes", "citations"}
        assert slide["bullets"] and slide["notes"]
        assert len(slide["citations"]) == 1


def test_the_infographic_has_stats_points_and_steps(client):
    artifact = _generate(client, _notebook(client), tool="infographic")
    content = artifact["content"]
    assert set(content) == {"title", "subtitle", "stats", "points", "steps"}
    assert [s["value"] for s in content["stats"]] == ["7850 kg/m³", "2700 kg/m³"]
    assert len(content["points"]) == 2 and len(content["steps"]) == 2
    assert artifact["item_count"] == 6


# --- the infographic's layout ------------------------------------------------------


@pytest.mark.parametrize("tool", ["audio", "video", "slides"])
def test_only_the_infographic_carries_an_infographic_layout(client, tool):
    artifact = _generate(client, _notebook(client), tool=tool)
    assert artifact["infographic"] is None
    assert artifact["layout"] is None


@pytest.mark.parametrize(
    ("fmt", "width"), [(None, 1200), ("paisagem", 1200), ("retrato", 800), ("quadrado", 1000)]
)
def test_the_infographic_layout_is_the_one_the_backend_computes(client, fmt, width):
    payload = {"tool": "infographic"} | ({"format": fmt} if fmt else {})
    artifact = _generate(client, _notebook(client), **payload)
    orientation = fmt or "paisagem"
    assert artifact["format"] == orientation
    assert artifact["layout"] is None  # that one is the mind map's
    sheet = artifact["infographic"]
    expected = infographic.layout(artifact["content"], artifact["format"]).to_dict()
    assert sheet == _json(expected)
    assert sheet["orientation"] == orientation and sheet["width"] == width
    kinds = [block["kind"] for block in sheet["blocks"]]
    assert kinds.count("stat") == 2 and kinds.count("point") == 2 and kinds.count("step") == 2
    assert len(sheet["connectors"]) == 1  # two steps, one arrow


# --- choices: refused, resolved, and what reaches the model ------------------------


@pytest.mark.parametrize(
    ("payload", "detail"),
    [
        ({"tool": "audio", "format": "podcast"}, "Resumo em áudio não tem formato para escolher."),
        (
            {"tool": "audio", "template": "entrevista"},
            "Modelo desconhecido para Resumo em áudio: entrevista. "
            "Aceitos: conversa, resumo, critica, debate.",
        ),
        (
            # No "Crie o seu" for the phase 4 tools.
            {"tool": "audio", "template": "personalizado"},
            "Modelo desconhecido para Resumo em áudio: personalizado. "
            "Aceitos: conversa, resumo, critica, debate.",
        ),
        (
            {"tool": "audio", "count": "mais"},
            "Quantidade desconhecida para Resumo em áudio: mais. Aceitos: curto, padrao, longo.",
        ),
        (
            {"tool": "audio", "difficulty": "facil"},
            "Resumo em áudio não tem dificuldade para escolher.",
        ),
        (
            {"tool": "video", "format": "longo"},
            "Formato desconhecido para Resumo em vídeo: longo. Aceitos: explicativo, resumo.",
        ),
        ({"tool": "video", "count": "padrao"}, "Resumo em vídeo não tem quantidade para escolher."),
        (
            {"tool": "video", "template": "conversa"},
            "Resumo em vídeo não tem modelo para escolher.",
        ),
        (
            {"tool": "video", "instructions": "Fale devagar."},
            "Resumo em vídeo não aceita instruções de modelo; use o campo de foco.",
        ),
        (
            {"tool": "slides", "format": "paisagem"},
            "Apresentação de slides não tem formato para escolher.",
        ),
        (
            {"tool": "slides", "count": "curto"},
            "Quantidade desconhecida para Apresentação de slides: curto. "
            "Aceitos: menos, padrao, mais.",
        ),
        (
            {"tool": "slides", "template": "personalizado"},
            "Modelo desconhecido para Apresentação de slides: personalizado. "
            "Aceitos: detalhada, apresentador.",
        ),
        (
            {"tool": "infographic", "format": "panorama"},
            "Formato desconhecido para Infográfico: panorama. "
            "Aceitos: paisagem, retrato, quadrado.",
        ),
        (
            {"tool": "infographic", "count": "mais"},
            "Quantidade desconhecida para Infográfico: mais. "
            "Aceitos: conciso, padrao, detalhado.",
        ),
        (
            {"tool": "infographic", "template": "cartaz"},
            "Infográfico não tem modelo para escolher.",
        ),
        (
            {"tool": "infographic", "instructions": "Use azul."},
            "Infográfico não aceita instruções de modelo; use o campo de foco.",
        ),
        ({"tool": "infographic", "columns": ["A"]}, "Infográfico não tem colunas para escolher."),
    ],
)
def test_a_choice_the_new_tool_does_not_have_is_refused_in_portuguese(client, payload, detail):
    notebook_id = _notebook(client)
    response = client.post(f"/api/notebooks/{notebook_id}/studio", json=payload)
    assert response.status_code == 400
    assert response.json()["detail"] == detail
    # Refused before anything was written or counted.
    listed = client.get(f"/api/notebooks/{notebook_id}/studio").json()
    assert listed["artifacts"] == [] and listed["usage"]["used"] == 0


@pytest.mark.parametrize(
    ("payload", "field"),
    [
        ({"tool": "video", "format": "x" * 33}, "format"),
        ({"tool": "slides", "count": "x" * 17}, "count"),
        ({"tool": "audio", "template": "x" * 65}, "template"),
        ({"tool": "audio", "instructions": "x" * 2001}, "instructions"),
        ({"tool": "podcast"}, "tool"),
    ],
)
def test_a_malformed_choice_is_a_422(client, payload, field):
    response = client.post(f"/api/notebooks/{_notebook(client)}/studio", json=payload)
    assert response.status_code == 422
    assert [error["loc"][-1] for error in response.json()["detail"]] == [field]


@pytest.mark.parametrize(
    ("payload", "amount"),
    [
        ({"tool": "video"}, 8),  # the first format, explicativo
        ({"tool": "video", "format": "explicativo"}, 8),
        ({"tool": "video", "format": "resumo"}, 5),
        ({"tool": "audio", "count": "curto"}, 12),
        ({"tool": "audio"}, 24),
        ({"tool": "audio", "count": "longo"}, 40),
        ({"tool": "slides", "count": "menos"}, 6),
        ({"tool": "slides", "count": "mais"}, 15),
        ({"tool": "infographic", "count": "conciso"}, 3),
        ({"tool": "infographic", "count": "detalhado"}, 7),
    ],
)
def test_the_amount_asked_for_is_stored_and_sent(client, monkeypatch, payload, amount):
    spy = _use(monkeypatch, _Spy())
    artifact = _generate(client, _notebook(client), **payload)
    assert artifact["status"] == "pronto", artifact["error"]
    assert artifact["options"]["amount"] == amount
    (request,) = spy.requests
    assert request.count == amount
    if payload["tool"] == "video":
        # The video's length is its format; it has no count of its own.
        assert artifact["options"]["count"] is None


@pytest.mark.parametrize(
    ("payload", "expected"),
    [
        ({"tool": "audio"}, CATALOG["audio"].template("conversa").instructions),
        (
            {"tool": "audio", "template": "debate"},
            CATALOG["audio"].template("debate").instructions,
        ),
        (
            {"tool": "audio", "template": "critica", "instructions": "Seja duro com as fontes."},
            "Seja duro com as fontes.",
        ),
        (
            {"tool": "slides", "template": "apresentador"},
            CATALOG["slides"].template("apresentador").instructions,
        ),
        ({"tool": "video", "format": "resumo"}, None),
        ({"tool": "infographic", "format": "retrato"}, None),
    ],
)
def test_instructions_reach_the_model_only_from_a_template_that_has_them(
    client, monkeypatch, payload, expected
):
    spy = _use(monkeypatch, _Spy())
    artifact = _generate(client, _notebook(client), **payload)
    assert artifact["status"] == "pronto", artifact["error"]
    (request,) = spy.requests
    assert request.instructions == expected
    assert artifact["options"]["instructions"] == expected
    if expected:
        assert CATALOG[payload["tool"]].template(request.template).instructions


def test_the_infographic_orientation_is_layout_only(client, monkeypatch):
    spy = _use(monkeypatch, _Spy())
    _generate(client, _notebook(client), tool="infographic", format="quadrado")
    (request,) = spy.requests
    # The request carries it for the service; the prompt never writes it (D-98).
    prompt = studio_system(request) + studio_user(request)
    assert "quadrado" not in prompt.lower()


# --- save as note -------------------------------------------------------------------


@pytest.mark.parametrize(
    ("tool", "marker"),
    [
        ("audio", "Apresentador(a) 1: "),
        ("slides", "Notas: "),
        ("video", "Narração: "),
        ("infographic", "DADOS EM DESTAQUE"),
    ],
)
def test_each_new_tool_saves_as_a_note_with_its_citations(client, tool, marker):
    notebook_id = _notebook(client)
    artifact = _generate(client, notebook_id, tool=tool)
    response = client.post(f"{_base(notebook_id, artifact)}/note")
    assert response.status_code == 201, response.text
    note = response.json()
    assert note["origin"] == "estudio"
    assert note["title"] == artifact["title"]
    assert note["body"] == plain_text(tool, artifact["content"])
    assert marker in note["body"] and "[1]" in note["body"] and "[2]" in note["body"]
    assert [(c["number"], c["excerpt"]) for c in note["citations"]] == [
        (c["number"], c["excerpt"]) for c in artifact["citations"]
    ]
    notebook = client.get(f"/api/notebooks/{notebook_id}").json()
    assert [n["id"] for n in notebook["notes"]] == [note["id"]]


# --- exports ------------------------------------------------------------------------


def _docx_text(data: bytes) -> str:
    return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)


def _pptx_text(data: bytes) -> tuple[str, str]:
    slides = list(Presentation(io.BytesIO(data)).slides)
    text = "\n".join(
        shape.text_frame.text for s in slides for shape in s.shapes if shape.has_text_frame
    )
    notes = "\n".join(s.notes_slide.notes_text_frame.text for s in slides if s.has_notes_slide)
    return text, notes


def _svg_text(data: bytes) -> str:
    root = ElementTree.fromstring(data.decode("utf-8"))
    # Wrapped over several <text> lines: compare word by word.
    return " ".join(" ".join(t.text or "" for t in root.iter(f"{SVG}text")).split())


def _words(sentence: str) -> str:
    return " ".join(sentence.split())


@pytest.mark.parametrize(
    ("tool", "fmt"),
    [
        (tool, fmt)
        for tool in NEW_TOOLS
        for fmt in CATALOG[tool].exports  # docx, txt, pptx, pptx, svg
    ],
)
def test_each_export_is_the_file_with_the_notices_and_the_references(client, tool, fmt):
    notebook_id = _notebook(client)
    artifact = _generate(client, notebook_id, tool=tool)
    response = client.get(f"{_base(notebook_id, artifact)}/export.{fmt}")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == MEDIA[fmt]
    assert response.headers["content-security-policy"] == "default-src 'none'"
    assert response.headers["content-disposition"] == (
        f'attachment; filename="{studio_service._slug(artifact["title"])}.{fmt}"'
    )
    data = response.content
    assert data

    if fmt == "docx":
        text = _docx_text(data)
        assert "Roteiro" in text
        assert f"{SPEAKER_LABEL} 1: Vamos ver o que a fonte “Aula 1”" in text
    elif fmt == "txt":
        text = data.decode("utf-8")
        assert text.index(LIMITATION_NOTICE) < text.index(f"{SPEAKER_LABEL} 1:")
    elif fmt == "pptx":
        text, notes = _pptx_text(data)
        assert "Aula 1" in text and "Aula 2" in text
        assert (NARRATION_LABEL in notes) is (tool == "video")
        assert "Fontes:" in notes
    else:
        text = _svg_text(data)
        assert "foreignObject" not in data.decode("utf-8")
        assert "7850 kg/m³" in text

    for sentence in (LIMITATION_NOTICE, AI_NOTICE):
        assert _words(sentence) in _words(text)
    assert "Referências" in text and "[1] Aula 1" in text and "[2] Aula 2" in text


@pytest.mark.parametrize(
    ("tool", "fmt"),
    [(tool, fmt) for tool in NEW_TOOLS for fmt in EVERY_FORMAT if fmt not in CATALOG[tool].exports],
)
def test_a_format_the_new_tool_does_not_export_is_refused_naming_the_accepted(client, tool, fmt):
    notebook_id = _notebook(client)
    artifact = _generate(client, notebook_id, tool=tool)
    response = client.get(f"{_base(notebook_id, artifact)}/export.{fmt}")
    assert response.status_code == 400
    spec = CATALOG[tool]
    accepted = ", ".join(e.upper() for e in spec.exports)
    assert response.json()["detail"] == (
        f"{spec.label} não se exporta em {fmt.upper()}. Aceitos: {accepted}."
    )


def test_a_format_no_tool_has_is_a_422(client):
    notebook_id = _notebook(client)
    artifact = _generate(client, notebook_id, tool="slides")
    assert client.get(f"{_base(notebook_id, artifact)}/export.pdf").status_code == 422


def test_the_exact_refusals_the_student_reads(client):
    notebook_id = _notebook(client)
    audio = _generate(client, notebook_id, tool="audio")
    response = client.get(f"{_base(notebook_id, audio)}/export.pptx")
    assert (
        response.json()["detail"] == "Resumo em áudio não se exporta em PPTX. Aceitos: DOCX, TXT."
    )
    board = _generate(client, notebook_id, tool="infographic")
    response = client.get(f"{_base(notebook_id, board)}/export.docx")
    assert response.json()["detail"] == "Infográfico não se exporta em DOCX. Aceitos: SVG."


# --- the figure check: a statistic by the strict rule and by its unit ---------------

#: 45% is in no passage; 210 MPa is — as 210 GPa, another unit.
_INVENTED_PERCENT = {"value": "45%", "label": "da produção mundial", "citations": [1]}
_WRONG_UNIT = {"value": "210 MPa", "label": "módulo de elasticidade do aço", "citations": [1]}
_GOOD_STAT = {"value": "7850 kg/m³", "label": "densidade do aço carbono", "citations": [1]}
_POINT = {"heading": "Leveza", "text": "O alumínio aparece em estruturas leves.", "citations": [1]}
_STEP = {"text": "Comparar a densidade dos dois metais", "citations": [1]}

PERCENT_SENTENCE = (
    "Um dado em destaque foi omitido porque citava números que não aparecem nos trechos "
    "citados: 45."
)
UNIT_SENTENCE = (
    "Um dado em destaque foi omitido porque sua unidade não aparece no trecho citado: 210 MPa."
)


def _board(*stats, points=(_POINT,), steps=(_STEP,)) -> dict:
    return {
        "title": "Aço e alumínio",
        "subtitle": "O que a aula diz",
        "stats": list(stats),
        "points": list(points),
        "steps": list(steps),
    }


def test_an_invented_statistic_is_withheld_and_the_rest_kept(client, monkeypatch):
    provider = _use(monkeypatch, _Scripted([_board(_INVENTED_PERCENT, _GOOD_STAT)]))
    notebook_id = _one_passage(client)
    artifact = _generate(client, notebook_id, tool="infographic")
    assert provider.calls == 2  # one retry, naming the figure
    assert provider.notes[0] is None and "45" in provider.notes[1]
    assert artifact["status"] == "pronto"
    content = artifact["content"]
    assert [s["value"] for s in content["stats"]] == ["7850 kg/m³"]
    assert content["points"] == [{**_POINT}]
    assert content["steps"] == [{**_STEP}]
    assert artifact["item_count"] == 3
    assert artifact["withheld"] == [PERCENT_SENTENCE]
    # The poster has only the stat that passed.
    stats = [b for b in artifact["infographic"]["blocks"] if b["kind"] == "stat"]
    assert [b["heading_lines"] for b in stats] == [["7850 kg/m³"]]
    # And the exported image says what was left out.
    svg = client.get(f"{_base(notebook_id, artifact)}/export.svg").content
    assert _words(PERCENT_SENTENCE) in _svg_text(svg)


def test_a_small_integer_does_not_ground_a_statistic_as_it_does_a_point(client, monkeypatch):
    # "45" in the student's focus would ground a point; a statistic, never.
    point = {"heading": "Foco", "text": "As 45 ligas do curso.", "citations": [1]}
    _use(monkeypatch, _Scripted([_board(_INVENTED_PERCENT, points=(point,))]))
    artifact = _generate(client, _one_passage(client), tool="infographic", topic="as 45 ligas")
    assert artifact["status"] == "pronto"
    assert artifact["content"]["stats"] == []
    assert [p["text"] for p in artifact["content"]["points"]] == ["As 45 ligas do curso."]
    assert artifact["withheld"] == [PERCENT_SENTENCE]


def test_a_statistic_in_another_unit_is_withheld(client, monkeypatch):
    provider = _use(monkeypatch, _Scripted([_board(_WRONG_UNIT, _GOOD_STAT)]))
    artifact = _generate(client, _one_passage(client), tool="infographic")
    assert provider.calls == 2 and "210 MPa" in provider.notes[1]
    assert artifact["status"] == "pronto"
    assert [s["value"] for s in artifact["content"]["stats"]] == ["7850 kg/m³"]
    assert artifact["withheld"] == [UNIT_SENTENCE]


def test_both_reasons_are_written_each_in_its_own_sentence(client, monkeypatch):
    _use(monkeypatch, _Scripted([_board(_INVENTED_PERCENT, _WRONG_UNIT, _GOOD_STAT)]))
    artifact = _generate(client, _one_passage(client), tool="infographic")
    assert sorted(artifact["withheld"]) == sorted([PERCENT_SENTENCE, UNIT_SENTENCE])
    assert artifact["item_count"] == 3


def test_a_clean_retry_withholds_nothing(client, monkeypatch):
    provider = _use(
        monkeypatch,
        _Scripted(
            [
                _board(_INVENTED_PERCENT, _WRONG_UNIT, _GOOD_STAT),
                _board(_GOOD_STAT),
            ]
        ),
    )
    artifact = _generate(client, _one_passage(client), tool="infographic")
    assert provider.calls == 2
    assert "45" in provider.notes[1] and "210 MPa" in provider.notes[1]
    assert artifact["status"] == "pronto"
    assert artifact["withheld"] == []
    assert [s["value"] for s in artifact["content"]["stats"]] == ["7850 kg/m³"]


def test_an_invented_headline_is_replaced_retried_and_said_so(client, monkeypatch):
    # The review's reproduction: "45%" above the stats band, at 30 px.
    board = {
        **_board(_GOOD_STAT),
        "title": "45% das falhas são por fadiga",
        "subtitle": "Em 12 anos, 80% dos casos",
    }
    provider = _use(monkeypatch, _Scripted([board]))
    notebook_id = _one_passage(client)
    artifact = _generate(client, notebook_id, tool="infographic", topic="45% das falhas")
    assert provider.calls == 2
    assert "O título e o subtítulo não citam trechos" in provider.notes[1]
    assert artifact["status"] == "pronto"
    assert artifact["title"] == CATALOG["infographic"].label
    assert artifact["content"]["subtitle"] == ""
    assert artifact["withheld"] == [
        "O título gerado foi trocado por um título neutro porque trazia números que não "
        "aparecem nos trechos citados: 45.",
        "O subtítulo gerado foi omitido porque trazia números que não aparecem nos "
        "trechos citados: 12, 80.",
    ]
    kinds = [b["kind"] for b in artifact["infographic"]["blocks"]]
    assert "subtitle" not in kinds
    title = next(b for b in artifact["infographic"]["blocks"] if b["kind"] == "title")
    assert title["heading_lines"] == [CATALOG["infographic"].label]
    svg = _svg_text(client.get(f"{_base(notebook_id, artifact)}/export.svg").content)
    assert "45%" not in svg and "80%" not in svg


def test_a_renamed_infographic_draws_the_new_title(client):
    notebook_id = _notebook(client)
    artifact = _generate(client, notebook_id, tool="infographic")
    base = _base(notebook_id, artifact)
    renamed = client.patch(base, json={"title": "Pôster da aula"}).json()
    title = next(b for b in renamed["infographic"]["blocks"] if b["kind"] == "title")
    assert title["heading_lines"] == ["Pôster da aula"]
    root = ElementTree.fromstring(client.get(f"{base}/export.svg").content.decode("utf-8"))
    strings = [t.text for t in root.iter(f"{SVG}text")]
    assert strings[0] == "Pôster da aula"


def test_every_statistic_withheld_keeps_the_points_and_drops_the_band(client, monkeypatch):
    _use(monkeypatch, _Scripted([_board(_INVENTED_PERCENT)]))
    artifact = _generate(client, _one_passage(client), tool="infographic")
    assert artifact["status"] == "pronto"
    assert artifact["content"]["stats"] == []
    assert artifact["item_count"] == 2
    kinds = [b["kind"] for b in artifact["infographic"]["blocks"]]
    assert "stat" not in kinds and "point" in kinds and "step" in kinds


def test_nothing_left_after_the_check_is_a_failure_with_the_reason(client, monkeypatch):
    _use(monkeypatch, _Scripted([_board(_INVENTED_PERCENT, points=(), steps=())]))
    notebook_id = _one_passage(client)
    artifact = _generate(client, notebook_id, tool="infographic")
    assert artifact["status"] == "falhou"
    assert artifact["error"] == (
        "A IA não devolveu nenhum item que passasse na conferência com as fontes. "
        "Números sem lastro nos trechos citados: 45."
    )
    assert artifact["content"] is None and artifact["infographic"] is None
    assert client.get(f"/api/notebooks/{notebook_id}/studio").json()["usage"]["used"] == 0
    # A failed artifact neither exports nor becomes a note.
    base = _base(notebook_id, artifact)
    assert client.get(f"{base}/export.svg").status_code == 400
    assert client.post(f"{base}/note").status_code == 400


@pytest.mark.parametrize(
    ("tool", "good", "bad", "sentence"),
    [
        (
            "audio",
            {
                "title": "Áudio",
                "lines": [{"speaker": 1, "text": "O aço tem 210 GPa.", "citations": [1]}],
            },
            {"speaker": 2, "text": "E o titânio tem 999 GPa.", "citations": [1]},
            "Uma fala foi omitida porque citava números que não aparecem nos trechos citados: 999.",
        ),
        (
            "slides",
            {
                "title": "Slides",
                "slides": [
                    {"title": "Aço", "bullets": ["7850 kg/m³"], "notes": "", "citations": [1]}
                ],
            },
            {"title": "Titânio", "bullets": ["999 GPa"], "notes": "Nota.", "citations": [1]},
            "Um slide foi omitido porque citava números que não aparecem nos trechos citados: 999.",
        ),
        (
            "video",
            {
                "title": "Vídeo",
                "slides": [
                    {"title": "Aço", "bullets": [], "notes": "O aço tem 210 GPa.", "citations": [1]}
                ],
            },
            {"title": "Titânio", "bullets": [], "notes": "Tem 999 GPa.", "citations": [1]},
            "Uma cena foi omitida porque citava números que não aparecem nos trechos citados: 999.",
        ),
    ],
)
def test_a_line_slide_or_scene_with_an_invented_figure_goes_whole(
    client, monkeypatch, tool, good, bad, sentence
):
    key = "lines" if tool == "audio" else "slides"
    answer = {**good, key: [*good[key], bad]}
    provider = _use(monkeypatch, _Scripted([answer]))
    artifact = _generate(client, _one_passage(client), tool=tool)
    assert provider.calls == 2 and "999" in provider.notes[1]
    assert artifact["status"] == "pronto"
    assert len(artifact["content"][key]) == 1
    assert artifact["item_count"] == 1
    assert artifact["withheld"] == [sentence]


# --- isolation ----------------------------------------------------------------------


@pytest.mark.parametrize("tool", NEW_TOOLS)
def test_another_students_new_artifact_is_not_found_anywhere(client, login_as, other_user, tool):
    with login_as(other_user):
        notebook_id = _notebook(client, "Caderno do Bruno")
        artifact = _generate(client, notebook_id, tool=tool)
        assert artifact["status"] == "pronto"
    base = _base(notebook_id, artifact)
    responses = [
        client.get(f"/api/notebooks/{notebook_id}/studio"),
        client.post(f"/api/notebooks/{notebook_id}/studio", json={"tool": tool}),
        client.get(base),
        client.patch(base, json={"title": "Meu"}),
        client.post(f"{base}/note"),
        *(client.get(f"{base}/export.{fmt}") for fmt in CATALOG[tool].exports),
        client.delete(base),
    ]
    assert [r.status_code for r in responses] == [404] * len(responses)
    # Still there for its owner.
    with login_as(other_user):
        assert client.get(base).json()["status"] == "pronto"
