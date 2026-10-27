"""The eco audit as a file: CSV, XLSX, HTML, DOCX, PPTX (D-172).

What is pinned here, because each is the way this export could quietly mislead:

* the file carries the audit's own figures, by component and by leg, with units;
* an absent figure is the word ``ausente`` with a written reason — never a zero,
  never a blank — and a refused podium says ``recusado``;
* every datum read carries its source, quality, licence and ``is_demo``;
* the escape is per format: a formula-looking name is an apostrophe-prefixed text
  in a spreadsheet and stays literal in HTML, where markup is what gets escaped.
"""

from __future__ import annotations

import csv
import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.exporters.cells import FORMULA_PREFIXES
from app.models.material import Material
from app.models.process import Process
from app.tests.removed_notices import assert_no_removed_notice

MOBILE_USE = {
    "model": "movel",
    "distance_km": 200000.0,
    "mobile_intensity": 0.0025,
    "life_years": 12.0,
    "carbon_per_energy": 0.07,
}

HOSTILE_NAMES = ['=HYPERLINK("http://evil.example","x")', "@SUM(1+1)", "+cmd|'/c calc'!A1"]
MARKUP = "<script>alert(1)</script>"


def _material(db: Session, like: str) -> Material:
    return db.execute(select(Material).where(Material.name.like(like))).scalars().first()


def _process_id(db: Session, slug: str) -> int:
    return db.execute(select(Process).where(Process.slug == slug)).scalars().one().id


def _component(db: Session, **overrides):
    payload = {
        "name": "Carcaça",
        "material_id": _material(db, "%Alumínio%").id,
        "process_id": _process_id(db, "fundicao-areia"),
        "unit_mass": 2.0,
        "quantity": 1,
        "recycled_fraction": 0.0,
        "end_of_life": "reciclagem",
    }
    payload.update(overrides)
    return payload


def _brief(db: Session, components=None, legs=None, use=None):
    return {
        "components": components or [_component(db)],
        "transport_legs": legs or [{"mode": "rodoviario", "distance_km": 1500.0}],
        "use": use or MOBILE_USE,
    }


def _with_absence(db: Session):
    """A carcass in aluminium and an insulator whose recycling figure is not catalogued."""
    ceramic = _component(
        db,
        name="Isolador",
        material_id=_material(db, "%Cerâmica%").id,
        process_id=_process_id(db, "prensagem-sinterizacao"),
    )
    return _brief(db, components=[_component(db), ceramic])


def _sheets(content: bytes) -> dict[str, list[list[object]]]:
    """Every worksheet of an XLSX as plain rows."""
    workbook = load_workbook(io.BytesIO(content))
    return {
        ws.title: [[cell.value for cell in row] for row in ws.iter_rows()]
        for ws in workbook.worksheets
    }


def _table(rows: list[list[object]], first_header: str) -> tuple[list[object], list[list[object]]]:
    """The header and body of one worksheet, found by its first column title."""
    start = next(i for i, row in enumerate(rows) if row and row[0] == first_header)
    return rows[start], [r for r in rows[start + 1 :] if any(c is not None for c in r)]


def _xlsx(client: TestClient, body) -> dict[str, list[list[object]]]:
    response = client.post("/api/eco/exportar.xlsx", json=body)
    assert response.status_code == 200, response.text
    return _sheets(response.content)


# --- the file carries the audit -------------------------------------------


def test_csv_is_an_attachment_with_every_section_and_the_units(
    client: TestClient, db_session: Session
) -> None:
    response = client.post("/api/eco/exportar.csv", json=_brief(db_session))
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("text/csv")
    assert response.headers["content-disposition"].startswith("attachment;")
    assert response.headers["x-content-type-options"] == "nosniff"

    text = response.content.decode("utf-8-sig")
    for section in (
        "Resumo",
        "Fases do produto",
        "Componentes",
        "Componentes por fase",
        "Trechos de transporte",
        "Premissas de uso",
        "Proveniência",
    ):
        assert section in text
    assert "MJ" in text and "kg de CO₂" in text
    assert "Carcaça" in text
    assert_no_removed_notice(text)


def test_the_file_shows_the_figures_the_audit_returns(
    client: TestClient, db_session: Session
) -> None:
    body = _brief(db_session)
    audit = client.post("/api/eco/auditar", json=body).json()
    sheets = _xlsx(client, body)

    _, rows = _table(sheets["Resumo"], "Item")
    summary = {row[0]: row for row in rows}
    assert summary["Energia total"][1] == pytest.approx(audit["total_energy"])
    assert summary["Energia total"][2] == "MJ"
    assert summary["Carbono total"][1] == pytest.approx(audit["total_carbon"])
    assert summary["Carbono total"][2] == "kg de CO₂"
    assert summary["Massa na peça"][1] == pytest.approx(audit["mass_in_part"])
    assert summary["Fase dominante em energia"][1].startswith(audit["energy_dominance"]["label"])

    header, rows = _table(sheets["Fases do produto"], "Fase")
    assert header[:5] == ["Fase", "Grandeza", "Valor", "Unidade", "Situação"]
    assert len(rows) == 10  # five phases, energy and carbon each
    by_key = {(r[0], r[1]): r for r in rows}
    for phase in audit["phases"]:
        assert by_key[(phase["label"], "Energia")][2] == pytest.approx(phase["energy"])
        assert by_key[(phase["label"], "Carbono")][2] == pytest.approx(phase["carbon"])


def test_each_component_and_each_leg_has_its_own_rows(
    client: TestClient, db_session: Session
) -> None:
    body = _brief(
        db_session,
        components=[
            _component(db_session, name="Carcaça"),
            _component(db_session, name="Parafusos", unit_mass=0.01, quantity=8),
        ],
        legs=[
            {"mode": "maritimo", "distance_km": 8000.0},
            {"mode": "rodoviario", "distance_km": 300.0},
        ],
    )
    audit = client.post("/api/eco/auditar", json=body).json()
    sheets = _xlsx(client, body)

    _, components = _table(sheets["Componentes"], "Componente")
    assert [c[0] for c in components] == ["Carcaça", "Parafusos"]
    assert components[1][3] == 8  # quantity
    assert components[1][5] == pytest.approx(0.08)  # mass in part, kg

    header, by_phase = _table(sheets["Componentes por fase"], "Componente")
    assert header[2:6] == ["Fase", "Grandeza", "Valor", "Unidade"]
    assert len(by_phase) == 2 * 3 * 2  # components x phases they own x quantities
    assert {r[2] for r in by_phase} == {"Material", "Manufatura", "Fim de vida"}
    carcass_material = next(
        r for r in by_phase if r[0] == "Carcaça" and r[2:4] == ["Material", "Energia"]
    )
    assert carcass_material[4] == pytest.approx(audit["components"][0]["phases"][0]["energy"])

    header, legs = _table(sheets["Trechos de transporte"], "Trecho")
    assert len(legs) == 4  # two legs x two quantities
    assert [leg[2] for leg in legs[::2]] == [8000.0, 300.0]
    assert legs[0][3] == pytest.approx(audit["mass_in_part"])
    assert {leg[6] for leg in legs} == {"MJ", "kg de CO₂"}
    assert legs[0][5] == pytest.approx(audit["transport_legs"][0]["energy"])


def test_every_datum_read_carries_source_quality_licence_and_demo_marker(
    client: TestClient, db_session: Session
) -> None:
    sheets = _xlsx(client, _brief(db_session))
    header, rows = _table(sheets["Proveniência"], "Escopo")
    assert header[6:11] == [
        "Qualidade do dado",
        "Fonte",
        "Licença",
        "Demonstrativo",
        "Registro",
    ]
    data = {r[1].split(" (")[0]: r for r in rows if r[0] == "Carcaça"}
    embodied = next(r for r in rows if "energia_incorporada" in r[1])
    assert isinstance(embodied[2], (int, float))
    assert embodied[3] == "MJ/kg"
    assert embodied[6] in ("MEDIDO", "IMPORTADO", "ESTIMADO")
    assert embodied[7] and embodied[8]
    assert embodied[9] in ("sim", "não")
    assert embodied[10] in ("Catálogo", "Próprio")
    assert data  # the component's own data is listed under its name
    # The two transport intensities of the leg, with the unit written out.
    leg_rows = [r for r in rows if str(r[0]).startswith("Trecho 1")]
    assert {r[3] for r in leg_rows} == {"MJ/(t·km)", "kg de CO₂/(t·km)"}
    # Never an empty cell in a provenance row.
    assert all(cell not in (None, "") for row in rows for cell in row)


def test_use_premises_are_the_readers_own_and_say_so(
    client: TestClient, db_session: Session
) -> None:
    use = {"model": "estatico", "power_watts": 60.0, "duty_cycle": 0.25, "life_years": 10.0}
    sheets = _xlsx(client, _brief(db_session, use=use))
    _, rows = _table(sheets["Premissas de uso"], "Premissa")
    premises = {r[0]: r for r in rows}
    assert premises["Potência"][1:3] == [60.0, "W"]
    assert premises["Potência"][3] == "Informado pelo usuário"
    # Not typed: written as absent, with what that means for the phase.
    carbon = premises["Carbono por energia em serviço"]
    assert carbon[1] == "ausente"
    assert "sem carbono" in carbon[3]


# --- absence is written, never zero or blank ------------------------------


def test_an_absent_figure_is_the_word_ausente_with_its_reason(
    client: TestClient, db_session: Session
) -> None:
    sheets = _xlsx(client, _with_absence(db_session))

    _, rows = _table(sheets["Componentes por fase"], "Componente")
    end_of_life = [r for r in rows if r[0] == "Isolador" and r[2] == "Fim de vida"]
    assert [r[4] for r in end_of_life] == ["ausente", "ausente"]
    assert [r[6] for r in end_of_life] == ["ausente", "ausente"]
    assert all("Dados ausentes" in str(r[8]) for r in end_of_life)
    # The catalogued component next to it keeps its numbers.
    assert all(isinstance(r[4], (int, float)) for r in rows if r[0] == "Carcaça")

    _, phases = _table(sheets["Fases do produto"], "Fase")
    product_end = [r for r in phases if r[0] == "Fim de vida"]
    assert [r[2] for r in product_end] == ["ausente", "ausente"]
    assert "Isolador" in str(product_end[0][6])

    _, summary = _table(sheets["Resumo"], "Item")
    by_item = {r[0]: r for r in summary}
    assert by_item["Energia total"][1] == "ausente"
    assert "Isolador" in by_item["Energia total"][3]
    assert by_item["Fase dominante em energia"][1] == "recusado"
    assert by_item["Fase dominante em energia"][3] not in ("—", "", None)
    assert by_item["Fase dominante em carbono"][1] == "recusado"


def test_an_uncatalogued_input_is_a_written_row_not_a_missing_one(
    client: TestClient, db_session: Session
) -> None:
    sheets = _xlsx(client, _with_absence(db_session))
    _, rows = _table(sheets["Proveniência"], "Escopo")
    missing = [r for r in rows if r[0] == "Isolador" and r[2] == "ausente"]
    assert missing, "the absent recycling figure must appear as a row"
    for row in missing:
        # Not catalogued at all, or catalogued and flagged absent: either way written.
        assert row[6] == "ausente" or row[6] in ("MEDIDO", "IMPORTADO", "ESTIMADO")
        assert "não foi tratado como zero" in row[11]
        assert all(cell not in (None, "") for cell in row)


def test_the_same_absence_is_written_in_html_and_csv(
    client: TestClient, db_session: Session
) -> None:
    body = _with_absence(db_session)
    html = client.post("/api/eco/exportar.html", json=body).text
    csv_text = client.post("/api/eco/exportar.csv", json=body).content.decode("utf-8-sig")
    for text in (html, csv_text):
        assert "ausente" in text
        assert "recusado" in text
        assert "Dados ausentes" in text


# --- escape is per format -------------------------------------------------


def _hostile_brief(db: Session):
    return _brief(
        db,
        components=[
            _component(db, name=HOSTILE_NAMES[0]),
            _component(db, name=HOSTILE_NAMES[1], unit_mass=1.0),
            _component(db, name=HOSTILE_NAMES[2], unit_mass=1.0),
            _component(db, name=MARKUP, unit_mass=1.0),
        ],
    )


def _dangerous(value: object) -> bool:
    return isinstance(value, str) and value.startswith(FORMULA_PREFIXES)


def test_formula_injection_is_neutralised_in_csv(client: TestClient, db_session: Session) -> None:
    response = client.post("/api/eco/exportar.csv", json=_hostile_brief(db_session))
    assert response.status_code == 200, response.text
    reader = csv.reader(io.StringIO(response.content.decode("utf-8-sig")))
    cells = [cell for row in reader for cell in row]
    assert not [c for c in cells if _dangerous(c)]
    for name in HOSTILE_NAMES:
        assert f"'{name}" in cells


def test_formula_injection_is_neutralised_in_every_xlsx_cell(
    client: TestClient, db_session: Session
) -> None:
    sheets = _xlsx(client, _hostile_brief(db_session))
    for rows in sheets.values():
        for row in rows:
            assert not [c for c in row if _dangerous(c)]
    names = [r[0] for r in _table(sheets["Componentes"], "Componente")[1]]
    assert [f"'{n}" for n in HOSTILE_NAMES] == names[:3]
    # A negative number stays numeric: it is never a formula.
    assert all(
        isinstance(r[5], (int, float)) for r in _table(sheets["Componentes"], "Componente")[1]
    )


def test_a_hostile_catalogue_name_is_neutralised_too(
    client: TestClient, db_session: Session
) -> None:
    """Not only what the reader types: a material name reaches the file the same way."""
    material = _material(db_session, "%Alumínio%")
    component = _component(db_session, name="")
    material.name = "=cmd|'/c calc'!A1"
    db_session.flush()
    sheets = _xlsx(client, _brief(db_session, components=[component]))
    rows = _table(sheets["Componentes"], "Componente")[1]
    assert rows[0][0] == "'=cmd|'/c calc'!A1"
    assert rows[0][1] == "'=cmd|'/c calc'!A1"


def test_html_escapes_markup_and_leaves_formulas_literal(
    client: TestClient, db_session: Session
) -> None:
    response = client.post("/api/eco/exportar.html", json=_hostile_brief(db_session))
    assert response.status_code == 200, response.text
    assert response.headers["content-security-policy"].startswith("default-src 'none'")
    assert response.headers["content-disposition"].startswith("inline;")
    html = response.text
    assert MARKUP not in html
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html
    # A leading "=" is inert in HTML: the spreadsheet apostrophe would be corruption.
    assert "=HYPERLINK(&quot;http://evil.example&quot;,&quot;x&quot;)" in html
    assert "'=HYPERLINK" not in html
    assert "<script" not in html.lower()
    assert_no_removed_notice(html)


def test_a_markup_name_is_inert_in_the_spreadsheet_too(
    client: TestClient, db_session: Session
) -> None:
    sheets = _xlsx(client, _hostile_brief(db_session))
    names = [r[0] for r in _table(sheets["Componentes"], "Componente")[1]]
    assert MARKUP in names  # markup is data in a sheet; nothing renders it


# --- the other formats and the refusals -----------------------------------


@pytest.mark.parametrize("fmt", ["docx", "pptx"])
def test_docx_and_pptx_open_as_office_packages(
    client: TestClient, db_session: Session, fmt: str
) -> None:
    response = client.post(f"/api/eco/exportar.{fmt}", json=_with_absence(db_session))
    assert response.status_code == 200, response.text
    assert response.headers["content-disposition"].startswith("attachment;")
    archive = zipfile.ZipFile(io.BytesIO(response.content))
    assert archive.testzip() is None
    inner = "word/document.xml" if fmt == "docx" else "ppt/presentation.xml"
    assert inner in archive.namelist()


def test_an_unsupported_format_is_refused(client: TestClient, db_session: Session) -> None:
    response = client.post("/api/eco/exportar.pdf", json=_brief(db_session))
    assert response.status_code == 400, response.text
    assert "csv" in response.text


def test_the_single_part_shape_exports_like_the_audit_does(
    client: TestClient, db_session: Session
) -> None:
    single = {
        "material_id": _material(db_session, "%Alumínio%").id,
        "process_id": _process_id(db_session, "fundicao-areia"),
        "part_mass": 2.0,
        "transport_mode": "rodoviario",
        "transport_distance_km": 1500.0,
        "use": MOBILE_USE,
    }
    sheets = _xlsx(client, single)
    _, rows = _table(sheets["Componentes"], "Componente")
    assert len(rows) == 1


def test_a_brief_the_audit_refuses_is_refused_by_the_export(
    client: TestClient, db_session: Session
) -> None:
    unknown = _brief(db_session, components=[_component(db_session, material_id=999_999)])
    assert client.post("/api/eco/exportar.csv", json=unknown).status_code == 404
    mixed = _brief(db_session)
    mixed["material_id"] = 1
    assert client.post("/api/eco/exportar.csv", json=mixed).status_code == 422


def test_the_export_is_deterministic(client: TestClient, db_session: Session) -> None:
    body = _with_absence(db_session)
    first = client.post("/api/eco/exportar.html", json=body).content
    second = client.post("/api/eco/exportar.html", json=body).content
    assert first == second
