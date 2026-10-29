"""D-177: decompression limits on uploaded archives; XML-illegal text in exports."""

from __future__ import annotations

import csv
import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.domain import zip_guard
from app.domain.errors import ValidationError
from app.exporters.docx import to_docx
from app.exporters.html import to_html
from app.exporters.pptx import to_pptx
from app.exporters.report import Report, Sheet
from app.exporters.spreadsheet import to_csv, to_xlsx
from app.importers.readers import read_xlsx
from app.knowledge.readers import read_upload

# --- B1: zip bombs ----------------------------------------------------------


class ParserReached(BaseException):
    """Not an Exception, so a reader's ``except Exception`` cannot swallow it."""


_CONTENT_TYPES = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" '
    'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/word/document.xml" ContentType="application/'
    'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
    "</Types>"
)
_RELS = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/'
    '2006/relationships/officeDocument" Target="word/document.xml"/>'
    "</Relationships>"
)
_PARA = "<w:p><w:r><w:t>" + "a" * 50 + "</w:t></w:r></w:p>"


def _bomb_docx(inflated_mb: int) -> bytes:
    """A DOCX whose document.xml deflates by a factor of hundreds."""
    body = _PARA * ((inflated_mb * 1024 * 1024) // len(_PARA))
    document = (
        '<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="http://schemas.'
        'openxmlformats.org/wordprocessingml/2006/main"><w:body>' + body + "</w:body></w:document>"
    )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        archive.writestr("[Content_Types].xml", _CONTENT_TYPES)
        archive.writestr("_rels/.rels", _RELS)
        archive.writestr("word/document.xml", document)
    return out.getvalue()


def _real_docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Polímeros", level=1)
    document.add_paragraph("O PEAD tem densidade de 950 kg/m³.")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


@pytest.fixture
def parser_must_not_run(monkeypatch):
    import docx

    def _refuse(*_args, **_kwargs):
        raise ParserReached("python-docx")

    monkeypatch.setattr(docx, "Document", _refuse)


def test_notebook_upload_refuses_docx_bomb_before_parsing(parser_must_not_run):
    data = _bomb_docx(30)
    assert len(data) < 200 * 1024  # small on the wire
    with pytest.raises(ValidationError, match="descompress|descomprimid"):
        read_upload("bomba.docx", data, max_pages=400)


def test_docx_member_ratio_alone_is_refused(parser_must_not_run):
    # 5 MB inflated: under the member ceiling, but compressed ~1000x.
    with pytest.raises(ValidationError, match="bomba de descompressão"):
        read_upload("bomba.docx", _bomb_docx(5), max_pages=400)


def test_a_legitimate_docx_is_still_read():
    extracted = read_upload("aula.docx", _real_docx(), max_pages=400)
    assert "950 kg/m³" in extracted.pages[0]


def test_import_xlsx_bomb_is_refused_before_openpyxl(monkeypatch):
    import app.importers.readers as readers

    def _refuse(*_args, **_kwargs):
        raise ParserReached("openpyxl")

    monkeypatch.setattr(readers, "load_workbook", _refuse)
    with pytest.raises(ValidationError, match="XLSX"):
        read_xlsx(_bomb_docx(30), max_rows=5000)


def test_a_legitimate_xlsx_is_still_imported():
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.append(["nome", "densidade"])
    workbook.active.append(["PEAD", 950])
    out = io.BytesIO()
    workbook.save(out)
    table = read_xlsx(out.getvalue(), max_rows=5000)
    assert table.headers == ["nome", "densidade"]
    assert table.rows == [["PEAD", 950]]


def _stored_zip(members: dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_STORED) as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)
    return out.getvalue()


def test_zip_guard_counts_entries(monkeypatch):
    monkeypatch.setattr(zip_guard, "MAX_ENTRIES", 3)
    data = _stored_zip({f"p{i}.xml": b"x" for i in range(4)})
    with pytest.raises(ValidationError, match="partes internas"):
        zip_guard.check_zip_archive(data, kind="DOCX")


def test_zip_guard_caps_one_member_and_the_total(monkeypatch):
    monkeypatch.setattr(zip_guard, "MAX_MEMBER_BYTES", 1000)
    monkeypatch.setattr(zip_guard, "MAX_TOTAL_BYTES", 1500)
    with pytest.raises(ValidationError, match="parte 'a.xml'"):
        zip_guard.check_zip_archive(_stored_zip({"a.xml": b"x" * 1001}), kind="DOCX")
    with pytest.raises(ValidationError, match="descomprimido"):
        zip_guard.check_zip_archive(
            _stored_zip({"a.xml": b"x" * 800, "b.xml": b"x" * 800}), kind="DOCX"
        )
    zip_guard.check_zip_archive(_stored_zip({"a.xml": b"x" * 700}), kind="DOCX")


def test_zip_guard_refuses_a_header_that_lies_about_the_size(monkeypatch):
    """The bytes are counted, not the central directory believed."""
    monkeypatch.setattr(zip_guard, "MAX_MEMBER_BYTES", 1000)
    data = bytearray(_stored_zip({"a.xml": b"x" * 5000}))
    # Patch the declared uncompressed size (local header and central directory)
    # down to 10; zipfile then stops at 10 bytes or fails the CRC, never 5000.
    size = (5000).to_bytes(4, "little")
    for offset in [i for i in range(len(data) - 3) if data[i : i + 4] == size]:
        data[offset : offset + 4] = (10).to_bytes(4, "little")
    with pytest.raises(ValidationError):
        zip_guard.check_zip_archive(bytes(data), kind="DOCX")


def test_zip_guard_refuses_what_is_not_a_zip():
    with pytest.raises(ValidationError, match="inválido"):
        zip_guard.check_zip_archive(b"PK\x03\x04 not really", kind="DOCX")


# --- B3: control characters in exports --------------------------------------

_DIRTY = "Aço\x0b1020"
_INJECTION = "\x01=HYPERLINK(1)"


def _report() -> Report:
    return Report(
        title="Estudo\x01com controle",
        subtitle="sub\x02",
        notices=["aviso\x1f"],
        sheets=[
            Sheet(
                name="Candi\x03datos",
                header=["Material\x04", "E"],
                rows=[[_DIRTY, 200.0], [_INJECTION, -1.5]],
                notes=["nota\x05\tcom tab"],
            )
        ],
        responsible="Eng.\x06",
        narrative=["texto\x07"],
        narrative_caveats=["ressalva\x08"],
        narrative_note="nota\x0e",
    )


def _has_illegal(text: str) -> bool:
    return any(ch in text for ch in ("\x01", "\x0b", "\x03", "\x07"))


def test_csv_drops_controls_and_keeps_the_formula_apostrophe():
    text = to_csv(_report())
    assert not _has_illegal(text)
    rows = list(csv.reader(io.StringIO(text.lstrip("﻿"))))
    cells = [cell for row in rows for cell in row]
    assert "Aço1020" in cells
    assert "'=HYPERLINK(1)" in cells
    assert "nota\tcom tab" in cells  # TAB is legal and kept


def test_xlsx_with_control_characters_renders():
    workbook = load_workbook(io.BytesIO(to_xlsx(_report())))
    values = [
        c for ws in workbook.worksheets for row in ws.iter_rows(values_only=True) for c in row
    ]
    assert "Aço1020" in values
    assert "'=HYPERLINK(1)" in values
    assert -1.5 in values  # negative numbers stay numeric
    assert "Candidatos" in workbook.sheetnames


def test_docx_with_control_characters_renders():
    from docx import Document

    document = Document(io.BytesIO(to_docx(_report())))
    text = "\n".join(p.text for p in document.paragraphs)
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    assert "Estudo" in text and not _has_illegal(text)
    assert "Aço1020" in cells
    # DOCX is not a spreadsheet: no apostrophe there.
    assert "=HYPERLINK(1)" in cells


def test_pptx_with_control_characters_has_no_escaped_garbage():
    from pptx import Presentation

    deck = Presentation(io.BytesIO(to_pptx(_report())))
    texts = []
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if shape.has_table:
                texts.extend(c.text for row in shape.table.rows for c in row.cells)
    joined = "\n".join(texts)
    assert "_x000" not in joined
    assert "Aço1020" in joined


def test_html_with_control_characters_drops_them_and_still_escapes():
    report = _report()
    report.sheets[0].rows.append(["<script>\x01", 1.0])
    page = to_html(report)
    assert not _has_illegal(page)
    assert "&lt;script&gt;" in page
    assert "'=HYPERLINK" not in page


# --- B3 end to end: the reported 500s ---------------------------------------

STUDY = {
    "universe": "material",
    "stages": [
        {
            "kind": "limit",
            "label": "Etapa 1",
            "root_group": {
                "operator": "AND",
                "constraints": [
                    {
                        "operator": "gt",
                        "property_slug": "modulo_young",
                        "value": 1.0,
                        "unit": "GPa",
                    }
                ],
            },
        }
    ],
    "index": {
        "name": "Viga leve",
        "expression": "sqrt(modulo_young) / densidade",
        "goal": "maximize",
    },
    "criteria": [{"key": "__index__", "weight": 1.0}],
}


def _lenient(client: TestClient) -> TestClient:
    return TestClient(client.app, raise_server_exceptions=False, cookies=client.cookies)


def _study(client: TestClient, name: str) -> int:
    response = client.post("/api/selection/studies", json={**STUDY, "name": name})
    assert response.status_code == 201, response.text
    return response.json()["id"]


def test_control_char_in_study_name_exports_xlsx(client: TestClient) -> None:
    sid = _study(client, "Estudo\x01com controle")
    response = _lenient(client).get(f"/api/exports/estudos/{sid}.xlsx")
    assert response.status_code == 200, response.text[:200]


def test_control_char_in_study_name_exports_laudo_docx(client: TestClient) -> None:
    sid = _study(client, "Estudo\x01com controle")
    response = _lenient(client).get(f"/api/exports/estudos/{sid}/laudo.docx")
    assert response.status_code == 200, response.text[:200]


def test_shared_material_with_control_char_does_not_break_catalogue_xlsx(
    client: TestClient,
) -> None:
    listing = client.get("/api/materials").json()
    items = listing["items"] if isinstance(listing, dict) else listing
    material_id = items[0]["id"]
    patched = client.patch(f"/api/materials/{material_id}", json={"name": _DIRTY})
    assert patched.status_code == 200, patched.text
    response = _lenient(client).get("/api/exports/catalogo.xlsx")
    assert response.status_code == 200, response.text[:200]
    workbook = load_workbook(io.BytesIO(response.content))
    values = [
        c for ws in workbook.worksheets for row in ws.iter_rows(values_only=True) for c in row
    ]
    assert "Aço1020" in values
    # The stored name is untouched: the strip is on the way out.
    assert client.get(f"/api/materials/{material_id}").json()["name"] == _DIRTY
