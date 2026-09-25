"""Text extraction from reference documents.

Returns the same shape for every format — :class:`ExtractedText`, a list of
pages — so the chunker never learns what a PDF is. Today only PDF is
the Cérebro's only format, while a notebook (D-90) also takes DOCX, TXT and
Markdown — read from **bytes**, because an upload is never written to disk.

``pypdf`` was chosen over ``pymupdf``: it is pure Python, small, and
permissively licensed, where pymupdf is AGPL and would put a copyleft term on a
project that has none. The trade is real and worth naming — pymupdf reads
two-column layouts and tables noticeably better, which is exactly the shape of a
textbook page. If extraction quality proves to be the bottleneck, this module is
the only thing that has to change.

Nothing here raises on a page it cannot read. A PDF of scanned images yields
empty pages, and empty is the honest answer: no OCR runs, so no text is
invented to fill the gap.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

from app.domain.errors import ValidationError

SUPPORTED_EXTENSIONS = {".pdf"}


@dataclass
class ExtractedText:
    """Text of one document, one string per page.

    ``pages`` is 0-indexed; page *numbers* shown to a reader are 1-based, and
    the conversion happens once, in the chunker.
    """

    pages: list[str] = field(default_factory=list)

    @property
    def page_count(self) -> int:
        return len(self.pages)

    @property
    def is_empty(self) -> bool:
        """True when no page yielded any text — a scan, or a broken file."""
        return not any(page.strip() for page in self.pages)


def read_pdf(path: Path | io.BytesIO) -> ExtractedText:
    """Extract text from a PDF, page by page.

    Raises:
        ValidationError: the file cannot be opened or parsed at all. A file that
            opens but holds no text is *not* an error — it returns empty pages,
            and the caller records that as a document with nothing to index.
    """
    from pypdf import PdfReader

    try:
        reader = PdfReader(path if isinstance(path, io.BytesIO) else str(path))
        pages = [(page.extract_text() or "") for page in reader.pages]
    except Exception as exc:  # pypdf raises many types for malformed files
        raise ValidationError(f"Não foi possível ler o PDF: {exc}") from exc

    return ExtractedText(pages=pages)


def extract_text(path: Path) -> ExtractedText:
    """Dispatch to the reader for ``path``'s extension."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return read_pdf(path)
    raise ValidationError(f"Formato não suportado para extração: {suffix or path.name}")


# --- uploads (D-90) --------------------------------------------------------

#: What a notebook accepts, by extension. The content is checked too: an
#: extension is a claim the uploader makes, the first bytes are a fact.
UPLOAD_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}

_PDF_MAGIC = b"%PDF-"
_ZIP_MAGIC = b"PK\x03\x04"  # a DOCX is a zip container


def read_upload(filename: str, data: bytes, *, max_pages: int | None = None) -> ExtractedText:
    """Extract text from an uploaded file's bytes.

    The extension picks the reader, and the first bytes must agree with it: a
    ``.pdf`` that is not a PDF is refused with that said, rather than handed to
    a parser that would fail with a stack of pypdf internals.

    Raises:
        ValidationError: unsupported type, content that contradicts the
            extension, an unreadable file, or more pages than ``max_pages``.
    """
    suffix = Path(filename).suffix.lower()
    if suffix not in UPLOAD_EXTENSIONS:
        raise ValidationError(
            f"Formato não aceito: {suffix or filename}. Envie PDF, DOCX, TXT ou Markdown."
        )
    if suffix == ".pdf":
        if not data.startswith(_PDF_MAGIC):
            raise ValidationError("O arquivo tem extensão .pdf, mas o conteúdo não é um PDF.")
        extracted = read_pdf(io.BytesIO(data))
    elif suffix == ".docx":
        if not data.startswith(_ZIP_MAGIC):
            raise ValidationError("O arquivo tem extensão .docx, mas o conteúdo não é um DOCX.")
        extracted = _read_docx(data)
    else:
        extracted = ExtractedText(pages=[decode_text(data)])

    if max_pages is not None and extracted.page_count > max_pages:
        raise ValidationError(
            f"O arquivo tem {extracted.page_count} páginas; o limite por fonte é "
            f"{max_pages}. Divida o documento em partes."
        )
    return extracted


def decode_text(data: bytes) -> str:
    """Bytes of a plain-text file as text, whatever its encoding.

    A file with a NUL byte is binary under a text extension, and decoding it
    anyway would index garbage that no answer could ever cite honestly.
    """
    if b"\x00" in data[:4096]:
        raise ValidationError("O arquivo não parece ser texto: contém bytes binários.")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    from charset_normalizer import from_bytes

    best = from_bytes(data).best()
    if best is None:
        raise ValidationError("Não foi possível identificar a codificação do texto.")
    return str(best)


def _read_docx(data: bytes) -> ExtractedText:
    """A DOCX as one page: Word has no fixed pages, only a flow of paragraphs.

    Headings stay on their own line, so the chunker still recognises them, and
    tables are read row by row with cells separated by " | " — a value kept on
    the same line as its row label is a value that can still be cited.
    """
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several types for a bad zip
        raise ValidationError(f"Não foi possível ler o DOCX: {exc}") from exc

    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        lines.append("")
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return ExtractedText(pages=["\n\n".join(line for line in lines if line.strip())])
