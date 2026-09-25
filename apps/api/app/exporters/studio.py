"""Exporting what the Studio made (D-94): DOCX, CSV, XLSX and SVG.

The same two promises as every other export of this project:

* **Every file carries the limitation notice** (item 5 of the proposal), and a
  second one saying what an AI-made file is: made from the notebook's sources,
  each item citing the passage it rests on — to be checked there.
* **The escape is by format.** Spreadsheets go through ``Report``/``Sheet`` and
  :mod:`app.exporters.spreadsheet`, which neutralises formula injection cell by
  cell; the SVG escapes every string with ``html.escape``; python-docx writes
  text runs, never markup.

The mind map is drawn from the layout the screen draws
(``app.notebooks.mindmap``, handed over on ``ArtifactOut.layout``) — never laid
out again here.
"""

from __future__ import annotations

import io
from html import escape

from docx import Document
from docx.shared import Pt, RGBColor

from app.exporters.report import LIMITATION_NOTICE, Report, Sheet
from app.exporters.spreadsheet import to_csv, to_xlsx
from app.notebooks.mindmap import wrap
from app.notebooks.studio_content import cell_text
from app.schemas.notebook import ArtifactOut, CitationOut

AI_NOTICE = (
    "Gerado por IA a partir das fontes do caderno. Cada item cita o trecho em que se "
    "apoia; confira nas fontes antes de usar. Números só aparecem quando estão "
    "escritos no trecho citado."
)

NOTICES = [LIMITATION_NOTICE, AI_NOTICE]

_MEDIA = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "svg": "image/svg+xml",
}

_INK = RGBColor(15, 23, 42)
_SUBTLE = RGBColor(100, 116, 139)


def render(artifact: ArtifactOut, fmt: str, notebook_title: str) -> tuple[bytes, str]:
    """One ready artifact as the bytes of ``fmt`` and its media type."""
    subtitle = f"Caderno “{notebook_title}” — Estúdio do MaterialSelect AI"
    if fmt == "svg":
        return mindmap_svg(artifact, subtitle).encode("utf-8"), _MEDIA["svg"]
    if fmt == "docx":
        return _docx(artifact, subtitle), _MEDIA["docx"]
    report = _report(artifact, subtitle)
    if fmt == "csv":
        return to_csv(report).encode("utf-8"), _MEDIA["csv"]
    if fmt == "xlsx":
        return to_xlsx(report), _MEDIA["xlsx"]
    raise ValueError(f"Formato desconhecido: {fmt}")


def _marks(citations: list[int]) -> str:
    return " ".join(f"[{n}]" for n in citations)


def _reference(citation: CitationOut) -> str:
    where = citation.source_title
    if citation.page_start is not None:
        pages = (
            f"p. {citation.page_start}"
            if citation.page_end in (None, citation.page_start)
            else f"p. {citation.page_start}–{citation.page_end}"
        )
        where += f", {pages}"
    if citation.heading:
        where += f" — {citation.heading}"
    return where


# --- spreadsheets --------------------------------------------------------------


def _references_sheet(artifact: ArtifactOut) -> Sheet:
    return Sheet(
        name="Referências",
        header=["Nº", "Fonte", "Trecho citado"],
        rows=[[c.number, _reference(c), c.excerpt] for c in artifact.citations],
        notes=[] if artifact.citations else ["Nenhum trecho citado."],
    )


def _report(artifact: ArtifactOut, subtitle: str) -> Report:
    body = artifact.content or {}
    notes = list(artifact.withheld)
    if artifact.tool == "table":
        columns = body.get("columns", [])
        rows = []
        for row in body.get("rows", []):
            cells = row["cells"]
            sources = "; ".join(
                f"{column} {_marks(cell['citations'])}"
                for column, cell in zip(columns, cells, strict=False)
                if cell["citations"]
            )
            rows.append([cell_text(cell) for cell in cells] + [sources])
        main = Sheet(name="Tabela", header=[*columns, "Fontes"], rows=rows, notes=notes)
    elif artifact.tool == "flashcards":
        main = Sheet(
            name="Cartões",
            header=["Frente", "Verso", "Fontes"],
            rows=[
                [card["front"], card["back"], _marks(card["citations"])]
                for card in body.get("cards", [])
            ],
            notes=notes,
        )
    elif artifact.tool == "quiz":
        rows = []
        for number, question in enumerate(body.get("questions", []), start=1):
            options = question["options"]
            answer = question["answer_index"]
            rows.append(
                [
                    number,
                    question["prompt"],
                    " | ".join(f"{chr(97 + i)}) {o}" for i, o in enumerate(options)),
                    f"{chr(97 + answer)}) {options[answer]}",
                    question["explanation"],
                    _marks(question["citations"]),
                ]
            )
        main = Sheet(
            name="Questões",
            header=["Nº", "Enunciado", "Alternativas", "Resposta", "Explicação", "Fontes"],
            rows=rows,
            notes=notes,
        )
    else:  # pragma: no cover - the service offers only the tool's formats
        raise ValueError(f"{artifact.tool} não tem planilha")
    return Report(
        title=artifact.title,
        subtitle=subtitle,
        notices=list(NOTICES),
        sheets=[main, _references_sheet(artifact)],
    )


# --- DOCX ----------------------------------------------------------------------


def _docx(artifact: ArtifactOut, subtitle: str) -> bytes:
    doc = Document()
    title = doc.add_paragraph()
    run = title.add_run(artifact.title)
    run.font.size, run.font.bold, run.font.color.rgb = Pt(20), True, _INK
    sub = doc.add_paragraph()
    run = sub.add_run(subtitle)
    run.font.size, run.font.color.rgb = Pt(11), _SUBTLE
    for notice in [*NOTICES, *artifact.withheld]:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(notice)
        run.font.size, run.font.italic, run.font.color.rgb = Pt(9.5), True, _SUBTLE

    body = artifact.content or {}
    if artifact.tool == "report":
        bullets = artifact.format == "topicos"
        for section in body.get("sections", []):
            doc.add_heading(section["heading"], level=2)
            for paragraph in section["paragraphs"]:
                doc.add_paragraph(
                    f"{paragraph['text']} {_marks(paragraph['citations'])}".strip(),
                    style="List Bullet" if bullets else None,
                )
    elif artifact.tool == "flashcards":
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Frente", "Verso"
        for card in body.get("cards", []):
            cells = table.add_row().cells
            cells[0].text = card["front"]
            cells[1].text = f"{card['back']} {_marks(card['citations'])}".strip()
    elif artifact.tool == "quiz":
        questions = body.get("questions", [])
        doc.add_heading("Questões", level=2)
        for number, question in enumerate(questions, start=1):
            doc.add_paragraph(f"{number}. {question['prompt']}")
            for i, option in enumerate(question["options"]):
                doc.add_paragraph(f"{chr(97 + i)}) {option}").paragraph_format.left_indent = Pt(18)
        doc.add_heading("Gabarito", level=2)
        for number, question in enumerate(questions, start=1):
            answer = question["answer_index"]
            text = f"{number}. {chr(97 + answer)}) {question['options'][answer]}"
            if question["explanation"]:
                text += f" — {question['explanation']}"
            doc.add_paragraph(f"{text} {_marks(question['citations'])}".strip())
    else:  # pragma: no cover - the service offers only the tool's formats
        raise ValueError(f"{artifact.tool} não tem DOCX")

    doc.add_heading("Referências", level=2)
    if not artifact.citations:
        doc.add_paragraph("Nenhum trecho citado.")
    for citation in artifact.citations:
        paragraph = doc.add_paragraph()
        paragraph.add_run(f"[{citation.number}] {_reference(citation)}").bold = True
        excerpt = paragraph.add_run(f"\n{citation.excerpt}")
        excerpt.font.size, excerpt.font.color.rgb = Pt(9), _SUBTLE

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# --- SVG -----------------------------------------------------------------------

_SVG_INK = "#0f172a"
_SVG_SUBTLE = "#475569"
_SVG_EDGE = "#94a3b8"
_SVG_SURFACE = "#ffffff"
#: Fill and stroke per depth: the central theme strongest, leaves lightest.
_SVG_LEVELS = (
    ("#1e3a8a", "#1e3a8a", "#ffffff"),
    ("#dbeafe", "#1d4ed8", "#0f172a"),
    ("#eef2ff", "#6366f1", "#0f172a"),
    ("#f8fafc", "#94a3b8", "#0f172a"),
)
_HEADER_LINE = 16.0


def _t(text: object) -> str:
    return escape(str(text), quote=True)


def mindmap_svg(artifact: ArtifactOut, subtitle: str) -> str:
    """The mind map as a standalone SVG, drawn from the backend's layout."""
    layout = artifact.layout
    if layout is None:  # pragma: no cover - a ready mind map always has one
        raise ValueError("Mapa mental sem layout")
    width = max(layout.width, 640.0)
    chars = int((width - 32) // 6.4)
    header: list[tuple[str, int, str, str]] = [(artifact.title, 16, "600", _SVG_INK)]
    header += [(line, 11, "400", _SVG_SUBTLE) for line in wrap(subtitle, chars, 2)]
    for notice in [*NOTICES, *artifact.withheld]:
        header += [(line, 10, "400", _SVG_SUBTLE) for line in wrap(notice, chars, 4)]
    top = 16 + len(header) * _HEADER_LINE + 8
    # The references go below the map: an SVG carried alone still says where
    # every bracketed number points.
    references = [
        line
        for citation in artifact.citations
        for line in wrap(f"[{citation.number}] {_reference(citation)}", chars, 2)
    ]
    bottom = top + layout.height
    height = bottom + (len(references) * 14 + 24 if references else 0)

    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width:g} {height:g}" '
        f'width="{width:g}" height="{height:g}" role="img" aria-label="{_t(artifact.title)}" '
        'font-family="ui-sans-serif, system-ui, Segoe UI, Roboto, Arial, sans-serif">',
        f"<title>{_t(artifact.title)}</title>",
        f"<desc>{_t(AI_NOTICE)}</desc>",
        f'<rect x="0" y="0" width="{width:g}" height="{height:g}" fill="{_SVG_SURFACE}"/>',
    ]
    for index, (line, size, weight, colour) in enumerate(header):
        y = 16 + (index + 1) * _HEADER_LINE - 4
        parts.append(
            f'<text x="16" y="{y:g}" font-size="{size}" font-weight="{weight}" '
            f'fill="{colour}">{_t(line)}</text>'
        )
    parts.append(f'<g transform="translate(0 {top:g})">')
    for edge in layout.edges:
        bend = (edge.x2 - edge.x1) / 2
        parts.append(
            f'<path d="M {edge.x1:g} {edge.y1:g} C {edge.x1 + bend:g} {edge.y1:g}, '
            f'{edge.x2 - bend:g} {edge.y2:g}, {edge.x2:g} {edge.y2:g}" fill="none" '
            f'stroke="{_SVG_EDGE}" stroke-width="1.5"/>'
        )
    for node in layout.nodes:
        fill, stroke, ink = _SVG_LEVELS[min(node.depth, len(_SVG_LEVELS) - 1)]
        parts.append(
            f'<rect x="{node.x:g}" y="{node.y:g}" width="{node.width:g}" '
            f'height="{node.height:g}" rx="10" fill="{fill}" stroke="{stroke}"/>'
        )
        marks = _marks(node.citations)
        lines = list(node.lines)
        for index, line in enumerate(lines):
            y = node.y + 9 + 13 + index * 17
            parts.append(
                f'<text x="{node.x + 12:g}" y="{y:g}" font-size="13" fill="{ink}">'
                f"{_t(line)}</text>"
            )
        if marks:
            parts.append(
                f'<text x="{node.x + node.width - 6:g}" y="{node.y + node.height - 5:g}" '
                f'font-size="9" text-anchor="end" fill="{ink}" opacity="0.7">{_t(marks)}</text>'
            )
    parts.append("</g>")
    for index, line in enumerate(references):
        parts.append(
            f'<text x="16" y="{bottom + 8 + (index + 1) * 14:g}" font-size="10" '
            f'fill="{_SVG_SUBTLE}">{_t(line)}</text>'
        )
    parts.append("</svg>")
    return "".join(parts)
