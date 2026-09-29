"""The Studio's slide deck and video summary as a .pptx (D-98).

Both tools produce the same content — ``{"title", "slides": [{"title",
"bullets", "notes", "citations"}]}`` (``app.ai.studio.read_deck``) — and both
export as a deck. The video's ``notes`` are the narration the screen speaks, so
its speaker notes are labelled "Narração:" and a presenter can read them aloud.

The deck, in order:

1. a title slide, saying the deck was made by AI from the notebook's sources;
2. one slide per content slide: title, bullets, and a small box with the
   slide's citation marks ("[1] [3]") — the notes carry the narration or
   presenter notes, then "Fontes:" with the reference of every cited passage;
3. the references, eight per slide, with the source's address and the
   attribution its licence asks for (D-97) — a CC BY-SA passage keeps its
   credit in the file that quotes it. No citation at all is a slide that
   says so (D-24), never a missing slide;
4. the "Avisos" slide, drawn by :func:`app.exporters.pptx.add_notices_slide`;
5. what was withheld, when something was, on a slide of its own.

**The escape for this format is the control-character strip.** python-pptx
writes text runs and escapes markup itself, so a ``<script>`` stays literal
text; but it does not refuse a character XML cannot carry — it silently
writes ``_x0000_`` in its place. Every string goes through :func:`_clean`
before it reaches a shape.

This module does not import :mod:`app.exporters.studio` (which dispatches to
it); the notices are handed over by the caller.
"""

from __future__ import annotations

import re
from collections.abc import Sequence
from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.slide import Slide
from pptx.util import Inches, Pt

from app.exporters.pptx import MARGIN, SLIDE_HEIGHT, SLIDE_WIDTH, add_notices_slide
from app.schemas.notebook import CitationOut

DECK_SUBTITLE = "Gerado por IA a partir das fontes do caderno"
NARRATION_LABEL = "Narração:"
SOURCES_LABEL = "Fontes:"
REFERENCES_TITLE = "Referências"
NO_CITATIONS = "Nenhum trecho citado."
WITHHELD_TITLE = "Parte do que foi gerado foi omitida"
#: References per slide: at three lines each, eight fill the page at 12 pt.
REFERENCES_PER_SLIDE = 8

#: Characters XML cannot carry (the same set as ``studio._NOT_XML``).
_NOT_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")

_SUBTLE = RGBColor(71, 85, 105)
_WIDTH = SLIDE_WIDTH - 2 * MARGIN


def _clean(text: object) -> str:
    """``text`` as XML-legal text; the illegal characters are dropped."""
    return _NOT_XML.sub("", "" if text is None else str(text))


def _where(citation: CitationOut) -> str:
    """ "Title, p. X–Y — heading", as ``studio._reference`` writes it."""
    where = citation.source_title
    if citation.page_start is not None:
        pages = (
            f"p. {citation.page_start}"
            if citation.page_end in (None, citation.page_start)
            else f"p. {citation.page_start}–{citation.page_end}"
        )
        where += f", {pages}"
    if citation.heading:
        where += f" — {citation.heading}"
    return _clean(where)


def _unique(numbers: Sequence[object]) -> list[int]:
    seen: list[int] = []
    for number in numbers:
        if isinstance(number, int) and not isinstance(number, bool) and number not in seen:
            seen.append(number)
    return seen


def deck_to_pptx(
    title: str,
    content: dict | None,
    citations: Sequence[CitationOut],
    withheld: Sequence[str],
    *,
    notices: Sequence[str],
    narration: bool = False,
    subtitle: str | None = None,
) -> bytes:
    """A slides/video artifact as the bytes of a 16:9 .pptx.

    ``content`` is the artifact's stored content (``None`` or without
    ``slides`` gives a deck with no content slide). ``notices`` are the
    closing slide's paragraphs — the Studio passes ``studio.NOTICES``.
    ``narration`` is True for the video: the speaker notes are then labelled
    "Narração:". ``subtitle`` (the notebook line) goes under the AI line on
    the title slide.
    """
    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT
    by_number = {citation.number: citation for citation in citations}

    _title_slide(presentation, title, subtitle)
    for slide in (content or {}).get("slides", []):
        _content_slide(presentation, slide, by_number, narration)
    _reference_slides(presentation, list(citations))
    add_notices_slide(presentation, [_clean(notice) for notice in notices] or [""])
    if withheld:
        _text_slide(presentation, WITHHELD_TITLE, [_clean(line) for line in withheld])

    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _place_title(slide: Slide, text: str, top: int = Inches(0.4)) -> None:
    heading = slide.shapes.title
    heading.text = _clean(text)
    heading.left, heading.top, heading.width, heading.height = MARGIN, top, _WIDTH, Inches(1.1)
    for paragraph in heading.text_frame.paragraphs:
        paragraph.font.size = Pt(32)


def _title_slide(presentation: Presentation, title: str, subtitle: str | None) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    _place_title(slide, title, top=Inches(2.2))
    heading = slide.shapes.title
    heading.height = Inches(1.6)
    for paragraph in heading.text_frame.paragraphs:
        paragraph.font.size = Pt(40)
    below = slide.placeholders[1]
    below.left, below.top, below.width, below.height = MARGIN, Inches(4.0), _WIDTH, Inches(1.4)
    frame = below.text_frame
    frame.text = DECK_SUBTITLE
    if subtitle:
        frame.add_paragraph().text = _clean(subtitle)
    for paragraph in frame.paragraphs:
        paragraph.font.size = Pt(20)
        paragraph.font.color.rgb = _SUBTLE


def _content_slide(
    presentation: Presentation,
    content: dict,
    by_number: dict[int, CitationOut],
    narration: bool,
) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    _place_title(slide, content.get("title", ""))

    bullets = [_clean(bullet) for bullet in content.get("bullets", [])]
    body = slide.placeholders[1]
    body.left, body.top = MARGIN, Inches(1.6)
    body.width, body.height = _WIDTH, SLIDE_HEIGHT - Inches(2.5)
    if bullets:
        frame = body.text_frame
        frame.word_wrap = True
        frame.text = bullets[0]
        for bullet in bullets[1:]:
            frame.add_paragraph().text = bullet
        size = Pt(24) if len(bullets) <= 3 else Pt(20) if len(bullets) <= 5 else Pt(18)
        for paragraph in frame.paragraphs:
            paragraph.font.size = size
            paragraph.space_after = Pt(6)
    else:
        # A narration-only scene: an empty placeholder would show its prompt
        # text ("Click to add text") in an editor.
        body.element.getparent().remove(body.element)

    numbers = _unique(content.get("citations", []))
    if numbers:
        marks = slide.shapes.add_textbox(
            MARGIN, SLIDE_HEIGHT - Inches(0.8), _WIDTH, Inches(0.5)
        ).text_frame
        marks.text = " ".join(f"[{n}]" for n in numbers)
        marks.paragraphs[0].font.size = Pt(12)
        marks.paragraphs[0].font.color.rgb = _SUBTLE

    notes = slide.notes_slide.notes_text_frame
    spoken = _clean(content.get("notes", "")).strip()
    first = f"{NARRATION_LABEL} {spoken}" if narration and spoken else spoken
    notes.text = first
    cited = [by_number[n] for n in numbers if n in by_number]
    if cited:
        if first:
            notes.add_paragraph().text = ""
        notes.add_paragraph().text = SOURCES_LABEL
        for citation in cited:
            line = f"[{citation.number}] {_where(citation)}"
            url = _clean(citation.source_url).strip()
            if url:
                line += f" — {url}"
            notes.add_paragraph().text = line


def _reference_slides(presentation: Presentation, citations: list[CitationOut]) -> None:
    if not citations:
        _text_slide(presentation, REFERENCES_TITLE, [NO_CITATIONS])
        return
    pages = [
        citations[i : i + REFERENCES_PER_SLIDE]
        for i in range(0, len(citations), REFERENCES_PER_SLIDE)
    ]
    for index, page in enumerate(pages, start=1):
        heading = (
            REFERENCES_TITLE if len(pages) == 1 else f"{REFERENCES_TITLE} ({index}/{len(pages)})"
        )
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])  # title only
        _place_title(slide, heading)
        frame = slide.shapes.add_textbox(
            MARGIN, Inches(1.5), _WIDTH, SLIDE_HEIGHT - Inches(1.9)
        ).text_frame
        frame.word_wrap = True
        first = True
        for citation in page:
            paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            run = paragraph.add_run()
            run.text = f"[{citation.number}] {_where(citation)}"
            run.font.size, run.font.bold = Pt(12), True
            # Address and credit under the reference, never shortened: a
            # licence credit cut short is not the credit the licence asks for.
            for extra in (citation.source_url, citation.source_attribution):
                text = _clean(extra).strip()
                if text:
                    line = frame.add_paragraph()
                    line.text = text
                    line.font.size = Pt(10)
                    line.font.color.rgb = _SUBTLE


def _text_slide(presentation: Presentation, heading: str, lines: list[str]) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])  # title only
    _place_title(slide, heading)
    frame = slide.shapes.add_textbox(
        MARGIN, Inches(1.5), _WIDTH, SLIDE_HEIGHT - Inches(1.9)
    ).text_frame
    frame.word_wrap = True
    frame.text = lines[0] if lines else ""
    for line in lines[1:]:
        frame.add_paragraph().text = line
    for paragraph in frame.paragraphs:
        paragraph.font.size = Pt(16)
